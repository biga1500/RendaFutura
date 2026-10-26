from copy import deepcopy
import builtins
from datetime import date, datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import socket
import urllib.request

import openpyxl
import pytest

from .collect_market_data import (DividendHistory, financial_features, shift_month,
                                  stock_observations)
from .cvm_registry import parse_fca
from .primary_model_candidates import build_primary_candidate


CUTOFF = date(2026, 9, 30)
OBSERVED = datetime(2026, 10, 24, tzinfo=timezone.utc)
SPEC = {'PETR4': (9512, 'PN', 'BRPETRACNPR6', 10), 'VALE3': (4170, 'ON', 'BRVALEACNOR0', 3)}
NOTICE = (b'<title>Vale</title><p>A Vale informa o pagamento em 07 de agosto de 2020 '
          b'dos juros sobre capital proprio deliberados em 19 de dezembro de 2019, '
          b'aproximadamente R$ 1,41 por acao. Datas de corte 26 e 30 de dezembro de 2019.</p>')


class Cache:
    """Same read/receipt interface, with in-memory fixtures or existing disk bytes."""

    def __init__(self, folder, *, disk=False):
        self.folder, self.disk = Path(folder), disk
        self.payloads, self.receipts = {}, {}

    def put(self, relative, value):
        self.payloads[relative] = value if isinstance(value, bytes) else json.dumps(value).encode()

    def read(self, relative, role):
        if relative in self.payloads:
            body, observed = self.payloads[relative], OBSERVED
        elif self.disk:
            path = self.folder / relative
            body = path.read_bytes()
            observed = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc)
        else:
            raise ValueError(role + '_cache_missing')
        self.receipts[relative] = {'file': relative, 'sha256': hashlib.sha256(body).hexdigest(),
            'bytes': len(body), 'retrieved_at': observed.isoformat(), 'url': None,
            'role': role, 'access': 'existing_cache_read_only'}
        return body, observed

    def json(self, relative, role):
        return json.loads(self.read(relative, role)[0])


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def blocked(*args, **kwargs):
        pytest.fail('Primary candidate adapters must not acquire data.')

    monkeypatch.setattr(urllib.request, 'urlopen', blocked)
    monkeypatch.setattr(urllib.request.OpenerDirector, 'open', blocked)
    monkeypatch.setattr(socket, 'create_connection', blocked)


def stamp(day):
    return int(datetime(day.year, day.month, day.day, 16, tzinfo=timezone.utc).timestamp())


def workbook_bytes(rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.cell(2, 13, 'Acao Preferencial Nominativa (PN - PETR4)')
    headers = ['Ano fiscal de referencia', 'Trimestre', 'Situacao', 'Parcela', 'Data do pagamento',
               'Data da divulgacao', 'Data base da posicao acionaria', 'Valor bruto declarado',
               'Valor liquido pago', 'Tipo']
    for column, text in enumerate(headers, 13):
        sheet.cell(3, column, text)
    for number, row in enumerate(rows, 4):
        for column, value in enumerate(row, 13):
            sheet.cell(number, column, value)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def change_workbook(cache, change):
    relative = 'raw/income/petr-history.xlsx'
    workbook = openpyxl.load_workbook(io.BytesIO(cache.payloads[relative]))
    change(workbook.active)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    cache.put(relative, stream.getvalue())


def ri_page(events):
    header = ('<title>Vale</title><table><tr><th>Ano Fiscal de Referencia</th>'
              '<th>Data da divulgacao</th><th>Parcela</th><th>Data do inicio do pagamento</th>'
              '<th>Valor distribuido (R$ por acao)</th><th>US$ por acao</th><th>Tipo</th></tr>')
    rows = []
    for event in events:
        if not event.get('paymentDate'):
            continue
        day = date.fromisoformat(event['paymentDate'][:10])
        amount = f"{event['rate']:.3f}".replace('.', ',')
        values = [str(day.year), '01/08/' + str(day.year), '1a parcela', day.strftime('%d/%m/%Y'),
                  amount, '999,999', 'Dividendos' if event['label'] == 'DIVIDENDO' else event['label']]
        rows.append('<tr>' + ''.join('<td>' + value + '</td>' for value in values) + '</tr>')
    return (header + ''.join(rows) + '</table>').encode()


def fixture(folder, ticker='PETR4'):
    cvm, share, isin, sector = SPEC[ticker]
    row = {'ticker': ticker, 'cvm': cvm, 'share_class': share, 'isin': isin,
           'sector_code': sector, 'identity_status': 'official_cvm_fca', 'model_ready': False}
    cache = Cache(folder)
    days = [shift_month(CUTOFF, -i) for i in reversed(range(141))]
    raw_rows = [{'date': stamp(day), 'rawClose': 20., 'adjustedClose': 7., 'close': 999.} for day in days]
    cache.put(f'raw/brapi-{ticker}.json', {'results': [{'symbol': ticker, 'currency': 'BRL',
        'historicalDataPrice': [dict(item, close=888.) for item in raw_rows],
        'dividendsData': {'stockDividends': [], 'cashDividends': []}}]})
    cache.put(f'raw/income/brapi-raw-prices-{ticker}.json', {'results': [{'symbol': ticker,
        'requestedSymbol': ticker, 'changed': False, 'data': {'historicalDataPrice': raw_rows}}]})
    cache.put(f'raw/expanded/{ticker}-prices.json', {'chart': {'result': [{'meta': {
        'symbol': ticker + '.SA', 'currency': 'BRL'}, 'timestamp': [stamp(day) for day in days],
        'events': {}, 'indicators': {'quote': [{'close': [20.] * len(days)}],
                                   'adjclose': [{'adjclose': [123.] * len(days)}]}}]}})
    events = [{'paymentDate': f'{year}-09-02', 'lastDatePrior': f'{year}-08-10',
               'approvedOn': f'{year}-08-01', 'rate': 1.568705805, 'label': 'JCP',
               'assetIssued': isin, 'isinCode': isin, 'paymentDateIsDeadline': False}
              for year in range(2015, 2027)]
    if ticker == 'PETR4':
        rows = [[year, '1T', 'PAGO', '1a', datetime(year, 9, 2), None, None, 1.568705805, .01, 'JCP']
                for year in range(2015, 2027)]
        cache.put('raw/income/petr-history.xlsx', workbook_bytes(rows))
    else:
        cache.put('raw/income/brapi-v2-VALE3.json', {'results': [{'symbol': ticker,
            'requestedSymbol': ticker, 'changed': False, 'data': {'cashDividends': events}}]})
        cache.put('raw/income/vale-ri.html', ri_page(events))
        cache.put('raw/income/vale-jcp-2020.html', NOTICE)
    return row, cache


def edit_json(cache, relative, change):
    value = json.loads(cache.payloads[relative])
    change(value)
    cache.put(relative, value)


@pytest.mark.parametrize('ticker', SPEC)
def test_candidate_contract_long_history_gross_and_retrospective_availability(tmp_path, ticker):
    row, cache = fixture(tmp_path, ticker)
    before = deepcopy(cache.payloads)
    result = build_primary_candidate(row, cache, CUTOFF)
    assert result['financially_validated'] is False
    assert result['statement_spec'] == (row['cvm'], row['share_class'], row['sector_code'])
    assert result['financial_basis_start'] == '2014-01-01'
    assert result['training_since'] == date(2016, 1, 31)
    assert len(result['prices']) == 141
    assert result['prices'][CUTOFF] == (CUTOFF, 7., 20.)
    history = result['payments']
    assert isinstance(history, DividendHistory) and history.basis == 'payment'
    assert len(history) == 12 and all(day == known and value == 1.568705805 for day, known, value in history)
    assert history.blockers(date(2014, 1, 1), date(2015, 1, 1))
    assert not history.blockers(date(2015, 1, 1), CUTOFF)
    assert len(history.observed_payment_availability) == len(result['ledger']) == len(history)
    for item in [*result['ledger'], *history.observed_payment_availability]:
        assert item['published_at'] is None and item['publication_verified'] is False
        assert item['point_in_time_archive_verified'] is False and item['financially_confirmed'] is False
        assert item['amount_basis'] == 'gross_not_net'
    if ticker == 'PETR4':
        assert all(item['record_date'] is None for item in result['ledger'])
        assert result['independent_cash_matches'] == []
        assert result['checks']['primary_paid_status_count'] == 12
    assert cache.payloads == before
    assert all(item['file'].startswith('raw/') for item in cache.receipts.values())
    assert all(cache.receipts[name]['sha256'] == hashlib.sha256(before[name]).hexdigest() for name in cache.receipts)


@pytest.mark.parametrize('ticker', SPEC)
def test_snapshot_flags_are_irrelevant(tmp_path, ticker):
    row, cache = fixture(tmp_path, ticker)
    first = build_primary_candidate(row, cache, CUTOFF)
    row.update(model_ready=True, model_reference_date='1900-01-01', fundamentals={'status': 'invalid'})
    second = build_primary_candidate(row, cache, CUTOFF)
    assert first['prices'] == second['prices'] and first['ledger'] == second['ledger']


def test_build_has_no_filesystem_writes_including_Excel_temporaries(tmp_path, monkeypatch):
    row, cache = fixture(tmp_path)
    builtin_open, io_open = builtins.open, io.open

    def guard(opener):
        def checked(file, mode='r', *args, **kwargs):
            assert not any(flag in mode for flag in 'wax+'), 'No filesystem writes permitted.'
            return opener(file, mode, *args, **kwargs)
        return checked

    def blocked(*args, **kwargs):
        pytest.fail('No temporary worksheet files permitted.')

    monkeypatch.setattr(builtins, 'open', guard(builtin_open))
    monkeypatch.setattr(io, 'open', guard(io_open))
    monkeypatch.setattr(openpyxl.worksheet._writer, 'create_temporary_file', blocked)
    assert len(build_primary_candidate(row, cache, CUTOFF)['payments']) == 12


def test_parent_universal_financial_functions_consume_the_candidate_contract(tmp_path):
    row, cache = fixture(tmp_path)
    candidate = build_primary_candidate(row, cache, CUTOFF)
    statements = [{'cvm': row['cvm'], 'share_class': row['share_class'],
        'published_at': f'{year + 1}-03-01', 'fiscal_end': f'{year}-12-31', 'version': 1,
        'filing_id': 'fixture-' + str(year), 'profit_per_share': 5., 'free_cash_per_share': 4.,
        'debt_ratio': .3} for year in range(2014, 2026)]
    assert financial_features(statements, candidate['payments'], CUTOFF, candidate['statement_spec'])
    observations, coverage = stock_observations(row['ticker'], candidate['prices'], candidate['payments'],
                                               statements, CUTOFF, candidate['statement_spec'])
    assert observations and min(item['as_of'] for item in observations) >= '2016-01-01'
    assert coverage['current_cash_window']['complete']
    assert candidate['financially_validated'] is False


@pytest.mark.parametrize('field,value', [('cvm', 1), ('share_class', 'ON'), ('isin', 'BRVALEACNOR0'),
    ('identity_status', 'unverified'), ('sector_code', None), ('unit_composition', '1 ON + 1 PN')])
def test_exact_identity_is_required(tmp_path, field, value):
    row, cache = fixture(tmp_path)
    row[field] = value
    with pytest.raises(ValueError):
        build_primary_candidate(row, cache, CUTOFF)


def test_unsupported_format_has_no_cache_access(tmp_path):
    cache = Cache(tmp_path)
    assert build_primary_candidate({'ticker': 'OTHER3'}, cache, CUTOFF) is None
    assert not cache.receipts


@pytest.mark.parametrize('relative,field,value', [
    ('raw/brapi-PETR4.json', 'symbol', 'PETR3'),
    ('raw/brapi-PETR4.json', 'currency', 'USD'),
    ('raw/income/brapi-raw-prices-PETR4.json', 'symbol', 'PETR3'),
    ('raw/income/brapi-raw-prices-PETR4.json', 'requestedSymbol', 'PETR3'),
    ('raw/income/brapi-raw-prices-PETR4.json', 'changed', True),
    ('raw/income/brapi-raw-prices-PETR4.json', 'currency', 'USD')])
def test_provider_exact_symbol_and_currency(tmp_path, relative, field, value):
    row, cache = fixture(tmp_path)
    edit_json(cache, relative, lambda data: data['results'][0].update({field: value}))
    with pytest.raises(ValueError):
        build_primary_candidate(row, cache, CUTOFF)


@pytest.mark.parametrize('label', ['DESDOBRAMENTO', 'GRUPAMENTO', 'BONIFICACAO', 'CONVERSAO'])
@pytest.mark.parametrize('day', ['2022-01-01', '2026-12-01'])
def test_stock_history_actions_are_rejected_even_if_not_in_cash(tmp_path, label, day):
    row, cache = fixture(tmp_path)
    edit_json(cache, 'raw/brapi-PETR4.json', lambda data: data['results'][0]['dividendsData'].update(
        stockDividends=[{'label': label, 'lastDatePrior': day, 'factor': 2}]))
    with pytest.raises(ValueError, match='corporate_action'):
        build_primary_candidate(row, cache, CUTOFF)


@pytest.mark.parametrize('mutation,reason', [
    ('gap', '61_consecutive'), ('raw_missing', 'raw_or_adjusted'), ('duplicate', 'duplicate_raw'),
    ('scale', 'share_scale'), ('few_anchors', 'insufficient_independent'),
    ('yahoo_identity', 'identity_or_currency'), ('split', 'share_action')])
def test_price_evidence_failures(tmp_path, mutation, reason):
    row, cache = fixture(tmp_path)
    relative = 'raw/income/brapi-raw-prices-PETR4.json'
    payload = json.loads(cache.payloads[relative])
    records = payload['results'][0]['data']['historicalDataPrice']
    yahoo_file = 'raw/expanded/PETR4-prices.json'
    yahoo = json.loads(cache.payloads[yahoo_file])
    result = yahoo['chart']['result'][0]
    if mutation == 'gap':
        records.pop(-10)
    elif mutation == 'raw_missing':
        records[-10]['rawClose'] = None
    elif mutation == 'duplicate':
        records.append(dict(records[-1]))
    elif mutation == 'scale':
        records[-10]['rawClose'] *= 2
    elif mutation == 'few_anchors':
        result['timestamp'] = result['timestamp'][-3:]
        result['indicators']['quote'][0]['close'] = [20.] * 3
    elif mutation == 'yahoo_identity':
        result['meta']['symbol'] = 'PETR3.SA'
    elif mutation == 'split':
        result['events']['splits'] = {'x': {'date': stamp(date(2020, 1, 1)), 'numerator': 2, 'denominator': 1}}
    cache.put(relative, payload)
    cache.put(yahoo_file, yahoo)
    with pytest.raises(ValueError, match=reason):
        build_primary_candidate(row, cache, CUTOFF)


def test_old_price_gap_limits_training_suffix(tmp_path):
    row, cache = fixture(tmp_path)
    edit_json(cache, 'raw/income/brapi-raw-prices-PETR4.json', lambda data:
              data['results'][0]['data']['historicalDataPrice'].pop(38))
    result = build_primary_candidate(row, cache, CUTOFF)
    assert min(result['prices']) == date(2018, 4, 30)
    assert result['training_since'] == date(2019, 4, 30)


@pytest.mark.parametrize('status,day,label', [('A PAGAR', None, 'JCP'),
    ('A PAGAR', datetime(2027, 1, 1), 'JCP'), ('PAGO', datetime(2027, 1, 1), 'JCP'),
    ('PAGO', datetime(2026, 9, 2), 'PRINCIPAL')])
def test_petro_future_unknown_future_and_principal_are_not_cash(tmp_path, status, day, label):
    row, cache = fixture(tmp_path)
    def change(sheet):
        for column, value in enumerate([2026, '1T', status, '2a', day, None, None, 999., 999., label], 13):
            sheet.cell(16, column).value = value
    change_workbook(cache, change)
    result = build_primary_candidate(row, cache, CUTOFF)
    assert len(result['payments']) == 12
    assert not any(amount == 999. for _, _, amount in result['payments'])


def test_petro_paid_missing_date_blocks_but_missing_record_and_publication_stay_unknown(tmp_path):
    row, cache = fixture(tmp_path)
    change_workbook(cache, lambda sheet: setattr(sheet.cell(15, 17), 'value', None))
    with pytest.raises(ValueError, match='paid_primary_payment_date_missing'):
        build_primary_candidate(row, cache, CUTOFF)


def test_petro_extraordinary_dividend_uses_gross_not_net(tmp_path):
    row, cache = fixture(tmp_path)
    change_workbook(cache, lambda sheet: setattr(sheet.cell(15, 22), 'value', 'Dividendoss Extraordinarios'))
    result = build_primary_candidate(row, cache, CUTOFF)
    assert result['ledger'][-1]['label'] == 'DIVIDENDO'
    assert result['ledger'][-1]['gross_per_unit'] == 1.568705805


def test_petro_net_column_cannot_replace_gross(tmp_path):
    row, cache = fixture(tmp_path)
    change_workbook(cache, lambda sheet: setattr(sheet.cell(3, 20), 'value', 'Valor liquido pago'))
    with pytest.raises(ValueError):
        build_primary_candidate(row, cache, CUTOFF)


def test_vale_display_precision_does_not_verify_provider_precision(tmp_path):
    row, cache = fixture(tmp_path, 'VALE3')
    result = build_primary_candidate(row, cache, CUTOFF)
    matches = result['checks']['primary_RI_payment_matches']
    assert len(matches) == 12
    assert all(item['RI_displayed_amount'] == '1,569' and item['displayed_decimal_places'] == 3 for item in matches)
    assert all(item['full_provider_precision_verified'] is False for item in matches)
    assert all(item['provider_gross_per_share'] == 1.568705805 for item in matches)


@pytest.mark.parametrize('mutation,reason', [('few', 'insufficient_primary'), ('usd', 'not_gross_BRL'),
    ('deadline', 'deadline'), ('date_missing', 'date_missing'), ('class', 'share_class'),
    ('rate_nan', 'invalid_gross'), ('wrong_precision', 'insufficient_primary')])
def test_vale_rejects_incomplete_or_incompatible_cash(tmp_path, mutation, reason):
    row, cache = fixture(tmp_path, 'VALE3')
    relative = 'raw/income/brapi-v2-VALE3.json'
    data = json.loads(cache.payloads[relative])
    events = data['results'][0]['data']['cashDividends']
    if mutation == 'few':
        cache.put('raw/income/vale-ri.html', ri_page(events[-3:]))
    elif mutation == 'usd':
        cache.put('raw/income/vale-ri.html', cache.payloads['raw/income/vale-ri.html'].replace(b'R$ por acao', b'US$ por acao'))
    elif mutation == 'wrong_precision':
        cache.put('raw/income/vale-ri.html', cache.payloads['raw/income/vale-ri.html'].replace(b'1,569', b'1,568'))
    else:
        field, value = {'deadline': ('paymentDateIsDeadline', True), 'date_missing': ('paymentDate', None),
                        'class': ('isinCode', 'BRVALEACNPA3'), 'rate_nan': ('rate', float('nan'))}[mutation]
        events[-1][field] = value
        cache.put(relative, data)
    with pytest.raises(ValueError, match=reason):
        build_primary_candidate(row, cache, CUTOFF)


@pytest.mark.parametrize('bad_notice', [False, True])
def test_vale_2019_reconciliation_requires_exact_event_and_notice(tmp_path, bad_notice):
    row, cache = fixture(tmp_path, 'VALE3')
    relative = 'raw/income/brapi-v2-VALE3.json'
    def add(data):
        data['results'][0]['data']['cashDividends'].append({'paymentDate': None,
            'lastDatePrior': '2019-12-26', 'approvedOn': '2019-12-19', 'label': 'JCP', 'rate': 1.414364369,
            'isinCode': row['isin'], 'assetIssued': row['isin']})
    edit_json(cache, relative, add)
    if bad_notice:
        cache.put('raw/income/vale-jcp-2020.html', NOTICE.replace(b'26 e 30', b'25 e 30'))
        with pytest.raises(ValueError, match='notice_identity_or_calendar'):
            build_primary_candidate(row, cache, CUTOFF)
    else:
        result = build_primary_candidate(row, cache, CUTOFF)
        payment = next(item for item in result['ledger'] if item['primary_reconciliation'])
        assert payment['payment_date'] == '2020-08-07' and payment['record_date'] == '2019-12-26'
        assert payment['published_at'] is None and payment['gross_per_unit'] == 1.414364369


@pytest.mark.parametrize('kind', ['future_date', 'future_unknown_date', 'principal'])
def test_vale_future_and_principal_are_excluded(tmp_path, kind):
    row, cache = fixture(tmp_path, 'VALE3')
    def add(data):
        event = dict(data['results'][0]['data']['cashDividends'][-1], rate=999.)
        if kind == 'future_date':
            event['paymentDate'] = '2027-01-01'
        elif kind == 'future_unknown_date':
            event.update(paymentDate=None, lastDatePrior='2027-01-01')
        else:
            event['label'] = 'RESTITUICAO DE CAPITAL'
        data['results'][0]['data']['cashDividends'].append(event)
    edit_json(cache, 'raw/income/brapi-v2-VALE3.json', add)
    candidate = build_primary_candidate(row, cache, CUTOFF)
    assert len(candidate['payments']) == 12
    assert all(value != 999. for _, _, value in candidate['payments'])


@pytest.mark.parametrize('ticker', SPEC)
def test_missing_supported_cache_is_explicit_failure(tmp_path, ticker):
    row, cache = fixture(tmp_path, ticker)
    del cache.payloads[f'raw/income/brapi-raw-prices-{ticker}.json']
    with pytest.raises(ValueError, match='cache_missing'):
        build_primary_candidate(row, cache, CUTOFF)


@pytest.mark.parametrize('ticker', SPEC)
def test_real_caches_read_only_exact_FCA_identity_and_rates(ticker, tmp_path):
    folder = Path(os.environ.get('RENDAFUTURA_PRIMARY_CACHE', Path(__file__).parent / 'data/market')).resolve()
    required = [f'raw/brapi-{ticker}.json', f'raw/income/brapi-raw-prices-{ticker}.json',
                f'raw/expanded/{ticker}-prices.json', 'raw/inventory/fca_cia_aberta_2026.zip']
    required += (['raw/income/petr-history.xlsx'] if ticker == 'PETR4' else
                 ['raw/income/brapi-v2-VALE3.json', 'raw/income/vale-ri.html', 'raw/income/vale-jcp-2020.html'])
    if not all((folder / relative).is_file() for relative in required):
        pytest.skip('Existing real source caches unavailable; no acquisition permitted.')
    # Keep the read-only assertion independent of concurrent collector downloads.
    original_folder, folder = folder, tmp_path / 'primary-cache'
    for relative in required:
        target = folder / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(original_folder / relative, target)
        assert target.read_bytes() == (original_folder / relative).read_bytes()
    before = {path.relative_to(folder).as_posix(): (path.stat().st_size, path.stat().st_mtime_ns)
              for path in folder.rglob('*') if path.is_file()}
    digests = {relative: hashlib.sha256((folder / relative).read_bytes()).hexdigest() for relative in required}
    cache = Cache(folder, disk=True)
    fca = parse_fca(cache.read(required[3], 'official_CVM_FCA_identity')[0], 2026, CUTOFF)
    row = fca['instruments'][ticker]
    row['sector_code'] = SPEC[ticker][3]
    row['model_ready'] = False
    candidate = build_primary_candidate(row, cache, CUTOFF)
    assert len(candidate['payments']) == (65 if ticker == 'PETR4' else 33)
    assert candidate['training_since'] == date(2016, 1, 31)
    assert candidate['current_quote_source_file'] == f'raw/expanded/{ticker}-prices.json'
    assert len(candidate['prices']) == 141 and len(candidate['independent_price_matches']) >= 4
    assert candidate['payments'].source_issues or ticker == 'VALE3'
    if ticker == 'VALE3':
        matches = candidate['checks']['primary_RI_payment_matches']
        assert len(matches) >= 4
        assert any(item['RI_displayed_amount'] == '1,569' and item['provider_gross_per_share'] == 1.568705805 for item in matches)
        assert any(item['RI_displayed_amount'] == '0,462' and item['provider_gross_per_share'] == .462016093 for item in matches)
        assert candidate['checks']['primary_payment_date_reconciliations'] == 1
    assert before == {path.relative_to(folder).as_posix(): (path.stat().st_size, path.stat().st_mtime_ns)
                      for path in folder.rglob('*') if path.is_file()}
    assert digests == {relative: hashlib.sha256((folder / relative).read_bytes()).hexdigest() for relative in required}
    assert all('current.json' not in relative and 'snapshot' not in relative for relative in cache.receipts)
    print(json.dumps({'ticker': ticker, 'paid_events': len(candidate['payments']),
        'gross_BRL_per_share_total': sum(value for _, _, value in candidate['payments']),
        'raw_monthly_prices': len(candidate['prices']), 'independent_price_anchors': len(candidate['independent_price_matches']),
        'RI_cash_matches': len(candidate['independent_cash_matches']),
        'files': sorted(cache.receipts)}, sort_keys=True))
