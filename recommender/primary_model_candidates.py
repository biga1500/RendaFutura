"""Read-only adapters for documented primary cash formats, not model approvals.

The caller supplies the same authoritative FCA/registry identity as _one.
Only source formats are selected by ticker; snapshot readiness is never read.
All amounts are gross BRL/share. Availability is retrospective payment-date
observation, never verified publication, settlement, or a point-in-time archive.
The existing income readers' 2015 documentary scope is retained explicitly.
"""

from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import io
import json
import math
import posixpath
import re
from xml.etree import ElementTree as ET
import zipfile

from lxml import html
import openpyxl

from .collect_income_data import (PAYMENT_START, PETROBRAS_URL, VALE_URL,
                                  petrobras_payments, vale_payments,
                                  verified_raw_prices)
from .collect_market_data import DividendHistory, month_end, normalize, shift_month, stock_history
from .payment_sources import (_header, inspect_payment_page,
                              inspect_petrobras_workbook)


# Source identities, not approvals or model/financial exceptions.
_FORMATS = {'PETR4': (9512, 'PN', 'BRPETRACNPR6'),
            'VALE3': (4170, 'ON', 'BRVALEACNOR0')}
_PRICE_RTOL, _PRICE_ATOL = .01, .015


def _require(condition, reason):
    if not condition:
        raise ValueError(reason)


def _day(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if type(value) is date:
        return value
    _require(isinstance(value, str), 'invalid_reported_date')
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        raise ValueError('invalid_reported_date') from None


def _positive(value):
    return (type(value) in (int, float) and math.isfinite(value) and value > 0)


def _read(cache, relative, role, cutoff):
    body, observed = cache.read(relative, role)
    _require(isinstance(observed, datetime) and observed.date() >= cutoff,
             'primary_source_observation_precedes_cutoff')
    return body, observed


def _result(payload, ticker, *, currency=False, requested=False):
    _require(isinstance(payload, dict) and isinstance(payload.get('results'), list)
             and len(payload['results']) == 1, 'ambiguous_provider_result')
    result = payload['results'][0]
    _require(result.get('symbol') == ticker, 'exact_provider_symbol_mismatch')
    if currency:
        _require(result.get('currency') == 'BRL', 'provider_currency_not_BRL')
    if requested:
        _require(result.get('requestedSymbol') == ticker and result.get('changed') is False,
                 'provider_symbol_change_unreconciled')
    if 'currency' in result:
        _require(result['currency'] == 'BRL', 'provider_currency_not_BRL')
    return result


def _prices(row, cache, cutoff):
    ticker = row['ticker']
    quote_file = f'raw/brapi-{ticker}.json'
    quote, _ = _read(cache, quote_file, 'stock_history_and_actions', cutoff)
    result = _result(json.loads(quote), ticker, currency=True)
    # Reuse the established stock-history corporate-action gate before replacing
    # its ambiguous close field with the separately reported rawClose series.
    for action in result.get('dividendsData', {}).get('stockDividends', []):
        day = _day(action.get('lastDatePrior') or action.get('exDate') or action.get('approvedOn'))
        _require(day is not None, 'undated_stock_history_corporate_action')
        _require(day < date(2010, 1, 1), 'stock_history_corporate_action_requires_reconciliation')
    stock_history(quote, cutoff)

    relative = f'raw/income/brapi-raw-prices-{ticker}.json'
    raw, _ = _read(cache, relative, 'reported_raw_daily_prices', cutoff)
    raw_result = _result(json.loads(raw), ticker, requested=True)
    seen = set()
    for item in raw_result['data']['historicalDataPrice']:
        day = datetime.fromtimestamp(item['date'], timezone.utc).date()
        _require(day not in seen, 'duplicate_raw_price_date')
        seen.add(day)
        if day <= cutoff:
            _require(_positive(item.get('rawClose')) and _positive(item.get('adjustedClose')),
                     'reported_raw_or_adjusted_price_missing')
    prices = verified_raw_prices(raw, cutoff)
    required = [shift_month(cutoff, -offset) for offset in range(61)]
    _require(all(month in prices for month in required), 'less_than_61_consecutive_prices')
    # Extend only through a contiguous suffix in the same verified share regime.
    first = cutoff
    while shift_month(first, -1) in prices and shift_month(first, -1) >= month_end(2015, 1):
        first = shift_month(first, -1)
    prices = {month: value for month, value in prices.items() if first <= month <= cutoff}
    _require(all((month - value[0]).days <= 7 for month, value in prices.items()),
             'stale_month_end_price')

    yahoo_file = f'raw/expanded/{ticker}-prices.json'
    yahoo, _ = _read(cache, yahoo_file, 'independent_raw_price_anchors', cutoff)
    payload = json.loads(yahoo)
    _require(not payload['chart'].get('error') and len(payload['chart']['result']) == 1,
             'invalid_independent_price_result')
    independent = payload['chart']['result'][0]
    meta = independent['meta']
    _require(meta.get('symbol') == ticker + '.SA' and meta.get('currency') == 'BRL',
             'independent_price_identity_or_currency_mismatch')
    for action in independent.get('events', {}).get('splits', {}).values():
        day = datetime.fromtimestamp(action['date'], timezone.utc).date()
        _require(day < date(2010, 1, 1), 'independent_share_action_requires_reconciliation')
    stamps, values = independent['timestamp'], independent['indicators']['quote'][0]['close']
    _require(len(stamps) == len(values), 'independent_price_length_mismatch')
    daily, seen = {}, set()
    for stamp, value in zip(stamps, values):
        day = datetime.fromtimestamp(stamp, timezone.utc).date()
        _require(day not in seen, 'duplicate_independent_price_date')
        seen.add(day)
        if day <= cutoff and value is not None:
            _require(_positive(value), 'invalid_independent_raw_price')
            daily[day] = value
    matches = []
    for month, (day, _, raw_close) in sorted(prices.items()):
        if day in daily:
            _require(math.isclose(raw_close, daily[day], rel_tol=_PRICE_RTOL, abs_tol=_PRICE_ATOL),
                     'historical_raw_price_or_share_scale_mismatch')
            matches.append(day.isoformat())
    _require(sum(date.fromisoformat(day) >= required[-1] for day in matches) >= 4,
             'insufficient_independent_raw_price_anchors')
    return prices, matches, relative, yahoo_file


def _legacy_check(reader, expected, *, policy_error):
    """Reuse income parsers without imposing their old event-count/recency policy."""
    try:
        history, _ = reader()
    except ValueError as error:
        if str(error) != policy_error:
            raise
        return 'semantic_validation_legacy_count_or_recency_policy_not_applied'
    actual = Counter((day, amount) for day, _, amount in history)
    _require(actual == Counter(expected), 'income_reader_semantic_cash_mismatch')
    return 'income_reader_and_semantic_inspection_agree'


def _project_workbook(body, title, discarded, labels):
    """Project cached XLSX cells in memory; openpyxl.save creates temp files."""
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    relationships = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        sheets = ET.fromstring(archive.read('xl/workbook.xml')).find(ns + 'sheets')
        sheet = next(item for item in sheets if item.get('name') == title)
        relations = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        relation = next(item for item in relations if item.get('Id') == sheet.get(relationships + 'id'))
        _require(relation.get('TargetMode') != 'External', 'external_primary_worksheet')
        target = relation.get('Target', '')
        name = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
        _require(name.startswith('xl/') and name.endswith('.xml'), 'invalid_primary_worksheet_path')
        root = ET.fromstring(archive.read(name))
        for row in root.find(ns + 'sheetData'):
            number = int(row.get('r'))
            for cell in list(row):
                column = re.fullmatch(r'([A-Z]+)\d+', cell.get('r', ''))
                _require(column is not None, 'invalid_primary_cell_reference')
                column = column[1]
                if number in discarded and column in ('M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V'):
                    row.remove(cell)
                elif number in labels and column == 'V':
                    for child in list(cell):
                        cell.remove(child)
                    cell.set('t', 'inlineStr')
                    ET.SubElement(ET.SubElement(cell, ns + 'is'), ns + 't').text = labels[number]
                elif number in labels and column == 'T':
                    formula = cell.find(ns + 'f')
                    if formula is not None:
                        cell.remove(formula)
        with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as projected:
            for item in archive.infolist():
                projected.writestr(item, ET.tostring(root, encoding='utf-8') if item.filename == name else archive.read(item))
    return stream.getvalue()


def _petrobras(cache, cutoff):
    relative = 'raw/income/petr-history.xlsx'
    body, observed = _read(cache, relative, 'primary_reported_paid_cash', cutoff)
    inspected = inspect_petrobras_workbook(body, 'PETR4', cutoff, observed)
    _require(inspected['status'] != 'invalid' and inspected['report'].get('recognized_tables') == 1,
             'invalid_primary_workbook')
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.worksheets[0]
        title, rows = sheet.title, list(sheet.values)
        expected_headers = {15: 'reported_status', 17: 'payment_date', 18: 'published_at',
                            19: 'record_date', 20: 'gross_per_share', 22: 'label'}
        _require('PETR4' in str(sheet.cell(2, 13).value)
                 and all(_header(sheet.cell(3, column).value) == field
                         for column, field in expected_headers.items()), 'unsupported_primary_workbook_layout')
        future, exclusions, cached_amounts = 0, Counter(), 0
        discarded, labels = set(), {}
        formula_rows = {issue['row'] for issue in inspected['report']['issues']
                        if issue['code'] == 'invalid_amount_or_quantity'}
        for number, values in enumerate(rows[3:], 4):
            fiscal = values[12]
            if type(fiscal) not in (int, float):
                _require(normalize(str(values[14] or ''))
                         not in ('pago', 'pago prescrito', 'a pagar'), 'primary_fiscal_scope_missing')
                continue
            label = normalize(str(values[21] or ''))
            if re.fullmatch(r'dividendos*(?: extraordinarios)?', label) or label == 'div':
                labels[number] = 'DIVIDENDO'
            elif label in ('jcp', 'jscp', 'juros sobre capital proprio'):
                labels[number] = 'JCP'
            elif any(word in label for word in ('principal', 'amortiza', 'restituicao de capital')):
                discarded.add(number)
                exclusions['excluded_principal'] += 1
                continue
            else:
                _require(False, 'unsupported_primary_cash_label')
            status = normalize(str(values[14] or ''))
            day = _day(values[16]) if isinstance(values[16], (date, datetime)) else None
            if day is None and fiscal < PAYMENT_START.year:
                discarded.add(number)
                exclusions['excluded_before_documented_2015_scope'] += 1
                continue
            if status not in ('pago', 'pago prescrito') or day and day > cutoff:
                _require(status in ('a pagar', 'pago', 'pago prescrito'), 'unknown_primary_payment_status')
                future += int(day is None or day > cutoff)
                exclusions['excluded_not_reported_paid_by_cutoff'] += 1
                discarded.add(number)
                continue
            _require(day is not None, 'paid_primary_payment_date_missing')
            if day < PAYMENT_START:
                discarded.add(number)
                continue
            _require(_positive(values[19]), 'primary_gross_amount_missing')
            cached_amounts += int(number in formula_rows)
    finally:
        workbook.close()
    projected = _project_workbook(body, title, discarded, labels)
    parsed = inspect_petrobras_workbook(projected, 'PETR4', cutoff, observed)
    events = [event for event in parsed['events'] if _sheet_scope(event, PAYMENT_START, cutoff)]
    included_rows = {event['source_row'] for event in events}
    _require(not any(issue['blocking'] and issue.get('row') in included_rows
                     and issue['code'] != 'missing_record_date'
                     for issue in parsed['report']['issues']), 'primary_payment_inspection_unresolved')
    expected = [(date.fromisoformat(event['payment_date']), event['gross_per_share']) for event in events]
    policy = _legacy_check(lambda: petrobras_payments(projected, cutoff), expected,
                           policy_error='Historico Petrobras ausente ou desatualizado.')
    source_issues = []
    for issue in inspected['report']['issues']:
        item = dict(issue)
        if item['code'] == 'invalid_amount_or_quantity' and item.get('row') in included_rows:
            item.update(blocking=False, resolution='existing_cached_gross_scalar_no_formula_evaluation')
        elif item['code'] == 'missing_record_date' and item.get('row') in included_rows:
            item.update(blocking=False, resolution='unknown_record_date_preserved_primary_paid_cash_basis')
        source_issues.append(item)
    return events, observed, relative, {'primary_paid_status_count': len(events),
        'future_excluded': future, 'excluded_cash': dict(exclusions), 'income_reader': policy,
        'cached_formula_gross_amounts': cached_amounts,
        'formula_basis': 'existing_cached_numeric_value_no_formula_evaluation'}, source_issues


def _sheet_scope(event, first, cutoff):
    return (event['reported_status'] in ('pago', 'pago prescrito') and event['payment_date'] is not None
            and first.isoformat() <= event['payment_date'] <= cutoff.isoformat())


def _vale_notice(evidence):
    tree = html.fromstring(evidence.decode('utf8'), parser=html.HTMLParser(no_network=True))
    paragraphs = [normalize(' '.join(node.text_content().split())) for node in tree.xpath('//p')]
    _require(any(all(token in text for token in ('a vale informa', '07 de agosto de 2020',
        'juros sobre capital proprio', '19 de dezembro de 2019', '1,41 por acao',
        '26 e 30 de dezembro de 2019')) for text in paragraphs), 'vale_2019_JCP_notice_identity_or_calendar_mismatch')


def _vale_corroboration(content, events, cutoff, observed):
    inspection = inspect_payment_page(content, 'VALE3', cutoff, observed)
    _require(inspection['status'] != 'invalid', 'invalid_primary_RI_page')
    tree = html.fromstring(content.decode('utf8'), parser=html.HTMLParser(no_network=True))
    matches, used = [], set()
    # This RI format names payment *start* explicitly and contains separate
    # BRL/share, USD/share and total columns. Only the first is comparable.
    for table_index, table in enumerate(tree.xpath('//table'), 1):
        rows = [[ ' '.join(' '.join(cell.itertext()).split()) for cell in row.xpath('./th|./td')]
                for row in table.xpath('./tr|./thead/tr|./tbody/tr')]
        if not rows:
            continue
        headers = [normalize(value) for value in rows[0]]
        if (len(headers) < 7 or 'ano fiscal' not in headers[0] or 'divulga' not in headers[1]
                or 'pagamento' not in headers[3] or 'valor distribu' not in headers[4]
                or headers[6] != 'tipo'):
            continue
        _require(not any(word in headers[4] for word in ('us$', 'liquido', 'total')),
                 'primary_RI_not_gross_BRL_per_share')
        if not (('r$' in headers[4] and 'acao' in headers[4])
                or any(len(values) > 4 and re.match(r'R\$ por a[c\u00e7][a\u00e3]o\s+', values[4])
                       for values in rows[1:])):
            continue
        for row_index, values in enumerate(rows[1:], 2):
            if len(values) < 7 or not re.fullmatch(r'\d{4}', values[0]):
                continue
            if int(values[0]) < PAYMENT_START.year - 1:
                continue
            label = normalize(values[6])
            label = 'DIVIDENDO' if label in ('dividendo', 'dividendos') else 'JCP' if label == 'jcp' else None
            if label is None or not re.fullmatch(r'\d{2}/\d{2}/\d{4}', values[3]):
                continue
            day = datetime.strptime(values[3], '%d/%m/%Y').date()
            if not PAYMENT_START <= day <= cutoff:
                continue
            # Ambiguous decimal-point/thousands notation is not promoted to an anchor.
            amount = re.fullmatch(r'(?:R\$ por a[c\u00e7][a\u00e3]o\s+)?(\d+(?:,\d+)?)', values[4])
            if amount is None:
                continue
            displayed = Decimal(amount[1].replace(',', '.'))
            quantum = Decimal(1).scaleb(displayed.as_tuple().exponent)
            for event in events:
                index = event['source_row']
                if (index not in used and event['payment_date'] == day.isoformat() and event['label'] == label
                        and Decimal(str(event['gross_per_share'])).quantize(quantum, rounding=ROUND_HALF_UP) == displayed):
                    used.add(index)
                    matches.append({'source_event': index, 'payment_date': day.isoformat(), 'label': label,
                        'provider_gross_per_share': event['gross_per_share'], 'RI_displayed_amount': amount[1],
                        'displayed_decimal_places': max(0, -displayed.as_tuple().exponent),
                        'table': table_index, 'row': row_index, 'basis': 'rounded_to_RI_display_precision',
                        'full_provider_precision_verified': False})
                    break
    _require(sum(date.fromisoformat(item['payment_date']) >= shift_month(cutoff, -60)
                 for item in matches) >= 4, 'insufficient_primary_RI_payment_matches')
    return matches


def _vale(cache, cutoff):
    relative = 'raw/income/brapi-v2-VALE3.json'
    body, observed = _read(cache, relative, 'reported_historical_cash', cutoff)
    payload = json.loads(body)
    data = _result(payload, 'VALE3', requested=True)
    evidence, _ = _read(cache, 'raw/income/vale-jcp-2020.html', 'primary_payment_date_reconciliation', cutoff)
    events, expected, retained = [], [], []
    future, excluded, reconciled = 0, 0, 0
    for index, item in enumerate(data['data']['cashDividends']):
        if item['label'] not in ('DIVIDENDO', 'JCP'):
            _require(any(word in normalize(item['label']) for word in ('principal', 'amortiza', 'capital')),
                     'unsupported_provider_cash_label')
            excluded += 1
            continue
        day, record = _day(item.get('paymentDate')), _day(item.get('lastDatePrior'))
        if day and day > cutoff or day is None and record and record > cutoff:
            future += 1
            continue
        if day and day < PAYMENT_START or day is None and record and record < PAYMENT_START:
            excluded += 1
            continue
        for field in ('isinCode', 'assetIssued'):
            _require(item.get(field) in ('BRVALEACNOR0', 'VALE3'), 'exact_cash_share_class_mismatch')
        notice_event = (record == date(2019, 12, 26) and item['label'] == 'JCP'
                        and Decimal(str(item['rate'])) == Decimal('1.414364369'))
        special = notice_event and day is None
        if notice_event:
            _require(_day(item.get('approvedOn')) == date(2019, 12, 19)
                     and item['isinCode'] == 'BRVALEACNOR0'
                     and day in (None, date(2020, 8, 7)), 'vale_2019_JCP_event_identity_or_calendar_mismatch')
            _vale_notice(evidence)
            day = date(2020, 8, 7)
            reconciled += int(special)
        _require(day is not None, 'historical_payment_date_missing')
        _require(not item.get('paymentDateIsDeadline'), 'payment_deadline_is_not_payment_date')
        _require(_positive(item.get('rate')), 'invalid_gross_cash_rate')
        _require(record is None or record <= day, 'payment_before_record_date')
        retained.append(item)
        if day > cutoff:
            future += 1
            continue
        expected.append((day, float(item['rate'])))
        events.append({'payment_date': day.isoformat(), 'record_date': record.isoformat() if record else None,
            'label': item['label'], 'gross_per_share': float(item['rate']), 'source_row': index,
            'reported_status': 'reported_historical_payment', 'primary_reconciliation': special})
    data['data']['cashDividends'] = retained
    policy = _legacy_check(lambda: vale_payments(json.dumps(payload).encode(), evidence, cutoff), expected,
                           policy_error='Historico Vale ausente ou desatualizado.')
    primary, primary_observed = _read(cache, 'raw/income/vale-ri.html', 'primary_display_precision_cash_corroboration', cutoff)
    matches = _vale_corroboration(primary, events, cutoff, primary_observed)
    return events, observed, relative, {'primary_RI_payment_matches': matches,
        'primary_RI_precision_limit': 'displayed_rounding_only_not_full_provider_decimal_precision',
        'primary_RI_corroboration_scope': 'partial_history_not_exhaustive_or_full_precision_verification',
        'primary_RI_corroboration_ratio': len(matches) / len(events) if events else 0.,
        'primary_payment_date_reconciliations': reconciled, 'future_excluded': future,
        'excluded_principal_or_before_scope': excluded, 'income_reader': policy}, []


def build_primary_candidate(row, cache, cutoff):
    """Return an _one candidate, None for unsupported formats, ValueError(reason) otherwise.

    cache.folder/read(relative, role), json and receipts belong to the caller.
    No acquisitions, pointer reads/writes, snapshot approvals, or financial gates.
    The parent must apply universal DFP financial_features/stock_observations.
    """
    if row.get('ticker') not in _FORMATS:
        return None
    # The parent can import this adapter before defining its identity helper.
    from .inventory_candidates import _identity

    cutoff = _day(cutoff)
    _require(type(cutoff) is date and cutoff == month_end(cutoff.year, cutoff.month)
             and cutoff <= date.today(), 'cutoff_must_be_past_or_current_month_end')
    ticker = row['ticker']
    try:
        detail, codes = _identity(row, cache, cutoff)
        cvm, share, isin = _FORMATS[ticker]
        _require((row['cvm'], row['share_class']) == (cvm, share), 'primary_source_exact_issuer_class_mismatch')
        _require(row.get('isin') in (None, isin) and codes.get(ticker, isin) == isin,
                 'primary_source_exact_ISIN_mismatch')
        prices, price_matches, price_file, yahoo_file = _prices(row, cache, cutoff)
        events, observed, cash_file, cash_checks, issues = (_petrobras(cache, cutoff) if ticker == 'PETR4'
                                                          else _vale(cache, cutoff))
        _require(bool(events), 'reported_historical_cash_unavailable')
        _require(min(date.fromisoformat(event['payment_date']) for event in events) <= shift_month(cutoff, -60),
                 'reported_cash_does_not_span_return_window')
        history = DividendHistory(cutoff, len(events) + cash_checks['future_excluded'], basis='payment')
        ledger, seen = [], set()
        for event in sorted(events, key=lambda item: (item['payment_date'], item['source_row'])):
            day, amount = date.fromisoformat(event['payment_date']), event['gross_per_share']
            key = day, amount, event['label'], event.get('record_date')
            _require(key not in seen, 'duplicate_primary_cash_event')
            seen.add(key)
            _require(_positive(amount), 'invalid_primary_gross_cash')
            entry = {'ticker': ticker, 'record_date': event.get('record_date'), 'payment_date': day.isoformat(),
                'gross_per_unit': amount, 'gross_per_share': amount, 'label': event['label'],
                'source_row': event['source_row'], 'source_file': cash_file, 'status': 'included',
                'reported_payment_status': event['reported_status'], 'published_at': None,
                'observed_at': observed.isoformat(), 'publication_date_reported': False,
                'publication_verified': False, 'approval_publication_verified': False,
                'point_in_time_archive_verified': False, 'financially_confirmed': False,
                'financial_confirmation': 'reported_historical_cash_not_settlement_confirmed',
                'availability_basis': 'payment_date_retrospective_observation_not_publication',
                'share_basis_date': cutoff.isoformat(), 'currency': 'BRL', 'amount_basis': 'gross_not_net',
                'primary_reconciliation': event.get('primary_reconciliation', False),
                'primary_evidence': ([{'file': 'raw/income/vale-jcp-2020.html',
                    'sha256': cache.receipts['raw/income/vale-jcp-2020.html']['sha256'],
                    'basis': 'exact_event_identity_payment_date_only_not_full_amount_precision'}]
                    if event.get('primary_reconciliation') else [])}
            ledger.append(entry)
            history.append((day, day, amount))
            history.observed_payment_availability.append({**entry, 'event_index': event['source_row'],
                'availability_date': day.isoformat(), 'available_at': day.isoformat(), 'approval_date': None})
        history.coverage_start = PAYMENT_START
        history.source_observed_at = observed.isoformat()
        history.future_events = cash_checks['future_excluded']
        history.source_issues, history.installments, history.reconciled_events = issues, ledger, ledger
        history.issues.append({'event_index': -1, 'interval_start': date.min.isoformat(),
            'interval_end': PAYMENT_START.isoformat(), 'reason': 'outside_documented_primary_payment_coverage'})
        history.cash_payment_issues = history.issues
        history.filter_counts = {'accepted_events': len(history), 'future_payment_date': history.future_events}
        matches = cash_checks.get('primary_RI_payment_matches', [])
        checks = {'identity': 'authoritative_FCA_or_registry_exact_CVM_class_and_provider_symbol',
            'prices': '61_consecutive_monthly_raw_anchors', 'independent_price_matches': len(price_matches),
            'raw_price_reconciliation': {'relative_tolerance': _PRICE_RTOL, 'absolute_tolerance_BRL': _PRICE_ATOL,
                'basis': 'all_overlapping_reported_month_end_dates_not_adjusted_close',
                'source_file': yahoo_file}, 'independent_cash_matches': len(matches),
            'cash_history_qualification': 'primary_reported_paid' if ticker == 'PETR4' else 'historical_reported_RI_display_precision_corroborated',
            'point_in_time_archive_verified': False, **cash_checks}
        return {'ticker': ticker, 'name': row.get('name') or (detail or {}).get('companyName') or ticker,
            'cvm': cvm, 'share_class': share, 'isin': isin, 'sector_code': row['sector_code'],
            'industry_classification': (detail or {}).get('industryClassification'), 'prices': prices,
            'payments': history, 'ledger': ledger,
            'independent_cash_matches': [item['payment_date'] for item in matches],
            'independent_price_matches': price_matches, 'share_actions': [],
            'share_normalization': {'share_basis_date': cutoff.isoformat(), 'price_source_basis': 'reported_rawClose',
                'price_output_basis': 'cutoff_share_units_no_dividend_reinvestment', 'payments': [], 'actions': [],
                'calendar_basis': 'reported_payment_dates_not_ex_dates'},
            'payment_coverage_start': PAYMENT_START.isoformat(), 'financial_basis_start': '2014-01-01',
            'training_since': shift_month(min(prices), 12), 'statement_spec': (cvm, share, row['sector_code']),
            'payment_url': PETROBRAS_URL if ticker == 'PETR4' else VALE_URL,
            'price_url': cache.receipts[price_file].get('url'), 'price_source_file': price_file,
            'current_quote_source_file': yahoo_file, 'checks': checks,
            'readiness_scope': 'reconciled_sources_only_financial_checks_required', 'financially_validated': False}
    except (KeyError, TypeError, IndexError, OverflowError, OSError) as error:
        raise ValueError('incomplete_primary_cached_source') from error
