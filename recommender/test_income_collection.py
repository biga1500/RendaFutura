from datetime import date, datetime
import hashlib
import io
import json

import openpyxl
import pytest

from .collect_income_data import petrobras_payments, select_bonds, treasury_payments, vale_payments
from .collect_market_data import month_end, shift_month
from .datasets import real_dataset
from .test_income import paid_cash_files


CUTOFF = date(2026, 9, 30)


def petro_sheet(*extra):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.cell(2, 13, 'Acao Preferencial PETR4')
    sheet.cell(3, 13, 'Ano fiscal')
    for index in range(40):
        day = shift_month(CUTOFF, -index)
        row = [None] * 23
        row[12:22] = [day.year, '1T', 'PAGO', str(index), datetime.combine(day, datetime.min.time()),
                       datetime(2022, 1, 1), datetime(2022, 1, 2), .1, .085, 'JCP']
        sheet.append(row)
    for change in extra:
        row = [None] * 23
        row[12:22] = [2026, '2T', change.get('status', 'A PAGAR'), 'future', change.get('day', datetime(2026, 12, 21)),
                       datetime(2026, 8, 6), datetime(2026, 8, 21), .67407131, .55610883, 'JCP']
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_primary_paid_status_not_record_date_or_future_promise():
    payments, audit = petrobras_payments(petro_sheet({}), CUTOFF)
    assert len(payments) == 40
    assert all(amount == .1 for _, _, amount in payments)
    assert payments[0][0] == CUTOFF
    assert audit[-1]['status'] == 'excluded_not_paid_by_reference'


def test_primary_current_payment_without_date_is_blocked():
    with pytest.raises(ValueError, match='data de pagamento'):
        petrobras_payments(petro_sheet({'day': None}), CUTOFF)


def vale_payload(*extra):
    events = []
    for index in range(20):
        day = shift_month(CUTOFF, -index).isoformat()
        events.append({'label': 'DIVIDENDO', 'rate': .1, 'paymentDate': day,
                       'lastDatePrior': day, 'paymentDateIsDeadline': False})
    return json.dumps({'results': [{'symbol': 'VALE3', 'data': {'cashDividends': events + list(extra)}}]}).encode()


def test_missing_vale_payment_requires_exact_primary_evidence():
    event = {'label': 'JCP', 'rate': 1.414364369, 'paymentDate': None, 'lastDatePrior': '2019-12-26'}
    with pytest.raises(ValueError, match='Evidencia'):
        vale_payments(vale_payload(event), b'<p>No payment information</p>', CUTOFF)
    payments, audit = vale_payments(vale_payload(event), b'<p>pagamento em 07 de agosto de 2020</p>', CUTOFF)
    assert (date(2020, 8, 7), date(2020, 8, 7), 1.414364369) in payments
    assert audit[-1]['primary_reconciliation'] is True


@pytest.mark.parametrize('change', [{'paymentDate': None}, {'paymentDateIsDeadline': True}])
def test_unknown_or_deadline_does_not_become_observed_cash(change):
    event = {'label': 'DIVIDENDO', 'rate': .1, 'paymentDate': '2026-09-02',
             'lastDatePrior': '2026-08-11', **change}
    with pytest.raises(ValueError):
        vale_payments(vale_payload(event), b'', CUTOFF)


def test_coupon_resource_contains_only_observed_coupon_unit_values():
    content = ('Tipo Titulo;Vencimento do Titulo;Data Resgate;PU;Quantidade;Valor\n'
               'Tesouro Prefixado com Juros Semestrais;01/01/2031;01/07/2026;48,808850;2,00;97,62\n'
               'Tesouro Prefixado com Juros Semestrais;01/01/2031;04/01/2027;48,808850;2,00;97,62\n').encode()
    payments = treasury_payments(content, CUTOFF)
    assert payments[('Tesouro Prefixado com Juros Semestrais', date(2031, 1, 1))] == {date(2026, 7, 1): 48.80885}


def test_missing_expected_coupon_excludes_bond_instead_of_zero_filling():
    key = ('Tesouro Prefixado com Juros Semestrais', date(2031, 1, 1))
    histories = {key: {shift_month(CUTOFF, -i): (shift_month(CUTOFF, -i), 1000) for i in range(61)}}
    payments = {key: {shift_month(CUTOFF, -i): 48.80885 for i in range(60) if shift_month(CUTOFF, -i).month in (1, 7)}}
    selected, _ = select_bonds(histories, payments, CUTOFF)
    assert selected == [key]
    payments[key].pop(date(2026, 7, 31))
    with pytest.raises(ValueError, match='cupons efetivos completos'):
        select_bonds(histories, payments, CUTOFF)


def test_cash_file_hash_detects_changes(paid_cash_files):
    folder, metadata = paid_cash_files
    metadata['cashflows_sha256'] = hashlib.sha256((folder / 'cashflows.csv').read_bytes()).hexdigest()
    (folder / 'assets.json').write_text(json.dumps(metadata), encoding='utf8')
    real_dataset(folder)
    with (folder / 'cashflows.csv').open('a') as stream:
        stream.write('\n')
    with pytest.raises(ValueError, match='Integridade'):
        real_dataset(folder)
