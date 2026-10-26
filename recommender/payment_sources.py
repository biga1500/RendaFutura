"""Offline payment-source inspection with an opt-in bounded primary collector.

The inspect_* functions never fetch or write files. collect_primary_payments
alone acquires allowlisted public files into an explicit immutable raw cache.

Contract: inspect_payment_page(content, ticker, as_of, observed_at) returns a JSON
compatible dict with status, events, field_status, report and source_format.
Status is available, partial, invalid, not_found or zero_window_confirmed. It
describes source readability, NOT eligibility for the cash model or completeness
of a five-year history. There are no minimum event-count/history-length gates.

Events use ISO dates and gross BRL per share on the REPORTED share basis; no
split adjustment, reinvestment, net-to-gross conversion or date inference occurs.
record_date is data-com, never payment_date. Missing dates/amounts remain None.
announced/past_scheduled describe schedules; passage of time is not proof of
payment. paid_reported requires an explicit paid status in the source, not just
a past date. Even paid_reported is NOT independent financial confirmation:
cash_model_eligible and financially_confirmed are always False. Publication
after as_of is retained with available_as_of=False; missing publication yields
None (not historical availability inferred from a later retrieval).

inspect_petrobras_workbook is a viable PRIMARY issuer fallback for the exact
PETR3/PETR4 blocks in the linked native XLSX. inspect_cvm_proventos_index is a
PRIMARY document-discovery fallback, not a payment extractor: IPE reference and
submission dates are never interpreted as cash dates. Parent owns acquisition,
identity reconciliation, share basis, verification and cash-model integration.
SOURCE_DISCOVERIES records bounded public checks; no B3 automation is needed.
"""

import calendar
from collections import Counter
import csv
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
from itertools import islice
import math
from pathlib import Path
import re
import unicodedata
from urllib.parse import parse_qsl, urlsplit
import zipfile

from lxml import etree, html


SOURCE_DISCOVERIES = (
    {'source': 'CVM IPE', 'primary': True,
     'url': 'https://dados.cvm.gov.br/dados/CIA_ABERTA/DOC/IPE/DADOS/ipe_cia_aberta_2026.zip',
     'documentation_url': 'https://dados.cvm.gov.br/dataset/cia_aberta-doc-ipe',
     'checked_on': '2026-10-24', 'http_status': 200, 'bytes': 1749658,
     'sha256': 'd2cd1be0b0db8be3cf7a51985649dce142c34ed96b4bc38296daa9721354780c',
     'proventos_rows': 546, 'category': 'Relat\u00f3rio Proventos',
     'limitation': 'Index only: no ISIN, gross amount or paid confirmation columns.'},
    {'source': 'CVM Formulario de Proventos', 'primary': True,
     'url': 'https://www.gov.br/cvm/pt-br/assuntos/regulados/consultas-por-participante/companhias/envio-de-informacoes-enet/06-formulario-de-proventos',
     'limitation': 'ISIN, gross/unit, installments and protocol are declaration fields, not financial settlement evidence.'},
    {'source': 'CVM public document sample', 'primary': True,
     'url': 'https://www.rad.cvm.gov.br/ENET/frmDownloadDocumento.aspx?Tela=ext&descTipo=IPE&CodigoInstituicao=1&numProtocolo=1477025&numSequencia=1001731&numVersao=1',
     'checked_on': '2026-10-24', 'http_status': 200,
     'limitation': 'text/html; incomplete transfer on both attempts; no document validated. No bypass or authentication attempted.'},
    {'source': 'Petrobras RI native XLSX', 'primary': True,
     'landing_url': 'https://www.investidorpetrobras.com.br/acoes-dividendos-e-dividas/dividendos-e-jcp/',
     'url': 'https://filemanager-cdn.mziq.com/published/25fdf098-34f5-4608-b7fa-17d60b2de47d/c5272574-8ecd-4862-af72-f5fdb9f9008d',
     'checked_on': '2026-10-24', 'http_status': 200, 'bytes': 33728,
     'sha256': 'b8c8c8f1ae240e8905f21d1e823dceaebfe27704e79420dc236a03d21b0abd27',
     'tickers': ['PETR3', 'PETR4'],
     'limitation': 'Reported gross declaration and PAGO/A PAGAR; not independent settlement evidence or a general issuer fallback.'},
)

_FIELDS = ('record_date', 'payment_date', 'gross_per_share', 'published_at', 'reference_date')
_MAX_BYTES = 20_000_000
_TICKER_PATTERN = r'[A-Z][A-Z0-9]{3}\d{1,2}'


def _norm(value):
    value = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().lower()
    return ' '.join(re.sub(r'[^a-z0-9/]+', ' ', value).split())


def _day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or '').strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError('invalid_date')


def _context(ticker, as_of, observed_at):
    if not isinstance(ticker, str) or not re.fullmatch(_TICKER_PATTERN, ticker):
        raise ValueError('An exact uppercase ticker is required.')
    cutoff = _day(as_of)
    if isinstance(observed_at, str):
        observed_at = datetime.fromisoformat(observed_at.replace('Z', '+00:00'))
    if not isinstance(observed_at, datetime) or observed_at.utcoffset() is None:
        raise ValueError('observed_at must be a timezone-aware timestamp.')
    if observed_at > datetime.now(timezone.utc) + timedelta(minutes=5):
        raise ValueError('observed_at is in the future.')
    if cutoff > observed_at.date():
        raise ValueError('as_of cannot be later than observed_at.')
    return cutoff, observed_at


def _result(content, source_format, cutoff, observed):
    body = content.encode('utf-8') if isinstance(content, str) else content
    if not isinstance(body, bytes) or len(body) > _MAX_BYTES:
        raise ValueError('Expected bounded HTML, CSV, ZIP or XLSX content.')
    previous = cutoff.replace(year=cutoff.year - 1,
                              day=min(cutoff.day, calendar.monthrange(cutoff.year - 1, cutoff.month)[1]))
    return {'status': 'not_found', 'events': [], 'field_status': {key: 'not_found' for key in _FIELDS},
            'source_format': source_format,
            'report': {'issues': [], 'counts': {}, 'zero_claims': [],
                       'requested_window': {'start': (previous + timedelta(days=1)).isoformat(),
                                            'end': cutoff.isoformat(), 'scope': ['DIVIDENDO', 'JCP']},
                       'as_of': cutoff.isoformat(), 'observed_at': observed.isoformat(),
                       'sha256': hashlib.sha256(body).hexdigest(), 'bytes': len(body),
                       'financially_confirmed': False, 'cash_model_eligible': False,
                       'history_complete': False, 'share_basis': 'reported_unverified'}}


def _issue(result, code, *, row=None, field=None, blocking=True):
    item = {'code': code, 'blocking': blocking}
    if row is not None:
        item['row'] = row
    if field is not None:
        item['field'] = field
    # Do not copy page text, arbitrary URLs or sensitive query strings into logs.
    if item not in result['report']['issues']:
        result['report']['issues'].append(item)


def _br_number(value):
    if isinstance(value, bool):
        raise ValueError('invalid_number')
    if isinstance(value, (int, float, Decimal)):
        result = Decimal(str(value))
    else:
        text = str(value or '').strip().replace('\xa0', ' ')
        text = re.sub(r'^R\$\s*', '', text).strip()
        # HTML strings follow explicit Brazilian notation, not guessed US decimals.
        if not re.fullmatch(r'[+-]?(?:\d+|[1-9]\d{0,2}(?:\.\d{3})+)(?:,\d+)?', text):
            raise ValueError('invalid_brazilian_number')
        result = Decimal(text.replace('.', '').replace(',', '.'))
    if not result.is_finite() or not math.isfinite(float(result)):
        raise ValueError('non_finite_number')
    return result


def _label(value):
    text = _norm(value)
    if text in ('div', 'dividendo', 'dividendos', 'dividend'):
        return 'DIVIDENDO'
    if text in ('jcp', 'jscp', 'juros sobre capital proprio', 'juros sobre o capital proprio'):
        return 'JCP'
    return None


def _header(value):
    text = _norm(value)
    if text in ('ticker', 'codigo de negociacao', 'ativo'):
        return 'ticker'
    if text in ('isin', 'codigo isin'):
        return 'isin'
    if text in ('tipo', 'tipo de provento', 'provento', 'natureza', 'tipo provento'):
        return 'label'
    if text in ('situacao', 'status', 'situacao de pagamento', 'status do pagamento'):
        return 'reported_status'
    if text in ('data ex', 'data ex dividendos', 'ex dividendos', 'ex date'):
        return 'ex_date'
    if text in ('data com', 'data base', 'data record', 'record date', 'data de registro', 'data registro'):
        return 'record_date'
    if ('data' in text and ('posi' in text and 'acion' in text or 'ultimo dia' in text and 'direito' in text)):
        return 'record_date'
    if text in ('pagamento', 'data pagamento', 'data de pagamento', 'data do pagamento', 'payment date'):
        return 'payment_date'
    if 'data' in text and any(part in text for part in ('divulga', 'publica')):
        return 'published_at'
    if text in ('data aprovacao', 'data de aprovacao', 'data deliberacao', 'data da deliberacao'):
        return 'approval_date'
    if text in ('data referencia', 'data de referencia'):
        return 'reference_date'
    if text in ('qtd cotacao', 'quantidade base', 'quantidade de acoes por unidade', 'qtd acoes por unidade'):
        return 'quoted_quantity'
    if text in ('qtd', 'quantidade', 'quantidade de acoes', 'qtd acoes'):
        return 'ambiguous_quantity'
    if 'liquido' in text or 'montante' in text or 'total' in text:
        return None
    if text == 'valor' or 'valor' in text and any(part in text for part in ('bruto', 'acao', 'acoes', 'unitario', 'unidade')):
        return 'gross_per_share'
    return None


def _columns(values):
    found = {}
    for index, value in enumerate(values):
        key = _header(value)
        if key:
            if key in found:
                return None
            found[key] = index
    # An ex-date alone is recognizable but must not become a payment date.
    return found if 'gross_per_share' in found and any(key in found for key in ('record_date', 'payment_date', 'ex_date')) else None


def _date_field(result, value, key, row):
    if value in (None, '') or str(value).strip() in ('-', '--'):
        if key in ('record_date', 'payment_date'):
            _issue(result, 'missing_' + key, row=row, field=key)
        return None
    try:
        return _day(value).isoformat()
    except (ValueError, TypeError):
        _issue(result, 'invalid_' + key, row=row, field=key)
        return None


def _event(result, values, columns, ticker, cutoff, row, *, amount_header=''):
    def get(key):
        index = columns.get(key)
        return values[index] if index is not None and index < len(values) else None

    row_ticker = str(get('ticker') or '').strip()
    if 'ticker' in columns and row_ticker != ticker:
        _issue(result, 'other_or_missing_ticker_row', row=row, blocking=False)
        return
    label = _label(get('label'))
    if label is None:
        _issue(result, 'unsupported_or_missing_income_label', row=row)
        return
    event = {'ticker': ticker, 'label': label,
             'cash_model_eligible': False, 'financially_confirmed': False,
             'share_basis': 'reported_unverified', 'gross_per_share': None,
             'quoted_quantity': None, 'reported_amount': None,
             'source_row': row, 'isin': str(get('isin') or '').strip() or None}
    for key in (*_FIELDS[:2], 'published_at', 'reference_date', 'approval_date'):
        event[key] = _date_field(result, get(key), key, row)
    if event['record_date'] and event['payment_date'] and event['payment_date'] < event['record_date']:
        _issue(result, 'payment_before_record_date', row=row)
        event['date_order_valid'] = False
    else:
        event['date_order_valid'] = True
    try:
        amount_text = get('gross_per_share')
        quantity = get('quoted_quantity') if 'quoted_quantity' in columns else 1
        if 'ambiguous_quantity' in columns:
            raise ValueError('ambiguous_quantity_basis')
        # Explicit per-N-share headers are not total declared cash amounts.
        unit_header = unicodedata.normalize('NFKD', str(amount_header)).encode('ascii', 'ignore').decode().lower()
        match = re.search(r'por\s+(\d[\d.]*)\s+(?:acao|acoes)\b', unit_header)
        if match:
            if 'quoted_quantity' in columns:
                raise ValueError('ambiguous_quantity_basis')
            quantity = match.group(1)
        divisor = _br_number(quantity)
        if divisor <= 0 or divisor != divisor.to_integral_value():
            raise ValueError('invalid_quote_quantity')
        amount = _br_number(amount_text)
        if amount < 0:
            raise ValueError('negative_income')
        event.update(gross_per_share=float(amount / divisor), reported_amount=float(amount),
                     quoted_quantity=int(divisor))
    except (ValueError, InvalidOperation, OverflowError):
        _issue(result, 'invalid_amount_or_quantity', row=row, field='gross_per_share')
    status = _norm(get('reported_status'))
    event['reported_status'] = status or None
    event['payment_status_observed_at'] = result['report']['observed_at']
    event['paid_status_as_of_verified'] = False
    payment = event['payment_date']
    paid = status in ('pago', 'paga', 'pagos', 'pagas', 'paid', 'liquidado', 'liquidada', 'pago prescrito')
    event['reported_expired'] = status == 'pago prescrito'
    event['status'] = ('paid_reported' if paid and payment and payment <= cutoff.isoformat()
                       else 'announced' if payment and payment > cutoff.isoformat()
                       else 'past_scheduled' if payment else 'unresolved')
    if paid and payment and payment > cutoff.isoformat():
        _issue(result, 'paid_status_after_as_of', row=row, blocking=False)
    if status in ('cancelado', 'cancelada', 'cancelled'):
        event['status'] = 'cancelled'
    event['available_as_of'] = event['published_at'] <= cutoff.isoformat() if event['published_at'] else None
    if event['available_as_of'] is False:
        _issue(result, 'publication_after_as_of', row=row, blocking=False)
    if event['published_at'] and event['published_at'] > result['report']['observed_at'][:10]:
        _issue(result, 'publication_after_observation', row=row, field='published_at')
    if event['published_at'] is None:
        _issue(result, 'publication_not_reported', row=row, blocking=False)
    result['events'].append(event)


def _finish(result):
    unique, seen = [], set()
    amounts = {}
    for event in result['events']:
        identity = tuple(event.get(key) for key in ('ticker', 'label', 'record_date', 'payment_date', 'published_at', 'isin'))
        signature = identity + (event['gross_per_share'], event['reported_status'])
        if signature in seen:
            _issue(result, 'duplicate_event', row=event['source_row'], blocking=False)
            continue
        seen.add(signature)
        if identity in amounts and amounts[identity] != event['gross_per_share']:
            _issue(result, 'conflicting_event_amount', row=event['source_row'])
        amounts[identity] = event['gross_per_share']
        unique.append(event)
    result['events'] = unique
    for key in _FIELDS:
        present = sum(event.get(key) is not None for event in unique)
        invalid = any(issue.get('field') == key and issue['code'].startswith('invalid_')
                      for issue in result['report']['issues'])
        result['field_status'][key] = ('partial' if 0 < present < len(unique) else 'available' if present
                                       else 'invalid' if invalid else 'not_found')
    result['report']['counts'] = dict(Counter(event['status'] for event in unique))
    result['report']['counts']['events'] = len(unique)
    blocking = any(issue['blocking'] for issue in result['report']['issues'])
    if unique:
        result['status'] = 'partial' if blocking else 'available'
    elif blocking:
        result['status'] = 'invalid' if result['report'].get('recognized_tables') else 'partial'
    return result


def _table_rows(table):
    """Expand table spans so multirow headers keep their column alignment."""
    active = {}
    for row in table.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr'):
        values, column = [], 0
        for cell in row.xpath('./td|./th'):
            while column in active:
                values.append(active[column][0])
                column += 1
            try:
                colspan, rowspan = int(cell.get('colspan', 1)), int(cell.get('rowspan', 1))
            except ValueError:
                raise ValueError('invalid_table_span') from None
            if not 1 <= colspan <= 100 or not 1 <= rowspan <= 100:
                raise ValueError('invalid_table_span')
            text = ' '.join(' '.join(cell.itertext()).split())
            for _ in range(colspan):
                values.append(text)
                if rowspan > 1:
                    active[column] = (text, rowspan)
                column += 1
        while column in active:
            values.append(active[column][0])
            column += 1
        active = {key: (value, left - 1) for key, (value, left) in active.items() if left > 1}
        yield values


def _find_header(rows):
    selected, previous = None, None
    for index, values in enumerate(rows[:8]):
        if selected and 'label' in selected[0] and selected[0]['label'] < len(values) and _label(values[selected[0]['label']]):
            break
        columns = _columns(values)
        headers = values
        if not columns and previous and len(previous) == len(values):
            headers = [child if _header(child) else f'{parent} {child}'
                       for parent, child in zip(previous, values)]
            columns = _columns(headers)
        if columns:
            selected = columns, index, headers
        previous = values
    return selected or (None, None, None)


def _zero_claims(tree, ticker, cutoff, observed, result):
    # Small publisher paragraphs only, never user comments, scripts or testimonials.
    nodes = tree.xpath('//p|//h2|//h3|//div[not(div) and not(p)]')
    for node in nodes:
        ancestors = [node, *node.iterancestors()]
        if any(re.search(r'comment|coment|review|testimonial|forum|user', (item.get('id', '') + ' ' + item.get('class', '')).lower())
               or item.tag in ('script', 'noscript', 'footer') for item in ancestors):
            continue
        text = _norm(' '.join(node.itertext()))
        if len(text) > 700:
            continue
        if not re.search(r'nao (?:pagou|distribuiu)\b|sem pagamentos? de|nenhum pagamento de', text):
            continue
        if re.search(r'\b(?:formula|caso|empresas que|se a empresa)\b', text):
            continue
        tickers = set(re.findall(r'\b' + _TICKER_PATTERN.lower() + r'\b', text))
        if tickers and tickers != {ticker.lower()}:
            continue
        broad = 'proventos' in text or 'dividendos e jcp' in text or 'dividendos e juros sobre' in text
        if not broad and 'dividendo' not in text:
            continue
        scope = ['DIVIDENDO', 'JCP'] if broad else ['DIVIDENDO']
        dates = re.findall(r'\b\d{2}/\d{2}/\d{4}\b', text)
        start = end = None
        anchored = False
        if len(dates) == 2 and 'entre ' in text:
            try:
                start, end = (_day(value) for value in dates)
                anchored = True
            except ValueError:
                _issue(result, 'invalid_zero_claim_dates')
        elif '12 meses' in text or 'ultimo ano' in text:
            end = observed.date()
            reference = tree.xpath('//meta[@name="date" or @name="reference_date"]/@content')
            if len(reference) == 1:
                try:
                    end = _day(reference[0])
                    anchored = True
                except ValueError:
                    _issue(result, 'invalid_page_reference_date', blocking=False)
            start = end.replace(year=end.year - 1, day=min(end.day, calendar.monthrange(end.year - 1, end.month)[1])) + timedelta(days=1)
        claim = {'scope': scope, 'start': start.isoformat() if start else None,
                 'end': end.isoformat() if end else None, 'explicit_date_anchor': anchored,
                 'ticker_identity_verified': tickers == {ticker.lower()} or result['report'].get('page_ticker_verified', False),
                 'source_assertion_only': True}
        if claim not in result['report']['zero_claims']:
            result['report']['zero_claims'].append(claim)
    if result['report']['zero_claims']:
        requested = result['report']['requested_window']
        complete = any(claim['explicit_date_anchor'] and claim['ticker_identity_verified'] and claim['scope'] == requested['scope']
                       and claim['start'] and claim['start'] <= requested['start']
                       and claim['end'] >= requested['end'] and claim['end'] <= observed.date().isoformat()
                       for claim in result['report']['zero_claims'])
        if complete and not result['events'] and not any(issue['blocking'] for issue in result['report']['issues']):
            result['status'] = 'zero_window_confirmed'
        elif not result['events']:
            result['status'] = 'partial'
            _issue(result, 'zero_claim_does_not_confirm_requested_scope_and_window', blocking=False)
        else:
            for event in result['events']:
                conflicts = any(claim['start'] and claim['end'] and event['payment_date']
                                and claim['start'] <= event['payment_date'] <= claim['end']
                                and event['label'] in claim['scope'] for claim in result['report']['zero_claims'])
                if conflicts and event['gross_per_share']:
                    _issue(result, 'zero_claim_conflicts_with_event')
                    result['status'] = 'partial'


def inspect_payment_page(content, ticker, as_of, observed_at):
    """Read semantic HTML headers, including the old four-column table shape.

    Amount strings must use Brazilian notation; native workbook numbers are
    numeric. Explicit per-N-share quote units are divided by N, never by a total
    share count. No ID selector or external request is used. An unbounded empty
    table, no table, and a dividend-only zero claim are not confirmed cash zero.
    """
    cutoff, observed = _context(ticker, as_of, observed_at)
    result = _result(content, 'html_tables', cutoff, observed)
    try:
        if not content or not str(content).strip():
            raise ValueError('empty_html')
        tree = html.fromstring(content, parser=html.HTMLParser(no_network=True))
    except (etree.ParserError, ValueError, TypeError):
        _issue(result, 'invalid_html')
        result['status'] = 'invalid'
        return result
    titles = tree.xpath('//title/text()|//meta[@property="og:title"]/@content')
    title_tickers = set(re.findall(r'\b' + _TICKER_PATTERN + r'\b', ' '.join(titles)))
    if title_tickers and ticker not in title_tickers:
        _issue(result, 'page_ticker_mismatch')
        result['status'] = 'invalid'
        return result
    result['report']['page_ticker_verified'] = title_tickers == {ticker}
    if not tree.xpath('//table') and any(token in _norm(tree.text_content()) for token in ('access denied', 'captcha', 'acesso negado')):
        _issue(result, 'blocked_or_challenge_content')
        result['status'] = 'invalid'
        return result
    recognized = 0
    for table_index, table in enumerate(tree.xpath('//table'), 1):
        try:
            rows = list(_table_rows(table))
        except ValueError:
            _issue(result, 'invalid_table_span')
            continue
        columns, header_index, headers = _find_header(rows)
        if not columns:
            continue
        recognized += 1
        for index, values in enumerate(rows[header_index + 1:], header_index + 2):
            if not any(values) or _columns(values):
                continue
            # Table-wide empty messages are claims, not fabricated event rows.
            if len(set(values)) == 1 and re.search(r'nenhum|sem registro|nao ha', _norm(values[0])):
                _issue(result, 'unbounded_empty_table_message', blocking=False)
                continue
            _event(result, values, columns, ticker, cutoff, f'table:{table_index}:row:{index}',
                   amount_header=headers[columns['gross_per_share']])
    result['report']['recognized_tables'] = recognized
    _finish(result)
    _zero_claims(tree, ticker, cutoff, observed, result)
    if not recognized and not result['report']['zero_claims']:
        _issue(result, 'payment_table_not_found', blocking=False)
    return result


def _bounded_zip(content):
    archive = zipfile.ZipFile(io.BytesIO(content))
    entries = archive.infolist()
    if len(entries) > 500 or sum(entry.file_size for entry in entries) > 100_000_000:
        archive.close()
        raise ValueError('archive_too_large')
    if any(entry.flag_bits & 1 or entry.filename.startswith(('/', '\\'))
           or '..' in entry.filename.replace('\\', '/').split('/') for entry in entries):
        archive.close()
        raise ValueError('unsafe_archive')
    return archive


def inspect_petrobras_workbook(content, ticker, as_of, observed_at):
    """Inspect the primary RI XLSX, selecting the exact named PETR3/PETR4 block.

    The URL in SOURCE_DISCOVERIES is linked by issuer RI, not the ADR workbook.
    This accepts bytes only and does not open hyperlinks, macros or external
    workbooks. Native Excel dates are required, not arbitrary serial guessing.
    Formula cells are not silently accepted as a cached declared gross amount.
    """
    from openpyxl import load_workbook

    cutoff, observed = _context(ticker, as_of, observed_at)
    result = _result(content, 'petrobras_ri_xlsx', cutoff, observed)
    result['report']['source_manifest'] = dict(SOURCE_DISCOVERIES[-1])
    result['report']['matches_checked_source_hash'] = result['report']['sha256'] == SOURCE_DISCOVERIES[-1]['sha256']
    if ticker not in ('PETR3', 'PETR4'):
        _issue(result, 'unsupported_exact_issuer_ticker')
        result['status'] = 'invalid'
        return result
    workbook = None
    try:
        with _bounded_zip(content):
            pass
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        blocks = []
        for sheet in workbook:
            if sheet.max_row > 10000 or sheet.max_column > 100:
                raise ValueError('worksheet_too_large')
            rows = list(sheet.values)
            for index, row in enumerate(rows[:8]):
                for start, value in enumerate(row):
                    if not isinstance(value, str) or not re.search(r'\b' + ticker + r'\b', value):
                        continue
                    end = next((column for column in range(start + 1, len(row))
                                if re.search(r'\b' + _TICKER_PATTERN + r'\b', str(row[column] or ''))), len(row))
                    # Exact block heading prevents ON/PN cross-contamination.
                    for header_index in range(index + 1, min(index + 4, len(rows))):
                        headers = list(rows[header_index][start:end])
                        columns = _columns(headers)
                        if columns and 'reported_status' in columns and 'label' in columns:
                            blocks.append((sheet.title, rows, header_index, start, end, headers, columns))
                            break
        if len(blocks) != 1:
            raise ValueError('missing_or_ambiguous_exact_ticker_block')
        sheet, rows, header_index, start, end, headers, columns = blocks[0]
        result['report']['recognized_tables'] = 1
        for index, row in enumerate(rows[header_index + 1:], header_index + 2):
            values = list(row[start:end])
            if not any(value is not None for value in values):
                continue
            reported_label = values[columns['label']] if columns['label'] < len(values) else None
            # The native issuer workbook contains this duplicated-s spelling.
            # Keep this correction local to its exact ON/PN block, not all HTML.
            extraordinary = _norm(reported_label) in (
                'dividendos extraordinarios', 'dividendoss extraordinarios')
            if extraordinary:
                values[columns['label']] = 'DIVIDENDO'
            if not _label(values[columns['label']] if columns['label'] < len(values) else None):
                _issue(result, 'non_income_or_footnote_row', row=index, blocking=False)
                continue
            event_count = len(result['events'])
            _event(result, values, columns, ticker, cutoff, index,
                   amount_header=headers[columns['gross_per_share']])
            if extraordinary and len(result['events']) > event_count:
                result['events'][-1].update(reported_income_label=reported_label,
                    income_label_normalization='exact_Petrobras_extraordinary_dividend_label')
        result['report']['worksheet'] = sheet
        return _finish(result)
    except (ValueError, KeyError, OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        _issue(result, 'invalid_primary_workbook')
        result['status'] = 'invalid'
        return result
    finally:
        if workbook is not None:
            workbook.close()


_CVM_HEADERS = {'Codigo_CVM', 'Data_Referencia', 'Categoria', 'Data_Entrega',
                'Protocolo_Entrega', 'Versao', 'Link_Download'}


def _cvm_url(value):
    parts = urlsplit(value)
    allowed = {'Tela', 'descTipo', 'CodigoInstituicao', 'numProtocolo', 'numSequencia', 'numVersao'}
    query = parse_qsl(parts.query, keep_blank_values=True)
    return (parts.scheme == 'https' and parts.hostname == 'www.rad.cvm.gov.br'
            and not parts.username and not parts.password and parts.port in (None, 443)
            and not parts.fragment and parts.path == '/ENET/frmDownloadDocumento.aspx'
            and bool(query) and len({key for key, _ in query}) == len(query)
            and all(key in allowed and re.fullmatch(r'[A-Za-z0-9]+', item) for key, item in query))


def inspect_cvm_proventos_index(content, cvm, as_of, observed_at):
    """Discover public CVM filings by exact caller-provided CVM identity.

    Accept the official semicolon Latin-1/UTF-8 CSV or its annual ZIP, in memory.
    Documents published after as_of remain separately flagged, never historical
    cash. Preserve every version and the actual (possibly blank) protocol field.
    No filings, or reference dates in the index, do not establish zero income.
    No archive entries are extracted to disk and no document URLs are fetched.
    """
    if isinstance(cvm, bool) or not re.fullmatch(r'\d+', str(cvm)) or int(cvm) <= 0:
        raise ValueError('An exact positive CVM code is required.')
    cutoff, observed = _context('CVMX3', as_of, observed_at)
    result = _result(content, 'cvm_ipe_index', cutoff, observed)
    result['documents'] = []
    result['report']['source_manifest'] = [dict(item) for item in SOURCE_DISCOVERIES[:3]]
    try:
        body = content.encode('utf-8') if isinstance(content, str) else content
        if body.startswith(b'PK'):
            with _bounded_zip(body) as archive:
                members = archive.namelist()
                if len(members) != 1 or not re.fullmatch(r'ipe_cia_aberta_\d{4}\.csv', members[0]):
                    raise ValueError('unexpected_cvm_archive_members')
                body = archive.read(members[0])
        try:
            text = body.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = body.decode('latin-1')
        reader = csv.DictReader(io.StringIO(text), delimiter=';')
        if not reader.fieldnames or not _CVM_HEADERS.issubset(reader.fieldnames) or len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError('invalid_cvm_index_schema')
        for index, row in enumerate(reader, 2):
            if _norm(row.get('Categoria')) != 'relatorio proventos':
                continue
            code = row.get('Codigo_CVM', '')
            if not re.fullmatch(r'\d+', code):
                _issue(result, 'invalid_cvm_identity', row=index)
                continue
            if int(code) != int(cvm):
                continue
            if None in row or any(value is None for value in row.values()):
                _issue(result, 'invalid_cvm_csv_row', row=index)
                continue
            try:
                reference = _day(row['Data_Referencia'][:10])
                published = _day(row['Data_Entrega'][:10])
                if not _cvm_url(row['Link_Download']):
                    raise ValueError('unsafe_document_url')
                if not re.fullmatch(r'\d+', row['Versao']):
                    raise ValueError('invalid_document_version')
                result['documents'].append({'cvm': int(cvm), 'reference_date': reference.isoformat(),
                                            'published_at': published.isoformat(),
                                            'protocol': row['Protocolo_Entrega'] or None,
                                            'version': int(row['Versao']), 'url': row['Link_Download'],
                                            'available_as_of': published <= cutoff,
                                            'status': 'document_reference_only',
                                            'cash_model_eligible': False})
            except (ValueError, TypeError):
                _issue(result, 'invalid_cvm_document_metadata', row=index)
        if result['documents']:
            result['status'] = 'partial'
            _issue(result, 'document_content_not_inspected', blocking=False)
        elif result['report']['issues']:
            result['status'] = 'invalid'
        else:
            _issue(result, 'no_cvm_proventos_document_found', blocking=False)
        result['report']['counts'] = {'documents': len(result['documents']), 'events': 0}
        return result
    except (ValueError, OSError, zipfile.BadZipFile, csv.Error):
        _issue(result, 'invalid_cvm_index')
        result['status'] = 'invalid'
        return result


class _FetchFailure(Exception):
    def __init__(self, code, http_status=None):
        super().__init__(code)
        self.http_status = http_status


def _fetch_primary(url, *, timeout, max_bytes):
    import http.client
    import ssl
    import urllib.request

    class SameHostRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, request, response, code, message, headers, newurl):
            old, new = urlsplit(request.full_url), urlsplit(newurl)
            if new.scheme != 'https' or old.hostname != new.hostname or new.username or new.password:
                raise _FetchFailure('unapproved_redirect', code)
            return super().redirect_request(request, response, code, message, headers, newurl)

    opener = urllib.request.build_opener(SameHostRedirect(),
                                         urllib.request.HTTPSHandler(context=ssl.create_default_context()))
    request = urllib.request.Request(url, headers={'User-Agent': 'RendaFutura-Academic/1.0'})
    with opener.open(request, timeout=timeout) as response:
        status = response.status
        try:
            body = response.read(max_bytes + 1)
        except (http.client.IncompleteRead, OSError):
            raise _FetchFailure('incomplete_response', status) from None
        if len(body) > max_bytes:
            raise _FetchFailure('response_size_limit', status)
        return status, body


def collect_primary_payments(folder, specs, as_of, *, observed_at=None, refresh=False,
                             max_requests=3, max_documents=4, max_bytes=8_000_000,
                             timeout=20, fetch=None):
    """Optional orchestrator API; no existing collector or active pointer touched.

    specs is a list of {ticker: exact B3 ticker, cvm: exact CVM code,
    sources?: ['petrobras_ri_xlsx', 'cvm_ipe_index'], years?: [2026]}.
    The default is issuer XLSX for PETR3/PETR4 (CVM 9512, verified in existing
    read-only identity cache), otherwise the cutoff-year CVM index. At most two
    explicit index years/spec, not later than as_of, are allowed. No URLs or
    filesystem names come from specs. Network is restricted to the known RI
    XLSX and official annual CVM ZIPs; there is NO PDF/form crawling endpoint.

    Returns {by_ticker: {ticker: inspection bundle}, manifest: [acquisitions],
    requests_made, status}. All missing/failed providers remain partial, not
    fabricated zero. CVM references are capped at max_documents per ticker and
    partitioned by publication <= as_of; they are not cash observations. Dates
    for future scheduled payments are retained as announced, never paid.

    Deduplicates shared issuer/index requests; validates before caching. Cache
    files include URL/content hashes and are never overwritten. refresh=True
    attempts once and retains/reuses prior valid bytes on failure. A 403/429
    stops further networking in this call (cached reads still permitted).
    HTTPS uses normal certificate verification; redirects to other hosts are
    rejected. Each public request is bounded by timeout and max_bytes. The
    injected fetch(url, timeout=..., max_bytes=...) -> (HTTP_status, bytes)
    interface exists for tests/parent transport, not alternative endpoints.
    """
    observed = observed_at or datetime.now(timezone.utc)
    cutoff, observed = _context('CVMX3', as_of, observed)
    if (type(max_requests) is not int or not 0 <= max_requests <= 10
            or type(max_documents) is not int or not 1 <= max_documents <= 20
            or type(max_bytes) is not int or not 1 <= max_bytes <= _MAX_BYTES
            or not isinstance(timeout, (int, float)) or not math.isfinite(timeout) or not 0 < timeout <= 60):
        raise ValueError('Invalid bounded acquisition limits.')
    root = Path(folder).expanduser().resolve()
    if not isinstance(specs, (list, tuple)) or len(specs) > 1000:
        raise ValueError('Expected at most 1000 exact identity specs.')
    prepared, identities = [], set()
    for spec in specs:
        if not isinstance(spec, dict):
            raise ValueError('Each spec must contain ticker and CVM identity.')
        ticker, cvm = spec.get('ticker'), spec.get('cvm')
        _context(ticker, cutoff, observed)
        if isinstance(cvm, bool) or not re.fullmatch(r'\d+', str(cvm)) or int(cvm) <= 0 or ticker in identities:
            raise ValueError('Duplicate ticker or invalid exact CVM code.')
        identities.add(ticker)
        sources = spec.get('sources', ['petrobras_ri_xlsx'] if ticker in ('PETR3', 'PETR4') else ['cvm_ipe_index'])
        years = spec.get('years', [cutoff.year])
        if (not isinstance(sources, list) or not sources or len(sources) > 2 or len(set(sources)) != len(sources)
                or any(source not in ('petrobras_ri_xlsx', 'cvm_ipe_index') for source in sources)
                or not isinstance(years, list) or not 1 <= len(years) <= 2
                or len(set(years)) != len(years)
                or any(type(year) is not int or not 2003 <= year <= cutoff.year for year in years)):
            raise ValueError('Only documented primary providers and bounded past index years are accepted.')
        if 'petrobras_ri_xlsx' in sources and (ticker not in ('PETR3', 'PETR4') or int(cvm) != 9512):
            raise ValueError('Petrobras source requires the exact PETR3/PETR4 and CVM 9512 identity.')
        prepared.append((ticker, int(cvm), sources, years))
    output = {'by_ticker': {}, 'manifest': [], 'requests_made': 0, 'status': 'partial'}
    memo, network_stopped = {}, False
    transport = fetch or _fetch_primary

    def acquire(url, suffix, validate):
        nonlocal network_stopped
        if url in memo:
            return memo[url]
        prefix = 'primary-payments-' + hashlib.sha256(url.encode('ascii')).hexdigest()[:16] + '-'
        cached = []
        if root.is_dir():
            candidates = (item for item in root.glob(prefix + '*' + suffix) if not item.is_symlink() and item.is_file())
            cached = sorted(islice(candidates, 50), key=lambda item: item.stat().st_mtime, reverse=True)

        def previous():
            for path in cached:
                try:
                    if path.is_symlink() or path.resolve().parent != root or path.stat().st_size > max_bytes:
                        continue
                    body = path.read_bytes()
                    digest = hashlib.sha256(body).hexdigest()
                    retrieved = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc)
                    if path.name != prefix + digest + suffix or retrieved.date() < cutoff:
                        continue
                    if validate(body, retrieved)['status'] != 'invalid':
                        return body, retrieved, path.name
                except (ValueError, OSError):
                    continue
            return None

        prior = previous()
        entry = {'url': url, 'http_status': None, 'status': 'not_acquired'}
        output['manifest'].append(entry)
        selected = prior if not refresh else None
        if selected:
            entry['status'] = 'cached'
        elif output['requests_made'] >= max_requests or network_stopped:
            entry['status'] = 'request_budget_or_access_stop'
        else:
            output['requests_made'] += 1
            try:
                status, body = transport(url, timeout=timeout, max_bytes=max_bytes)
                entry['http_status'] = status
                if status in (403, 429):
                    network_stopped = True
                if status != 200 or not isinstance(body, bytes) or len(body) > max_bytes:
                    raise _FetchFailure('invalid_public_response', status)
                if validate(body, observed)['status'] == 'invalid':
                    raise _FetchFailure('invalid_public_content', status)
                root.mkdir(parents=True, exist_ok=True)
                name = prefix + hashlib.sha256(body).hexdigest() + suffix
                path = root / name
                # Exclusive creation cannot replace a previous download or pointer.
                try:
                    with path.open('xb') as stream:
                        stream.write(body)
                except FileExistsError:
                    if path.is_symlink() or path.read_bytes() != body:
                        raise _FetchFailure('immutable_cache_conflict', status) from None
                selected = body, observed, name
                entry['status'] = 'downloaded_valid'
            except Exception as error:
                # Never log exception text: URL/query/response text may be sensitive.
                entry['http_status'] = getattr(error, 'http_status', getattr(error, 'code', entry['http_status']))
                entry['error_class'] = type(error).__name__
                entry['status'] = 'refresh_failed'
                if entry['http_status'] in (403, 429):
                    network_stopped = True
        if selected is None and prior is not None:
            selected = prior
            entry['status'] += '_cached_prior_retained'
        if selected:
            entry.update(file=selected[2], sha256=hashlib.sha256(selected[0]).hexdigest(),
                         bytes=len(selected[0]), retrieved_at=selected[1].isoformat())
        memo[url] = selected
        return selected

    for ticker, cvm, sources, years in prepared:
        bundle = _result(b'', 'primary_source_bundle', cutoff, observed)
        bundle['report'].pop('sha256')
        bundle['report'].pop('bytes')
        bundle.update(documents=[], future_documents=[], future_events=[], source_results=[])
        failed = False
        for source in sources:
            for year in (years if source == 'cvm_ipe_index' else [None]):
                if source == 'petrobras_ri_xlsx':
                    url, suffix = SOURCE_DISCOVERIES[-1]['url'], '.xlsx'
                    reader = lambda body, seen: inspect_petrobras_workbook(body, ticker, cutoff, seen)
                else:
                    url = f'https://dados.cvm.gov.br/dados/CIA_ABERTA/DOC/IPE/DADOS/ipe_cia_aberta_{year}.zip'
                    suffix = '.zip'
                    reader = lambda body, seen: inspect_cvm_proventos_index(body, cvm, cutoff, seen)
                acquired = acquire(url, suffix, reader)
                if acquired is None:
                    failed = True
                    _issue(bundle, 'primary_source_unavailable')
                    continue
                acquisition = next(item for item in output['manifest'] if item['url'] == url)
                if acquisition['status'] not in ('downloaded_valid', 'cached'):
                    _issue(bundle, 'primary_refresh_failed_cached_prior_retained')
                parsed = reader(acquired[0], acquired[1])
                documents = parsed.get('documents', [])
                known = sorted((item for item in documents if item['available_as_of']),
                               key=lambda item: (item['published_at'], item['version'], item['reference_date']), reverse=True)
                future = sorted((item for item in documents if not item['available_as_of']),
                                key=lambda item: (item['published_at'], item['version']), reverse=True)
                remaining = max(0, max_documents - len(bundle['documents']))
                parsed['documents'], parsed['future_documents'] = known[:remaining], future[:max_documents]
                if len(known) > remaining:
                    _issue(parsed, 'document_reference_limit')
                bundle['documents'].extend(parsed['documents'])
                bundle['future_documents'].extend(parsed['future_documents'][:max(0, max_documents - len(bundle['future_documents']))])
                bundle['events'].extend(item for item in parsed['events'] if item['available_as_of'] is not False)
                bundle['future_events'].extend(item for item in parsed['events'] if item['available_as_of'] is False)
                bundle['report']['issues'].extend(parsed['report']['issues'])
                bundle['source_results'].append(parsed)
        _finish(bundle)
        if failed or not bundle['events'] or any(source['status'] == 'invalid' for source in bundle['source_results']):
            bundle['status'] = 'partial'
        output['by_ticker'][ticker] = bundle
    if output['by_ticker'] and all(item['status'] == 'available' for item in output['by_ticker'].values()):
        output['status'] = 'available'
    return output
