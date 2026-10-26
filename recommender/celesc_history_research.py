"""Offline Celesc history audit; no cash/date repair or model admission."""

import argparse
import csv
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile

import openpyxl

from .selected_issuer_sources import CACHE, catalog_request, history_request, source_body
from .selected_issuer_sources import links
from .dividend_scope import allowed_tickers


ROOT = Path(__file__).parent / 'data/market'
CUTOFF = date(2026, 9, 30)
START = date(2021, 9, 30)
WORKBOOK_SHA = 'debb630e8d00a7c3382ded0a3d92aac7fc04e8024d5b9ef8db434963623772dc'
HISTORY_SHA = '89b9d035338548ebe2c26b4bac0e61bb2aec59de0a8bdfdd42de140a98b69464'
HEADERS = ('Exerc\u00edcio Social', 'Tipo Provento', 'Data Aprova\u00e7\u00e3o AGO / RCA',
           'Data Ex-Dividendo', 'Data de In\u00edcio do Pagamento', 'Provento por A\u00e7\u00e3o ON (R$)',
           'Provento por A\u00e7\u00e3o PN (R$)', 'Valor Total Distribu\u00eddo (R$)')


def _require(condition, reason):
    if not condition:
        raise ValueError(reason)


def _day(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        _require(value.time().isoformat() == '00:00:00', 'non_midnight_workbook_date')
        return value.date().isoformat()
    if type(value) is date:
        return value.isoformat()
    _require(isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2} 00:00:00', value),
             'unrecognized_workbook_date')
    return datetime.fromisoformat(value).date().isoformat()


def inspect_workbook(content):
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        _require(workbook.sheetnames == ['Plan1'], 'unreviewed_celesc_worksheets')
        sheet = workbook.worksheets[0]
        _require(tuple(sheet.cell(11, i).value for i in range(1, 9)) == HEADERS
                 and sheet['A9'].value == 'Atualizado em 25/09/2026', 'unreviewed_celesc_workbook_headers')
        output, issues = [], []
        for number, values in enumerate(sheet.iter_rows(min_row=12, values_only=True), 12):
            fiscal, label, approval, ex_day, payment, on, pn, total = values[:8]
            _require(type(fiscal) in (int, float) and int(fiscal) == fiscal
                     and isinstance(label, str) and label.split()[0] in ('DIV', 'JCP'), 'invalid_cash_row_identity')
            amounts = []
            for amount in (on, pn, total):
                _require(type(amount) in (int, float) and math.isfinite(amount) and amount >= 0,
                         'invalid_reported_workbook_amount')
                amounts.append(str(Decimal(str(amount))))
            approval, ex_day, payment = map(_day, (approval, ex_day, payment))
            row_issues = []
            if approval and ex_day and ex_day < approval:
                row_issues.append('reported_ex_date_precedes_approval')
            if ex_day and payment and payment < ex_day:
                row_issues.append('reported_payment_start_precedes_ex_date')
            relevant = bool(payment and START.isoformat() <= payment <= CUTOFF.isoformat())
            for reason in row_issues:
                issues.append({'row': number, 'reason': reason, 'in_cash_window': relevant,
                               'approval_date': approval, 'reported_ex_date': ex_day,
                               'reported_payment_start': payment})
            output.append({'row': number, 'fiscal_year': int(fiscal), 'reported_label': label,
                           'cash_kind': label.split()[0], 'approval_date': approval,
                           'reported_ex_date': ex_day, 'record_date': None,
                           'reported_payment_start': payment, 'reported_ON_amount': amounts[0],
                           'reported_PN_amount': amounts[1], 'reported_total_brl': amounts[2],
                           'tax_basis': None, 'settlement_confirmed': False,
                           'in_cash_window': relevant, 'date_issues': row_issues,
                           'future_payment_excluded_from_window': bool(payment and payment > CUTOFF.isoformat())})
        return {'source_updated_date': '2026-09-25', 'events': output, 'issues': issues,
                'record_date_derived_from_ex_date': False, 'gross_inferred_from_amount': False,
                'window_rows': sum(row['in_cash_window'] for row in output),
                'future_rows': sum(row['future_payment_excluded_from_window'] for row in output)}
    finally:
        workbook.close()


def inspect_prices(content, secondary):
    payload = json.loads(content)
    data = payload.get('data', {})
    _require(payload.get('success') is True and data.get('errors') == []
             and data.get('source') == 'Enfoque' and data.get('link') == 'https://www.enfoque.com.br'
             and [row.get('ticker') for row in data.get('tickers', [])] == ['CLSC4'],
             'issuer_history_identity_or_provider_mismatch')
    daily, monthly = {}, {}
    for row in data['tickers'][0]['history']:
        day, close = row.get('date'), row.get('closed')
        _require(isinstance(day, str) and date.fromisoformat(day).isoformat() == day
                 and '2021-09-01' <= day <= CUTOFF.isoformat() and day not in daily
                 and type(close) in (int, float) and math.isfinite(close) and close > 0,
                 'invalid_or_duplicate_issuer_price')
        daily[day] = close
        if day > monthly.get(day[:7], ('', 0))[0]:
            monthly[day[:7]] = day, close
    expected = {f'{year:04d}-{month:02d}' for year in range(2021, 2027) for month in range(1, 13)
                if '2021-09' <= f'{year:04d}-{month:02d}' <= '2026-09'}
    _require(set(monthly) == expected and len(expected) == 61, 'issuer_history_missing_61_monthly_anchors')
    chart = json.loads(secondary)['chart']['result'][0]
    _require(chart['meta']['symbol'] == 'CLSC4.SA' and chart['meta']['currency'] == 'BRL',
             'secondary_price_identity_or_currency_mismatch')
    other = {datetime.fromtimestamp(stamp, timezone.utc).date().isoformat(): close
             for stamp, close in zip(chart['timestamp'], chart['indicators']['quote'][0]['close'])
             if close is not None}
    compared = [(day, value, other[day]) for day, value in sorted(daily.items()) if day in other]
    conflicts = [{'date': day, 'issuer_feed': first, 'secondary': second}
                 for day, first, second in compared
                 if not math.isclose(first, second, rel_tol=.01, abs_tol=.015)]
    return {'ticker': 'CLSC4', 'reported_provider': 'Enfoque', 'reported_adjusted_request': False,
            'source_rows': len(daily), 'first_date': min(daily), 'last_date': max(daily),
            'consecutive_monthly_anchors': len(monthly), 'compared_daily_quotes': len(compared),
            'conflicts': conflicts, 'same_existing_price_tolerances': {'rtol': .01, 'atol': .015},
            'monthly_anchors': [{'month': month, 'date': day, 'close': close}
                                for month, (day, close) in sorted(monthly.items())],
            'secondary_reported_split_count': len(chart.get('events', {}).get('splits', {})),
            'absence_of_all_corporate_actions_proven': False}


def candidate_cvm_notices(folder):
    """Locate latest eligible bodies across all cached categories, not fetch them."""
    from .cvm_proventos_documents import day, document_id
    index_root = folder / 'raw/primary-research/cvm-documents'
    manifest_file = index_root / 'index-sources.json'
    manifest_body = manifest_file.read_bytes()
    manifest = json.loads(manifest_body)
    found, sources = {}, []
    for source in manifest:
        if source.get('year') not in range(2021, 2027):
            continue
        year = source['year']
        expected_url = f'https://dados.cvm.gov.br/dados/CIA_ABERTA/DOC/IPE/DADOS/ipe_cia_aberta_{year}.zip'
        path = folder / source['file']
        _require(source['url'] == expected_url and not path.is_symlink()
                 and path.resolve().is_relative_to(folder.resolve()), 'invalid_CVM_notice_index')
        content = path.read_bytes()
        _require(hashlib.sha256(content).hexdigest() == source['sha256'], 'CVM_notice_index_changed')
        sources.append({'file': source['file'], 'sha256': source['sha256'], 'url': source['url'],
                        'bytes': len(content), 'retrieved_at': source.get('retrieved_at')})
        member = f'ipe_cia_aberta_{year}.csv'
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            with archive.open(member) as stream:
                for line, row in enumerate(csv.DictReader(io.TextIOWrapper(stream, encoding='latin1'), delimiter=';'), 2):
                    if not row['Codigo_CVM'].isdigit() or int(row['Codigo_CVM']) != 2461:
                        continue
                    reference, published = day(row['Data_Referencia']), day(row['Data_Entrega'])
                    if (not published or published > CUTOFF.isoformat() or not reference
                            or not ('2025-06-15' <= reference <= '2025-06-30'
                                    or '2025-09-15' <= reference <= '2025-09-30')):
                        continue
                    sequence, version = document_id(row['Link_Download'])
                    _require(version == int(row['Versao']), 'CVM_notice_index_version_mismatch')
                    entry = {'sequence': sequence, 'version': version, 'cvm': 2461,
                             'reference_date': reference, 'published_at': published,
                             'url': row['Link_Download'], 'category': row['Categoria'],
                             'subject': row['Assunto'], 'raw_row': row, 'csv_line': line,
                             'index_file': source['file'], 'index_sha256': source['sha256'],
                             'member': member}
                    if sequence not in found or version > found[sequence]['version']:
                        found[sequence] = entry
    for entry in found.values():
        entry['cached_body_paths_not_yet_validated'] = [path.relative_to(folder).as_posix()
            for path in sorted(index_root.glob(f"{entry['sequence']}-v{entry['version']}-*.pdf"))]
    sources.append({'file': manifest_file.relative_to(folder).as_posix(),
                    'sha256': hashlib.sha256(manifest_body).hexdigest(),
                    'bytes': len(manifest_body), 'authority': 'local_index_manifest'})
    return sorted(found.values(), key=lambda row: (row['reference_date'], row['sequence'])), sources


def audit(folder=ROOT):
    folder = Path(folder)
    selected = allowed_tickers(folder)
    _require(selected is not None and bool(selected & {'CLSC3', 'CLSC4'}), 'Celesc_outside_selected_positive_DY_scope')
    base = folder / CACHE
    receipts = {number: json.loads((base / f'response-{number:03d}.json').read_bytes())
                for number in (2, 3, 4, 7, 9, 11, 12)}
    _require(catalog_request(folder, receipts[2], receipts[4])[0] == receipts[7]['url']
             and history_request(folder, receipts[3], receipts[9]) == receipts[12]['url']
             and all(row['status'] == 200 and not row['stop'] for row in receipts.values())
             and receipts[11]['parent'] == receipts[7]
             and receipts[11]['url'] in {row['url'] for row in links(source_body(folder, receipts[7]), receipts[7]['url'])}
             and receipts[12]['history'] == {'page': receipts[3], 'bundle': receipts[9]}
             and receipts[11]['sha256'] == WORKBOOK_SHA and receipts[12]['sha256'] == HISTORY_SHA,
             'celesc_original_source_binding_changed')
    workbook = inspect_workbook(source_body(folder, receipts[11]))
    price_file = 'raw/expanded/CLSC4-prices.json'
    if not (folder / price_file).is_file():
        price_file = 'raw/inventory/daily-prices/CLSC4.json'
    raw_price = (folder / price_file).read_bytes()
    prices = inspect_prices(source_body(folder, receipts[12]), raw_price)
    notices, index_sources = candidate_cvm_notices(folder)
    return {'scope': 'CLSC3_CLSC4_positive_DY_source_research', 'reference_date': CUTOFF.isoformat(),
            'status': 'source_evidence_only_not_model_admission', 'cash_workbook': workbook,
            'price_concordance': prices, 'network_calls_by_audit': 0,
            'missing_data_imputed': False, 'model_admissions': [],
            'candidate_primary_notices_for_date_conflicts': notices,
            'required_before_cash_admission': [
                'Bind exact share class, record dates and tax basis to primary declarations.',
                'Resolve contradictory ex/approval/payment-start dates without guessing corrections.',
                'Crosscheck the full cash window and corporate-action/share basis under the normal parent checks.',
                'Do not treat a reported payment start as confirmed full settlement.'],
            'sources': [{'url': row['url'], 'file': row['file'], 'sha256': row['sha256'],
                         'bytes': row['bytes'], 'retrieved_at': row['completed_at']}
                        for row in receipts.values()] + index_sources + [
                            {'file': price_file, 'sha256': hashlib.sha256(raw_price).hexdigest(),
                             'bytes': len(raw_price), 'retrieved_at': None,
                             'authority': 'cached_secondary_price_source'}]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--folder', type=Path, default=ROOT)
    args = parser.parse_args()
    report = audit(args.folder)
    content = (json.dumps(report, ensure_ascii=True, sort_keys=True, indent=2, allow_nan=False) + '\n').encode()
    path = args.folder / CACHE / ('celesc-audit-' + hashlib.sha256(content).hexdigest() + '.json')
    if path.exists():
        _require(path.read_bytes() == content, 'immutable_celesc_audit_changed')
    else:
        with path.open('xb') as stream:
            stream.write(content)
    print(json.dumps({'file': str(path), 'price_days': report['price_concordance']['source_rows'],
                      'price_conflicts': len(report['price_concordance']['conflicts']),
                      'cash_window_rows': report['cash_workbook']['window_rows'],
                      'cash_issues': report['cash_workbook']['issues'], 'admissions': []}))


if __name__ == '__main__':
    main()
