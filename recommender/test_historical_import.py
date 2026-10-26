import csv
from datetime import date, datetime
import hashlib
import json
import zipfile

import openpyxl
import pytest

from .historical_import import (COLUMNS, ImportValidationError, emit_templates,
                                freeze_import_bundle, load_import_bundle, main)


CUTOFF = '2024-12-31'


def manifest():
    return {'schema_version': 1, 'source_name': 'Public test fixture',
            'source_url': 'https://example.org/history.csv', 'retrieved_at': '2025-01-02T12:00:00Z',
            'currency': 'BRL', 'price_basis': 'raw', 'cash_share_basis': 'record_date',
            'history_start': '2020-01-01', 'history_end': CUTOFF,
            'tax_basis': 'gross_before_individual_taxes', 'decimal_format': 'dot', 'csv_delimiter': ',',
            'tickers': {'TEST4': {'cvm': 123, 'share_class': 'PN', 'coverage': {
                table: {'start': '2020-01-01', 'end': CUTOFF, 'complete': True} for table in COLUMNS}}}}


def records():
    return {'prices': [{'ticker': 'TEST4', 'date': '2024-06-03', 'close': '10.25'}],
            'payments': [{'ticker': 'TEST4', 'record_date': '2024-06-10', 'payment_date': '2024-06-20',
                          'gross_per_share': '0.25', 'label': 'DIVIDENDO'}],
            'actions': [{'ticker': 'TEST4', 'effective_date': '2024-05-01',
                         'numerator': '2', 'denominator': '1', 'type': 'split'}],
            'fundamentals': [{'ticker': 'TEST4', 'cvm': '123', 'share_class': 'PN', 'fiscal_end': '2023-12-31',
                              'published_at': '2024-03-01', 'share_basis_date': CUTOFF, 'eps': '2.0',
                              'free_cash_per_share': '1.30', 'liabilities_to_assets': '0.7', 'sector_code': '3',
                              'scope': 'consolidated', 'source_ref': 'https://example.org/filing/2023'}]}


def write_manifest(folder, source):
    (folder / 'source.json').write_text(json.dumps(source), encoding='utf8')


def write_table(folder, table, rows, *, delimiter=',', encoding='utf8'):
    with (folder / (table + '.csv')).open('w', encoding=encoding, newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=COLUMNS[table], delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)


@pytest.fixture
def bundle(tmp_path):
    folder = tmp_path / 'input'
    folder.mkdir()
    write_manifest(folder, manifest())
    for table, rows in records().items():
        write_table(folder, table, rows)
    return folder


def hashes(folder):
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.iterdir() if p.is_file()}


def test_csv_interface_provenance_and_inputs_unchanged(bundle):
    before = hashes(bundle)
    result = load_import_bundle(bundle, date.fromisoformat(CUTOFF))
    assert result['prices'] == [{'ticker': 'TEST4', 'date': '2024-06-03', 'close': 10.25}]
    assert result['payments'][0]['gross_per_share'] == .25
    assert result['actions'][0]['numerator'] == 2
    assert result['fundamentals'][0]['cvm'] == 123
    assert result['fundamentals'][0]['free_cash_per_share'] == 1.3
    assert result['future_payments'] == []
    provenance = result['provenance']
    assert not provenance['cash_imputed']
    assert not provenance['share_adjustments_applied']
    assert not provenance['completeness_certified']
    assert provenance['missing_tables'] == provenance['empty_tables'] == []
    assert {row['file']: row['sha256'] for row in provenance['input_files']} == before
    assert provenance['per_ticker']['TEST4']['counts'] == {
        'prices': 1, 'payments': 1, 'actions': 1, 'fundamentals': 1, 'future_payments': 0}
    assert hashes(bundle) == before


def test_missing_and_empty_tables_are_distinct_not_zero_cash(bundle):
    (bundle / 'payments.csv').unlink()
    write_table(bundle, 'actions', [])
    result = load_import_bundle(bundle, CUTOFF)
    assert result['payments'] == result['actions'] == []
    assert result['provenance']['missing_tables'] == ['payments']
    assert result['provenance']['empty_tables'] == ['actions']
    assert not result['provenance']['completeness_certified']
    assert not result['provenance']['cash_imputed']


def test_all_tables_optional_but_reported(bundle):
    for table in COLUMNS:
        (bundle / (table + '.csv')).unlink()
    result = load_import_bundle(bundle, CUTOFF)
    assert result['provenance']['missing_tables'] == list(COLUMNS)
    assert result['provenance']['per_ticker']['TEST4']['tables_without_rows'] == list(COLUMNS)


@pytest.mark.parametrize('field,value', [
    ('schema_version', True), ('schema_version', 2), ('currency', 'USD'),
    ('price_basis', 'total_return'), ('cash_share_basis', 'unknown'),
    ('tax_basis', 'net'), ('decimal_format', 'auto'), ('csv_delimiter', '|'),
    ('retrieved_at', '2099-01-01T00:00:00Z'), ('retrieved_at', '2025-01-02T12:00:00'),
    ('history_start', '2025-01-01'), ('history_end', '2025-01-01'),
    ('source_url', 'http://example.org/'), ('source_url', 'https://user:SECRET@example.org/'),
    ('source_url', 'https://example.org/?token=SECRET'), ('source_url', 'https://example.org/#SECRET'),
    ('tickers', {}), ('files', {'prices': '../../secret.csv'}),
])
def test_invalid_manifest_rejected_without_echoing_values(bundle, field, value):
    source = manifest()
    source[field] = value
    write_manifest(bundle, source)
    with pytest.raises(ImportValidationError) as caught:
        load_import_bundle(bundle, CUTOFF)
    assert 'SECRET' not in str(caught.value)


@pytest.mark.parametrize('field', [field for field in manifest() if field != 'csv_delimiter'])
def test_required_metadata_cannot_be_guessed(bundle, field):
    source = manifest()
    del source[field]
    write_manifest(bundle, source)
    with pytest.raises(ImportValidationError, match='missing_or_unknown_fields'):
        load_import_bundle(bundle, CUTOFF)


def test_benign_url_query_allowed(bundle):
    source = manifest()
    source['source_url'] = 'https://example.org/history?year=2024&format=csv'
    write_manifest(bundle, source)
    assert load_import_bundle(bundle, CUTOFF)['provenance']['source']['source_url'] == source['source_url']


@pytest.mark.parametrize('field,value,date_field', [
    ('price_basis', 'split_adjusted', 'price_basis_date'),
    ('cash_share_basis', 'adjusted', 'cash_share_basis_date'),
])
def test_adjusted_basis_requires_exact_cutoff(bundle, field, value, date_field):
    source = manifest()
    source[field] = value
    write_manifest(bundle, source)
    with pytest.raises(ImportValidationError):
        load_import_bundle(bundle, CUTOFF)
    source[date_field] = '2024-12-30'
    write_manifest(bundle, source)
    with pytest.raises(ImportValidationError, match='must_equal_cutoff'):
        load_import_bundle(bundle, CUTOFF)
    source[date_field] = CUTOFF
    write_manifest(bundle, source)
    result = load_import_bundle(bundle, CUTOFF)
    assert result['provenance']['source'][date_field] == CUTOFF
    assert result['prices'][0]['close'] == 10.25


def test_unadjusted_basis_cannot_claim_adjustment_date(bundle):
    source = manifest()
    source['price_basis_date'] = CUTOFF
    write_manifest(bundle, source)
    with pytest.raises(ImportValidationError, match='unadjusted_basis'):
        load_import_bundle(bundle, CUTOFF)


def test_adjustment_basis_cannot_postdate_retrieval(bundle):
    source = manifest()
    source['price_basis'] = 'split_adjusted'
    source['price_basis_date'] = CUTOFF
    source['retrieved_at'] = '2024-07-01T00:00:00Z'
    source['history_end'] = '2024-06-30'
    for coverage in source['tickers']['TEST4']['coverage'].values():
        coverage['end'] = '2024-06-30'
    write_manifest(bundle, source)
    with pytest.raises(ImportValidationError, match='basis_after_retrieval'):
        load_import_bundle(bundle, CUTOFF)


def test_future_payments_separate_from_paid_cash(bundle):
    rows = records()['payments']
    rows.append({**rows[0], 'record_date': '2025-01-01', 'payment_date': '2025-01-15', 'gross_per_share': '9.00'})
    write_table(bundle, 'payments', rows)
    result = load_import_bundle(bundle, CUTOFF)
    assert len(result['payments']) == len(result['future_payments']) == 1
    assert sum(row['gross_per_share'] for row in result['payments']) == .25
    assert result['future_payments'][0]['gross_per_share'] == 9


@pytest.mark.parametrize('field', ['record_date', 'payment_date'])
def test_unknown_payment_dates_rejected_not_zero(bundle, field):
    rows = records()['payments']
    rows[0][field] = ''
    write_table(bundle, 'payments', rows)
    with pytest.raises(ImportValidationError, match='expected_ISO_date'):
        load_import_bundle(bundle, CUTOFF)


@pytest.mark.parametrize('table,field', [('prices', 'date'), ('actions', 'effective_date'),
                                       ('fundamentals', 'published_at')])
def test_future_training_rows_rejected(bundle, table, field):
    rows = records()[table]
    rows[0][field] = '2025-01-01'
    write_table(bundle, table, rows)
    with pytest.raises(ImportValidationError, match='future_training_data'):
        load_import_bundle(bundle, CUTOFF)


@pytest.mark.parametrize('value', ['', 'nan', 'NaN', 'inf', '-inf', '1e3', '1,234.50', '1.234,50', '0', '-1'])
def test_strict_positive_price_numbers(bundle, value):
    rows = records()['prices']
    rows[0]['close'] = value
    write_table(bundle, 'prices', rows)
    with pytest.raises(ImportValidationError):
        load_import_bundle(bundle, CUTOFF)


def test_explicit_comma_decimals_and_utf8_bom(bundle):
    source = manifest()
    source['decimal_format'] = 'comma'
    source['csv_delimiter'] = ';'
    write_manifest(bundle, source)
    for table, rows in records().items():
        for row in rows:
            for key in ('close', 'gross_per_share', 'eps', 'free_cash_per_share', 'liabilities_to_assets'):
                if key in row:
                    row[key] = row[key].replace('.', ',')
        write_table(bundle, table, rows, delimiter=';', encoding='utf-8-sig')
    result = load_import_bundle(bundle, CUTOFF)
    assert result['prices'][0]['close'] == 10.25
    assert result['fundamentals'][0]['free_cash_per_share'] == 1.3


@pytest.mark.parametrize('table', list(COLUMNS))
def test_duplicate_row_identities_rejected_even_when_amount_differs(bundle, table):
    rows = records()[table]
    rows.append(dict(rows[0]))
    if table == 'prices':
        rows[-1]['close'] = '12'
    write_table(bundle, table, rows)
    with pytest.raises(ImportValidationError, match='duplicate_row_identity'):
        load_import_bundle(bundle, CUTOFF)


@pytest.mark.parametrize('field,value', [('cvm', '999'), ('share_class', 'ON'), ('scope', 'individual')])
def test_mixed_fundamental_identity_or_scope_rejected(bundle, field, value):
    rows = records()['fundamentals']
    rows.append({**rows[0], 'published_at': '2024-04-01', field: value})
    write_table(bundle, 'fundamentals', rows)
    with pytest.raises(ImportValidationError, match='mixed_ticker_identity_or_scope'):
        load_import_bundle(bundle, CUTOFF)


def test_undeclared_ticker_is_not_inferred(bundle):
    rows = records()['prices']
    rows[0]['ticker'] = 'OTHER3'
    write_table(bundle, 'prices', rows)
    with pytest.raises(ImportValidationError, match='undeclared_ticker'):
        load_import_bundle(bundle, CUTOFF)


def test_empty_fcf_allowed_only_for_financial_sector_and_preserved_null(bundle):
    rows = records()['fundamentals']
    rows[0]['free_cash_per_share'] = ''
    write_table(bundle, 'fundamentals', rows)
    with pytest.raises(ImportValidationError, match='expected_finite_number'):
        load_import_bundle(bundle, CUTOFF)
    rows[0]['sector_code'] = '2'
    rows[0]['eps'] = '-2.0'
    rows[0]['liabilities_to_assets'] = '1.2'
    write_table(bundle, 'fundamentals', rows)
    row = load_import_bundle(bundle, CUTOFF)['fundamentals'][0]
    assert row['free_cash_per_share'] is None
    assert row['eps'] == -2.0
    assert row['liabilities_to_assets'] == 1.2


@pytest.mark.parametrize('label', ['AMORTIZACAO', 'RESGATE', 'PRINCIPAL', 'UNKNOWN'])
def test_principal_or_unknown_labels_cannot_enter_income(bundle, label):
    rows = records()['payments']
    rows[0]['label'] = label
    write_table(bundle, 'payments', rows)
    with pytest.raises(ImportValidationError, match='unsupported_value'):
        load_import_bundle(bundle, CUTOFF)


@pytest.mark.parametrize('numerator,denominator,kind', [('0', '1', 'split'), ('1', '0', 'split'),
                                                     ('1', '2', 'split'), ('2', '1', 'reverse_split'),
                                                     ('1', '1', 'bonus'), ('2.0', '1', 'split')])
def test_invalid_share_actions_rejected(bundle, numerator, denominator, kind):
    rows = records()['actions']
    rows[0].update(numerator=numerator, denominator=denominator, type=kind)
    write_table(bundle, 'actions', rows)
    with pytest.raises(ImportValidationError):
        load_import_bundle(bundle, CUTOFF)


def test_coverage_required_for_rows_and_not_upgraded(bundle):
    source = manifest()
    del source['tickers']['TEST4']['coverage']['payments']
    write_manifest(bundle, source)
    with pytest.raises(ImportValidationError, match='missing_ticker_table_coverage'):
        load_import_bundle(bundle, CUTOFF)
    source['tickers']['TEST4']['coverage']['payments'] = {'start': '2020-01-01', 'end': CUTOFF, 'complete': False}
    write_manifest(bundle, source)
    result = load_import_bundle(bundle, CUTOFF)
    assert not result['provenance']['per_ticker']['TEST4']['declared_coverage']['payments']['complete']


def write_workbook(path, all_rows):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for table, rows in all_rows.items():
        sheet = book.create_sheet(table)
        sheet.append(COLUMNS[table])
        for row in rows:
            values = []
            for field in COLUMNS[table]:
                value = row[field]
                if field in ('date', 'record_date', 'payment_date', 'effective_date',
                             'fiscal_end', 'published_at', 'share_basis_date'):
                    value = datetime.fromisoformat(value)
                elif field in ('close', 'gross_per_share', 'eps', 'free_cash_per_share', 'liabilities_to_assets'):
                    value = float(value)
                elif field in ('cvm', 'numerator', 'denominator', 'sector_code'):
                    value = int(value)
                values.append(value)
            sheet.append(values)
    book.save(path)
    book.close()


def test_xlsx_native_dates_and_numbers_match_csv(bundle):
    expected = load_import_bundle(bundle, CUTOFF)
    for table in COLUMNS:
        (bundle / (table + '.csv')).unlink()
    write_workbook(bundle / 'tables.xlsx', records())
    actual = load_import_bundle(bundle, CUTOFF)
    for table in (*COLUMNS, 'future_payments'):
        assert actual[table] == expected[table]
    assert [row['file'] for row in actual['provenance']['input_files']] == ['source.json', 'tables.xlsx']


def test_individual_xlsx_and_ambiguous_tables(bundle):
    write_workbook(bundle / 'prices.xlsx', {'prices': records()['prices']})
    with pytest.raises(ImportValidationError, match='ambiguous_table_files'):
        load_import_bundle(bundle, CUTOFF)
    (bundle / 'prices.csv').unlink()
    assert load_import_bundle(bundle, CUTOFF)['prices'][0]['close'] == 10.25


def test_xlsx_formula_is_never_evaluated_or_read_as_cached_value(bundle):
    (bundle / 'prices.csv').unlink()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(COLUMNS['prices'])
    sheet.append(['TEST4', '2024-06-03', '=1+2'])
    book.save(bundle / 'prices.xlsx')
    book.close()
    with pytest.raises(ImportValidationError, match='formula_or_excel_error'):
        load_import_bundle(bundle, CUTOFF)


def test_xlsx_oversized_formatted_range_rejected_before_row_scan(bundle):
    (bundle / 'prices.csv').unlink()
    write_workbook(bundle / 'prices.xlsx', {'prices': records()['prices']})
    book = openpyxl.load_workbook(bundle / 'prices.xlsx')
    book.active.cell(100000, 10000).number_format = '0.00'
    book.save(bundle / 'prices.xlsx')
    book.close()
    with pytest.raises(ImportValidationError, match='unexpected_excel_columns'):
        load_import_bundle(bundle, CUTOFF)


def test_xlsx_external_links_package_rejected(bundle):
    (bundle / 'prices.csv').unlink()
    write_workbook(bundle / 'prices.xlsx', {'prices': records()['prices']})
    with zipfile.ZipFile(bundle / 'prices.xlsx', 'a') as archive:
        archive.writestr('xl/externalLinks/externalLink1.xml', '<externalLink/>')
    with pytest.raises(ImportValidationError, match='unsafe_workbook_package'):
        load_import_bundle(bundle, CUTOFF)


@pytest.mark.parametrize('filename', ['../outside.xml', 'C:/outside.xml', '/outside.xml'])
def test_xlsx_traversal_members_rejected(bundle, filename):
    (bundle / 'prices.csv').unlink()
    write_workbook(bundle / 'prices.xlsx', {'prices': records()['prices']})
    with zipfile.ZipFile(bundle / 'prices.xlsx', 'a') as archive:
        archive.writestr(filename, '<data/>')
    with pytest.raises(ImportValidationError, match='unsafe_workbook_package'):
        load_import_bundle(bundle, CUTOFF)


def test_malformed_xlsx_errors_do_not_expose_bytes(bundle, capsys):
    (bundle / 'prices.csv').unlink()
    (bundle / 'prices.xlsx').write_bytes(b'INVALID_SECRET_DATA')
    assert main(['validate', str(bundle), '--cutoff', CUTOFF, '--workspace', str(bundle.parent), '--output', 'bad']) == 2
    logged = capsys.readouterr()
    assert 'INVALID_SECRET_DATA' not in logged.out + logged.err
    assert 'invalid_workbook' in logged.out


@pytest.mark.parametrize('body', [b'date,ticker,close\n', b'ticker,date,date\n',
                                b'ticker,date,close\nTEST4,2024-06-03,10,EXTRA\n', b'\xff\xfeX\x00'])
def test_csv_schema_and_encoding_errors(bundle, body):
    (bundle / 'prices.csv').write_bytes(body)
    with pytest.raises(ImportValidationError):
        load_import_bundle(bundle, CUTOFF)


def test_duplicate_json_keys_rejected(bundle):
    body = (bundle / 'source.json').read_text(encoding='utf8')
    (bundle / 'source.json').write_text(body[:-1] + ',"currency":"BRL"}', encoding='utf8')
    with pytest.raises(ImportValidationError, match='duplicate_key'):
        load_import_bundle(bundle, CUTOFF)


def test_normalized_rows_sorted_by_identity_and_date(bundle):
    rows = records()['prices']
    rows.insert(0, {**rows[0], 'date': '2024-12-30', 'close': '1'})
    write_table(bundle, 'prices', rows)
    assert [row['date'] for row in load_import_bundle(bundle, CUTOFF)['prices']] == ['2024-06-03', '2024-12-30']


def test_freeze_hash_addressed_idempotent_and_input_unchanged(bundle, tmp_path):
    before = hashes(bundle)
    frozen = freeze_import_bundle(bundle, CUTOFF, 'output', workspace=tmp_path)
    assert frozen.parent == tmp_path / 'output'
    assert frozen.name == hashlib.sha256((frozen / 'normalized.json').read_bytes()).hexdigest()
    assert json.loads((frozen / 'normalized.json').read_bytes())['provenance'] == json.loads((frozen / 'provenance.json').read_bytes())
    assert freeze_import_bundle(bundle, CUTOFF, 'output', workspace=tmp_path) == frozen
    assert hashes(bundle) == before


def test_freeze_rejects_tampered_existing_content(bundle, tmp_path):
    frozen = freeze_import_bundle(bundle, CUTOFF, 'output', workspace=tmp_path)
    normalized = frozen / 'normalized.json'
    normalized.chmod(0o666)
    normalized.write_text('{}', encoding='utf8')
    with pytest.raises(ImportValidationError, match='existing_frozen_content_mismatch'):
        freeze_import_bundle(bundle, CUTOFF, 'output', workspace=tmp_path)


@pytest.mark.parametrize('output', ['../escape', 'input/output'])
def test_output_cannot_escape_workspace_or_modify_input(bundle, tmp_path, output):
    before = hashes(bundle)
    with pytest.raises(ImportValidationError, match='outside_explicit_workspace|outside_input_bundle'):
        freeze_import_bundle(bundle, CUTOFF, output, workspace=tmp_path)
    assert hashes(bundle) == before


def test_linked_input_cannot_read_outside_bundle(bundle, tmp_path):
    target = tmp_path / 'outside.csv'
    (bundle / 'prices.csv').replace(target)
    try:
        (bundle / 'prices.csv').symlink_to(target)
    except OSError:
        pytest.skip('OS does not permit symlink creation')
    with pytest.raises(ImportValidationError, match='linked_input_not_allowed'):
        load_import_bundle(bundle, CUTOFF)


@pytest.mark.parametrize('format', ['csv', 'xlsx'])
def test_cli_empty_templates_have_no_fabricated_assets(tmp_path, capsys, format):
    target = tmp_path / ('templates-' + format)
    assert main(['templates', str(target), '--format', format]) == 0
    assert json.loads(capsys.readouterr().out)['status'] == 'templates_created'
    source = json.loads((target / 'source.json').read_bytes())
    assert source['tickers'] == {}
    assert source['source_name'] == source['source_url'] == ''
    assert 'Tipo Titulo;Vencimento do Titulo;Data Resgate;PU;Quantidade;Valor' in (target / 'README.md').read_text(encoding='utf8')
    if format == 'csv':
        for table, headers in COLUMNS.items():
            assert (target / (table + '.csv')).read_bytes() == (','.join(headers) + '\n').encode('ascii')
    else:
        book = openpyxl.load_workbook(target / 'tables.xlsx', read_only=True)
        assert set(book.sheetnames) == set(COLUMNS)
        assert all(list(book[table].values) == [tuple(headers)] for table, headers in COLUMNS.items())
        book.close()
    with pytest.raises(ImportValidationError, match='already_exists'):
        emit_templates(target)


def test_cli_validate_and_redacted_errors(bundle, tmp_path, capsys):
    args = ['validate', str(bundle), '--cutoff', CUTOFF, '--workspace', str(tmp_path), '--output', 'frozen']
    assert main(args) == 0
    output = json.loads(capsys.readouterr().out)
    assert output['status'] == 'validated'
    assert len(output['bundle_sha256']) == 64
    assert not output['completeness_certified']
    assert output['row_counts']['payments'] == 1
    source = manifest()
    source['source_url'] = 'https://user:VERY_SECRET@example.org/?api_key=HIDDEN_TOKEN'
    write_manifest(bundle, source)
    assert main(args) == 2
    logged = capsys.readouterr()
    assert 'VERY_SECRET' not in logged.out + logged.err
    assert 'HIDDEN_TOKEN' not in logged.out + logged.err
    assert json.loads(logged.out)['status'] == 'invalid'


def test_cli_reports_missing_payment_table(bundle, tmp_path, capsys):
    (bundle / 'payments.csv').unlink()
    assert main(['validate', str(bundle), '--cutoff', CUTOFF, '--workspace', str(tmp_path), '--output', 'partial']) == 0
    output = json.loads(capsys.readouterr().out)
    assert output['missing_tables'] == ['payments']
    assert not output['completeness_certified']
