"""Evidence-only recovery of the missing_sources_b manifest partition.

This module never changes an inventory, snapshot, admission gate or model.
Source text and semantic table locators survive parsing; missing dates remain
missing, and an issuer history is not proof of five-year completeness.
"""
import argparse
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
import math
from pathlib import Path
import re
import unicodedata
from urllib.parse import parse_qs, quote, urljoin, urlsplit

from lxml import html

from .bulk_pending_sources import PublicSourceCollector, approved


SCOPE = 'missing-b'
GROUP = 'missing_sources_b'
SEEDS = {
    24953: 'https://ri.estapar.com.br/pt/quem-somos/composicao-acionaria',
    26069: 'https://ri.armac.com.br/informacoes-financeiras/comunicados-e-fatos-relevantes/',
    25372: 'https://ri.assai.com.br/',
    1562: 'https://www.baumer.com.br/',
    25291: 'https://ri.bravaenergia.com/informacoes-aos-acionistas/proventos/',
    2100: 'https://ri.cambuci.com.br/',
    25984: 'https://www.cba.com.br/ri',
    5380: 'https://www.altona.com.br/ri/br/paginas/avisos-aos-acionistas/',
    18309: 'https://www.celpa.com.br/',
    25445: 'https://ri.espacolaser.com.br/',
    3069: 'https://www.ferbasa.com.br/informacoes-aos-investidores/historico-de-pagamentos-de-proventos/',
    6041: 'https://www.itau.com.br/relacoes-com-investidores/informacoes-financeiras/investimentos-bemge',
    25712: 'https://www.gpssa.com.br/',
    21393: 'https://ri.celgpar.com/',
    20877: 'https://ri.helbor.com.br/informacoes-financeiras/dividendos/',
    25496: 'https://ri.jalles.com/',
    25690: 'https://ri.materdei.com.br/servicos-aos-investidores/perguntas-frequentes/',
    25119: 'https://ri.melnick.com.br/dividendos',
    26034: 'https://ri.multilaser.com.br/pt/dividendos',
    13765: 'https://www.minupar.com.br/publicacoes/aviso-aos-acionistas',
    24902: 'https://ri.mitrerealty.com.br/',
    25534: 'https://ri.oceanpact.com/',
    18708: 'https://www.aliancaparticipacoes.com.br/',
    22608: 'https://ri.paguemenos.com.br/',
    25070: 'https://ri.planoeplano.com.br/',
    24821: 'https://ri.rededorsaoluiz.com.br/informacoes-financeiras/devivendo_old/',
    25780: 'https://ri.petroreconcavo.com.br/dividendos',
    4669: 'https://www.riachuelo.com.br/ri',
    24694: 'https://ri.gruposbf.com.br/',
    24260: 'https://investor.smartfit.com.br/dividendos-jcp/',
    10880: 'https://www.sondotecnica.com.br/sobre-nos/relacao-com-investidores/',
    25208: 'https://www.tfco.com.br/ri/',
    25518: 'https://ri.westwing.com.br/servicos-aos-investidores/fale-com-ri/',
}


def normalized(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', value)
                   if not unicodedata.combining(c)).lower()


def text(node):
    return ' '.join(node.text_content().split())


def native_brl(value):
    """Parse a published Brazilian number without rounding or scaling."""
    value = value.strip().replace('R$', '').replace('\xa0', '').strip()
    if not re.fullmatch(r'(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?', value):
        raise ValueError('ambiguous_or_missing_native_brl')
    try:
        return Decimal(value.replace('.', '').replace(',', '.'))
    except InvalidOperation as error:
        raise ValueError('invalid_native_brl') from error


def exact_date(value):
    if not re.fullmatch(r'\d{2}/\d{2}/\d{4}', value):
        raise ValueError('not_an_exact_published_date')
    return datetime.strptime(value, '%d/%m/%Y').date().isoformat()


def schedule_date(value):
    if re.fullmatch(r'\d{2}/\d{2}/2[0-6]', value):
        value = value[:6] + '20' + value[6:]
    return exact_date(value)


@dataclass(frozen=True)
class CashEvidence:
    ticker: str
    cvm: int
    share_class: str
    isin: str | None
    event_type: str
    amount_per_share: str
    amount_basis: str
    payment_date: str | None
    ex_date: str | None
    declaration_date: str | None
    source_url: str
    source_file: str
    source_sha256: str
    locator: str
    native_cells: tuple
    published_at: str | None = None
    model_admission: bool = False
    bank_settlement_verified: bool = False


def document_tree(body):
    try:
        decoded = body.decode('utf-8-sig')
    except UnicodeDecodeError:
        decoded = body.decode('latin1')
    return html.fromstring(decoded, parser=html.HTMLParser(no_network=True))


def table_grid(table):
    """Expand only explicitly declared HTML spans, retaining row locators."""
    occupied = {}
    result = []
    for row_index, row in enumerate(table.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')):
        cells = []
        column = 0
        for node in row.xpath('./th|./td'):
            while (row_index, column) in occupied:
                cells.append(occupied[row_index, column])
                column += 1
            value = text(node)
            across, down = int(node.get('colspan', '1')), int(node.get('rowspan', '1'))
            if not 1 <= across <= 30 or not 1 <= down <= 100:
                raise ValueError('unsupported_issuer_table_span')
            for offset in range(across):
                cells.append(value)
                for later in range(1, down):
                    occupied[row_index + later, column + offset] = value
            column += across
        while (row_index, column) in occupied:
            cells.append(occupied[row_index, column])
            column += 1
        result.append((cells, row.getroottree().getpath(row)))
    return result


def inspect_html(body, receipt):
    tree = document_tree(body)
    document_base = urljoin(receipt['url'], tree.xpath('string(//base/@href)') or receipt['url'])
    tables = []
    for index, table in enumerate(tree.xpath('//table'), 1):
        grid = table_grid(table)
        rows = [cells for cells, _ in grid]
        if rows:
            tables.append({'locator': tree.getroottree().getpath(table), 'rows': rows,
                           'row_locators': [locator for _, locator in grid],
                           'heading': table.xpath('preceding::*[self::h1 or self::h2 or self::h3][1]/text()')})
    links = []
    for tag, attribute in (('a', 'href'), ('iframe', 'src'), ('script', 'src')):
        for index, node in enumerate(tree.xpath(f'//{tag}[@{attribute}]'), 1):
            url = quote(urljoin(document_base, node.get(attribute)), safe=':/?=&%+;,@[]!$()*#')
            if approved(url):
                links.append({'url': url, 'text': text(node)[:220],
                              'locator': f'//{tag}[@{attribute}][{index}]', 'tag': tag})
    config = []
    for index, node in enumerate(tree.xpath('//script[not(@src)]'), 1):
        value = node.text or ''
        if re.search(r'(?i)stockInfoId|id_stockinfo|stockinfo-viewer|enfoque|mz-stock|mzStock|template:', value):
            config.append({'locator': f'//script[not(@src)][{index}]',
                           'text': value[:20000]})
    for node in tree.xpath('//script|//style|//nav|//footer'):
        node.drop_tree()
    content = text(tree)
    snippets = [content[max(0, m.start()-160):m.end()+650] for m in re.finditer(
        r'(?i)dividend|provento|capital pr[oó]prio|desdobramento|grupamento|cota[çc][aã]o', content)]
    return {'title': tree.xpath('string(//title)'), 'tables': tables, 'links': links,
            'quote_config_citations': config,
            'years_present': sorted(set(re.findall(r'\b202[1-6]\b', content))),
            'semantic_snippets': snippets[:35],
            'inline_quote_markers': sorted(set(re.findall(
                r'(?i)(?:https:)?//[^\s\x22\x27<>]*(?:cotac|quote|stock|chart|mziq)[^\s\x22\x27<>]*',
                body.decode('utf-8', errors='replace'))))[:35]}


def parse_supported_table(body, receipt, identity):
    """Known issuer layouts only; no class/suffix or predecessor inference."""
    url = receipt['url']
    host = urlsplit(url).hostname
    paths = {
        'ri.melnick.com.br': '/dividendos',
        'ri.helbor.com.br': '/informacoes-financeiras/dividendos',
        'ri.multilaser.com.br': '/pt/dividendos',
        'ri.bravaenergia.com': '/informacoes-aos-acionistas/proventos',
        'ri.armac.com.br': '/informacoes-financeiras/dividendos-e-jcp-2',
        'ri.mitrerealty.com.br': '/informacoes-aos-investidores/dividendos-2',
        'ri.paguemenos.com.br': '/informacoes-aos-investidores/remuneracao-ao-acionistas',
        'ri.petroreconcavo.com.br': '/dividendos',
        'ri.rededorsaoluiz.com.br': '/informacoes-financeiras/devivendo_old',
        'investor.smartfit.com.br': '/dividendos-jcp',
        'www.tfco.com.br': '/dividendos-e-jcp',
        'www.ferbasa.com.br': '/informacoes-aos-investidores/historico-de-pagamentos-de-proventos',
    }
    if paths.get(host) != urlsplit(url).path.rstrip('/'):
        return []
    layouts = {
        'ri.melnick.com.br': (25119, 'ON', 4, 3, 2, None, 1),
        'ri.helbor.com.br': (20877, 'ON', 1, 3, 2, None, None),
        'ri.multilaser.com.br': (26034, 'ON', 3, 0, None, None, None),
        'ri.bravaenergia.com': (25291, 'ON', 6, 4, 3, 2, 0),
        'ri.armac.com.br': (26069, 'ON', 5, 0, 3, None, 4),
        'ri.mitrerealty.com.br': (24902, 'ON', 5, 3, None, None, None),
        'ri.paguemenos.com.br': (22608, 'ON', 5, 4, 3, 2, 0),
        'ri.petroreconcavo.com.br': (25780, 'ON', 4, 3, None, None, 1),
        'ri.rededorsaoluiz.com.br': (24821, 'ON', 7, 4, None, 3, 1),
        'investor.smartfit.com.br': (24260, 'ON', 5, 2, None, None, 0),
        'www.tfco.com.br': (25208, 'PN', 4, 2, None, None, 1),
    }
    if host == 'www.ferbasa.com.br':
        if identity['share_class'] not in ('ON', 'PN'):
            raise ValueError('ferbasa_exact_share_class_not_supported')
        layouts[host] = (3069, identity['share_class'],
                         5 if identity['share_class'] == 'ON' else 6, 3, None, 2, 1)
    if host not in layouts:
        return []
    cvm, share_class, amount_col, payment_col, ex_col, declaration_col, kind_col = layouts[host]
    if identity['cvm'] != cvm or identity['share_class'] != share_class:
        raise ValueError('exact_manifest_identity_disagrees_with_typed_adapter')
    result = []
    for table in inspect_html(body, receipt)['tables']:
        heading = normalized(' '.join(table['heading']))
        if cvm == 25291 and 'enauta' in heading:
            continue
        headers = normalized(' '.join(' '.join(row) for row in table['rows'][:2]))
        if not ('pagamento' in headers and ('acao' in headers or cvm == 25780
                                           or (cvm == 25208 and 'bruto por pn' in headers))):
            continue
        for number, cells in enumerate(table['rows'], 1):
            if len(cells) <= max(amount_col, payment_col, ex_col or 0, declaration_col or 0):
                continue
            try:
                amount = native_brl(cells[amount_col])
                native_payment = cells[payment_col]
                if native_payment == 'A ser definida' or re.fullmatch(r'Dez/2[6-8]', native_payment):
                    payment = None
                else:
                    payment = schedule_date(native_payment)
                ex = schedule_date(cells[ex_col]) if ex_col is not None else None
                declaration = schedule_date(cells[declaration_col]) if declaration_col is not None else None
            except ValueError:
                continue
            dated = payment or ex or declaration
            if dated is not None and not '2021-01-01' <= dated <= '2026-12-31':
                continue
            kind = cells[kind_col] if kind_col is not None else 'Dividendos'
            if 'dividendo' in normalized(kind):
                event_type = 'dividend'
            elif normalized(kind) in ('jcp', 'jscp') or 'juros sobre' in normalized(kind):
                event_type = 'interest_on_equity'
            else:
                continue
            if amount <= 0:
                continue
            basis = 'gross' if 'bruto' in headers else 'as_reported_gross_net_unspecified'
            if cvm == 22608 and 'valor unitario liquido' in headers:
                basis = 'gross_distinct_from_separately_reported_net'
            result.append(CashEvidence(
                ticker=identity['ticker'], cvm=cvm, share_class=share_class, isin=identity['isin'],
                event_type=event_type, amount_per_share=str(amount), amount_basis=basis,
                payment_date=payment, ex_date=ex, declaration_date=declaration,
                source_url=url, source_file=receipt['file'], source_sha256=receipt['sha256'],
                locator=table['row_locators'][number-1], native_cells=tuple(cells)))
    return result


def parse_altona_declarations(body, receipt, identity):
    """Preserve proposed values and unequal installments without allocating."""
    if receipt['url'] != 'https://www.altona.com.br/ri/br/dividendos/' or identity['cvm'] != 5380:
        return []
    source_class = {'ON': 'Ordinarias', 'PNA': 'Preferenciais'}.get(identity['share_class'])
    if source_class is None:
        raise ValueError('altona_manifest_class_not_supported')
    tree = document_tree(body)
    result = []
    for table in tree.xpath('//div[@class="tabela"]'):
        headers = [text(node) for node in table.xpath('./div[@class="menu-tabela"]/ul/li')]
        if not any(normalized(value) == 'valor proposto' for value in headers):
            continue
        native = [([text(node) for node in row.xpath('./li')], row)
                  for row in table.xpath('./div[@class="itens-tabela"]/ul')]
        dates_text = next((cells[0] for cells, _ in native if cells and cells[0].startswith('Aprovado AGOE ')), '')
        dates = re.findall(r'\d{2}/\d{2}/\d{4}', dates_text)
        if len(dates) < 5 or '1/3 por pagamento' not in dates_text:
            continue
        for cells, row in native:
            if len(cells) != 4 or normalized(cells[0]) != normalized(source_class):
                continue
            result.append({'identity': {k: identity[k] for k in ('ticker', 'cvm', 'share_class', 'isin')},
                'source_class_label': cells[0], 'source_headers': headers, 'native_cells': cells,
                'native_proposed_value': str(native_brl(cells[2])),
                'proposed_value_unit_label': 'Valor Proposto', 'per_share_unit_not_explicit_in_header': True,
                'aggregate_amount': str(native_brl(cells[3])),
                'approval_date': exact_date(dates[0]), 'ex_date': exact_date(dates[4]),
                'installments': [{'payment_date': exact_date(date), 'amount_per_share': None}
                                 for date in dates[1:4]],
                'equal_installment_allocation_explicitly_disallowed_by_source': True,
                'source_class_matches_registered_class': identity['share_class'] == 'ON',
                'preferred_PNA_binding_not_inferred': identity['share_class'] == 'PNA',
                'source_file': receipt['file'], 'source_url': receipt['url'], 'sha256': receipt['sha256'],
                'locator': row.getroottree().getpath(row), 'native_schedule_note': dates_text,
                'model_admission': False, 'published_at': None})
    return result


def immutable_json(path, value):
    body = (json.dumps(value, ensure_ascii=True, indent=2, allow_nan=False) + '\n').encode()
    if path.exists():
        if path.read_bytes() != body:
            raise ValueError('immutable_missing_b_report_conflict')
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open('xb') as stream:
            stream.write(body)


def collect(root, manifest, cvms=None, follow=False):
    manifest_path = Path(manifest)
    data = json.loads(manifest_path.read_bytes())
    records = [r for r in data['records'] if r['task_group'] == GROUP]
    if len(records) != 41 or len({r['ticker'] for r in records}) != 41:
        raise ValueError('missing_b_manifest_partition_changed')
    collector = PublicSourceCollector(root, SCOPE)
    issuers = {}
    for record in records:
        issuers.setdefault(record['cvm'], []).append(record)
    for cvm, identities in issuers.items():
        if cvms and cvm not in cvms:
            continue
        if cvm not in SEEDS:
            print(json.dumps({'cvm': cvm, 'status': 'origin_explicitly_denied_no_network'}), flush=True)
            continue
        seed = SEEDS[cvm]
        try:
            receipt = collector.fetch(seed, discovery={'kind': 'FCA_website_or_web_search',
                'source_url': identities[0]['website_declared_in_FCA'],
                'issuer_cvm': cvm, 'publication_time_not_inferred': True})
            print(json.dumps({'cvm': cvm, 'url': seed, 'status': receipt['status'],
                              'error': receipt['error'], 'location': receipt['location']}), flush=True)
            if receipt['status'] != 200 or receipt['error']:
                continue
            inspection = inspect_html(collector.body(receipt), receipt)
            dest = collector.base / ('inspection-' + hashlib.sha256(seed.encode()).hexdigest() + '.json')
            immutable_json(dest, {'receipt': receipt, 'inspection': inspection})
            if follow:
                chosen = []
                for pattern in (r'dividendo|provento|remuneracao', r'cotac|quotes|grafico'):
                    candidates = [l for l in inspection['links'] if l['tag'] == 'a'
                                  and re.search(pattern, normalized(l['text']+' '+l['url']))
                                  and l['url'] != seed]
                    if candidates:
                        chosen.append(candidates[0])
                for link in chosen[:2]:
                    try:
                        follow_receipt = collector.fetch(link['url'], discovery={
                            'kind': 'issuer_published_link', 'source_url': seed,
                            'locator': link['locator'], 'link_text': link['text'], 'issuer_cvm': cvm})
                        print(json.dumps({'cvm': cvm, 'url': link['url'], 'status': follow_receipt['status'],
                                          'error': follow_receipt['error'], 'location': follow_receipt['location']}), flush=True)
                        if follow_receipt['status'] == 200 and not follow_receipt['error']:
                            detail = inspect_html(collector.body(follow_receipt), follow_receipt)
                            immutable_json(collector.base / ('inspection-' + hashlib.sha256(link['url'].encode()).hexdigest() + '.json'),
                                           {'receipt': follow_receipt, 'inspection': detail})
                    except ValueError as error:
                        immutable_json(collector.base / ('skip-' + hashlib.sha256(link['url'].encode()).hexdigest() + '.json'),
                                       {'url': link['url'], 'cvm': cvm, 'error': str(error), 'network_attempted': False})
        except ValueError as error:
            print(json.dumps({'cvm': cvm, 'url': seed, 'error': str(error)}), flush=True)
            immutable_json(collector.base / ('skip-' + hashlib.sha256(seed.encode()).hexdigest() + '.json'),
                           {'url': seed, 'cvm': cvm, 'error': str(error), 'network_attempted': False})


def collect_plan(root, plan):
    collector = PublicSourceCollector(root, SCOPE)
    for entry in json.loads(Path(plan).read_bytes()):
        url = entry['url']
        try:
            receipt = collector.fetch(url, discovery={k: v for k, v in entry.items() if k != 'url'})
            print(json.dumps({k: receipt.get(k) for k in ('url', 'status', 'error', 'location')}), flush=True)
            if receipt['status'] == 200 and not receipt['error']:
                body = collector.body(receipt)
                if body.lstrip().startswith((b'<!', b'<html', b'<HTML', b'<head', b'\xef\xbb\xbf<')):
                    inspection = inspect_html(body, receipt)
                    immutable_json(collector.base / ('inspection-v2-' + hashlib.sha256(url.encode()).hexdigest() + '.json'),
                                   {'receipt': receipt, 'inspection': inspection})
                    if re.search(r'(?i)pardon our interruption|just a moment|access denied|captcha', inspection['title']):
                        collector.prior_stops.add(urlsplit(url).hostname)
                        immutable_json(collector.base / ('stop-' + hashlib.sha256(url.encode()).hexdigest() + '.json'),
                                       {'url': url, 'error': 'public_content_challenge_stop',
                                        'source_sha256': receipt['sha256'], 'cvm': entry.get('issuer_cvm')})
        except ValueError as error:
            print(json.dumps({'url': url, 'error': str(error)}), flush=True)
            immutable_json(collector.base / ('skip-' + hashlib.sha256(url.encode()).hexdigest() + '.json'),
                           {'url': url, 'cvm': entry.get('issuer_cvm'), 'error': str(error), 'network_attempted': False})


def widget_request(page_body, page_receipt, script_body, identity):
    """Bind a public history template to literal, issuer-published parameters."""
    script_url = 'https://cdn.mziq.com/stockinfo-viewer/main/latest/js/main.js'
    parsed = inspect_html(page_body, page_receipt)
    if not any(l['url'] == script_url and l['tag'] == 'script' for l in parsed['links']):
        raise ValueError('issuer_did_not_publish_this_widget_script')
    script = script_body.decode('utf-8')
    if not all(marker in script for marker in (
            'https://api.mziq.com/mzstockinfo', '/history?from=', '&to=',
            '&tickers=', '&adjusted=', 'YYYY-MM-DD')):
        raise ValueError('public_widget_history_template_not_demonstrated')
    configs = '\n'.join(c['text'] for c in parsed['quote_config_citations'])
    ticker_lists = re.findall(r'''tickers\s*:\s*["']([^"']+)["']''', configs)
    if identity['ticker'] not in {t.strip() for group in ticker_lists for t in group.split(',')}:
        raise ValueError('exact_ticker_not_literal_in_issuer_widget_config')
    match = re.search(r'''(?:stockInfoId\s*=|stockinfoId\s*:)\s*["']([a-f0-9-]{36})["']''', configs)
    if not match:
        raise ValueError('issuer_widget_identifier_not_published')
    url = (f'https://api.mziq.com/mzstockinfo/{match[1]}/history'
           f'?from=2021-09-01&to=2026-09-30&tickers={identity["ticker"]}&adjusted=false')
    return url, {'kind': 'issuer_published_widget_template', 'source_url': page_receipt['url'],
                 'source_sha256': page_receipt['sha256'], 'script_url': script_url,
                 'script_sha256': hashlib.sha256(script_body).hexdigest(),
                 'issuer_cvm': identity['cvm'], 'ticker': identity['ticker'],
                 'identity': {k: identity[k] for k in ('ticker', 'cvm', 'share_class', 'isin')},
                 'adjusted_parameter': False, 'template_discovered_in_actual_script': True,
                 'historical_share_basis_reconciliation_required': True}


def collect_widgets(root, manifest):
    root = Path(root)
    base = root / 'raw/primary-research/all-pending' / SCOPE
    rows = [r for r in json.loads(Path(manifest).read_bytes())['records'] if r['task_group'] == GROUP]
    receipts = [json.loads(f.read_bytes()) for f in base.glob('response-*.json')]
    script = next(r for r in receipts if r['url'] == 'https://cdn.mziq.com/stockinfo-viewer/main/latest/js/main.js')
    collector = PublicSourceCollector(root, SCOPE)
    script_body = collector.body(script)
    seen = set()
    for page in receipts:
        if (page['status'] != 200 or page['error'] or not re.search(r'cotac|grafico', page['url'])):
            continue
        for row in rows:
            try:
                url, discovery = widget_request(collector.body(page), page, script_body, row)
            except ValueError:
                continue
            if url in seen:
                continue
            seen.add(url)
            try:
                result = collector.fetch(url, discovery=discovery)
                print(json.dumps({'ticker': row['ticker'], **{k: result[k] for k in ('status', 'error', 'bytes')}}), flush=True)
            except ValueError as error:
                print(json.dumps({'ticker': row['ticker'], 'error': str(error)}), flush=True)
                immutable_json(base / ('skip-' + hashlib.sha256(url.encode()).hexdigest() + '.json'),
                               {'url': url, 'cvm': row['cvm'], 'error': str(error), 'network_attempted': False})


def parse_widget_history(body, receipt, identity):
    """Return actual last observations for 61 months, without rescaling prices."""
    query = parse_qs(urlsplit(receipt['url']).query)
    if query.get('adjusted') != ['false'] or query.get('tickers') != [identity['ticker']]:
        raise ValueError('widget_request_identity_or_adjustment_disagrees')
    payload = json.loads(body)
    errors = payload.get('data', {}).get('errors', [])
    if (payload.get('success') is not True or not isinstance(errors, list)
            or any(not isinstance(error, dict) or error.get('ticker') == identity['ticker']
                   or error.get('status') != 'Invalid index / not configured' for error in errors)):
        raise ValueError('widget_did_not_return_success')
    rows = payload['data'].get('tickers')
    if not isinstance(rows, list) or len(rows) != 1 or rows[0].get('ticker') != identity['ticker']:
        raise ValueError('widget_response_exact_ticker_disagrees')
    if payload['data'].get('source') != 'Enfoque':
        raise ValueError('widget_underlying_provider_not_demonstrated')
    monthly, dates = {}, set()
    for index, row in enumerate(rows[0]['history']):
        observed = row.get('date', '')
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', observed):
            raise ValueError('widget_observation_date_not_native_ISO_date')
        datetime.strptime(observed, '%Y-%m-%d')
        close = row.get('closed')
        if type(close) not in (int, float) or not math.isfinite(close) or close <= 0:
            raise ValueError('widget_observation_price_not_positive_native_number')
        if observed in dates:
            raise ValueError('widget_duplicate_native_date')
        dates.add(observed)
        if not '2021-09-01' <= observed <= '2026-09-30':
            raise ValueError('widget_response_outside_requested_window')
        month = observed[:7]
        if month not in monthly or monthly[month]['date'] < observed:
            monthly[month] = {'month': month, 'date': observed, 'close': close,
                              'locator': f'$.data.tickers[0].history[{index}].closed'}
    expected = [f'{year:04d}-{month:02d}' for year in range(2021, 2027) for month in range(1, 13)
                if '2021-09' <= f'{year:04d}-{month:02d}' <= '2026-09']
    missing = [month for month in expected if month not in monthly]
    return {'identity': {k: identity[k] for k in ('ticker', 'cvm', 'share_class', 'isin')},
            'provider': 'Enfoque_via_issuer_published_MZ_widget', 'adjusted_parameter': False,
            'unrelated_index_configuration_errors_retained': errors,
            'daily_observations': len(dates), 'monthly_observations': len(monthly),
            'required_monthly_observations': 61, 'missing_months': missing,
            'has_61_consecutive_months': not missing and len(monthly) == 61,
            'monthly_prices': [monthly[m] for m in sorted(monthly)],
            'source_url': receipt['url'], 'source_file': receipt['file'], 'sha256': receipt['sha256'],
            'retrieved_at': receipt.get('retrieved_at'), 'published_at': None,
            'numeric_values_changed': False, 'price_actions_reconciled': False,
            'model_admission': False}


def verified_cached_body(root, receipt):
    root = Path(root).resolve()
    path = (root / receipt['file']).resolve()
    if not path.is_relative_to(root) or not path.is_file() or path.is_symlink():
        raise ValueError('unsafe_missing_b_cached_file')
    body = path.read_bytes()
    if hashlib.sha256(body).hexdigest() != receipt['sha256'] or len(body) != receipt.get('bytes', len(body)):
        raise ValueError('missing_b_cached_original_digest_changed')
    return body


def inspect_binary(body):
    import io
    if body.startswith(b'%PDF'):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(body))
        return {'format': 'PDF', 'pages': len(reader.pages), 'original_unchanged': True,
                'page_citations': [{'locator': f'PDF page {index+1}', 'text': page.extract_text()}
                                   for index, page in enumerate(reader.pages[:2])]}
    if body.startswith(b'PK'):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
        sheets = []
        for sheet in workbook:
            found = []
            for row in sheet.iter_rows(max_row=min(sheet.max_row, 500), max_col=min(sheet.max_column, 100)):
                if any(isinstance(c.value, str) and re.search(r'(?i)dividend|capital pr|\bjcp\b|jscp|cotac', c.value) for c in row):
                    found.append({'row': row[0].row, 'cells': [{'cell': c.coordinate, 'value': c.value}
                        for c in row if c.value is not None]})
            sheets.append({'sheet': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
                           'semantic_rows': found})
        workbook.close()
        return {'format': 'XLSX', 'sheets': sheets, 'original_unchanged': True,
                'no_per_share_cash_or_payment_dates_derived_from_account_balances': True}
    return None


ISSUER_FINDINGS = {
    24953: 'Issuer Next/SPA HTML acquired; published Location removes /pt while canonical location returns to /pt. No native dividend schedule in returned HTML; no redirect loop followed.',
    26069: 'Official dividend/JSCP table has 22 native events from 2021 through 2026. Three JSCP rows explicitly say A ser definida. Widget literals identify ARML3; unrelated SMAL11 index error retained separately.',
    25372: 'Official RI returned HTTP 200 with title Pardon Our Interruption and no public navigation. Content inspected as a challenge; no retry.',
    21610: 'B3SA3 issuer is itself a prohibited B3 origin. No network request was made; manifest CVM 21610 and null ISIN remain unchanged.',
    1562: 'Official FCA homepage and its published Nuxt state bootstrap acquired; homepage is a client shell and state payload contains no dividend calendar. No class values inferred for ON or PN.',
    25291: 'Official table separates Brava Energia from Enauta. Adapter extracts only the two Brava rows; all Enauta predecessor amounts are retained in original but excluded from BRAV3.',
    2100: 'Official 16-sheet financial-history XLSX acquired and inspected. DFC - CASH FLOW rows 33, 51 and 53 and BP - BS row 37 are company cash/account balances, not a per-share payment schedule. No EPS or per-share dividend derived.',
    25984: 'www.cba.com.br robots returned 403 and was stopped. Separate search-discovered official ri.cba.com.br FAQ acquired: states 25% minimum distribution policy, which does not establish individual payment dates.',
    5380: 'Official dividend HTML has three proposed-value schedules separately labeled Ordinarias and Preferenciais, approval 29/04/2026, ex 30/04/2026, payments 12/06, 28/08, 30/10/2026. Source explicitly says installments are not 1/3 each. No per-installment amount allocated; PNA binding not inferred from generic Preferenciais. Linked August PDF is robots-disallowed and was not requested.',
    17329: 'www.engie.com.br explicitly forbidden and previously stopped. No new network request. Manifest ON and exact ISIN retained.',
    18309: 'FCA-declared www.celpa.com.br robots failed with URLError before page access. No retry; EQPA3 ON and EQPA5 PNA remain distinct.',
    25445: 'Official ri.espacolaser.com.br robots failed with URLError before page access. No retry and no alternative cash/payment dates invented.',
    3069: 'Official table explicitly separates ON and PN amounts. HTML rowspans expanded only as published. ON source cell 0.20422 and multi-date installment rows are not silently corrected or divided; originals retained.',
    6041: 'FCA-declared Itau RI origin robots returned 403. Resource not requested. No Bemge dividend values substituted from Itau Unibanco.',
    25712: 'FCA homepage robots failed with URLError. Separate official ri.gpssa.com.br quote page and literal GGPS3 widget acquired, with 61 native monthly observations; dividend calendar not present.',
    21393: 'Official CELGPAR RI homepage and published dividend-policy page acquired; tables are site layout and do not provide an event calendar. Null manifest ISIN retained.',
    20877: 'Official ON dividend table provides five native payment rows through April 2025; 61-month native widget prices recovered. Source does not yet establish cash coverage through September 2026.',
    25496: 'Official Jalles robots returns 307 to /pt/robots.txt. Collector records redirect and cannot validate robots; no automatic follow or retry.',
    25690: 'Official Mater Dei FAQ and quote page acquired; published MATD3 widget supplies 61 native monthly observations. FAQ provides identity/policy, not an actual dividend calendar.',
    25119: 'Official Melnick history supplies 18 native events. Quote page publishes an Enfoque iframe with token query; it was inspected locally only and not requested under the credential-free collector policy.',
    26034: 'Official Multi dividend table contains two native payment rows in 2022 and January 2026. Alternate canonical path redirects back to /pt; no loop followed. No complete quote history recovered.',
    13765: 'Official notice PDF visually inspected: dividend payment 12/06/2026, holder date 14/04/2026 and ex date 15/04/2026. Document dated 18/05/2026; it omits per-share amount, so no cash amount is fabricated.',
    24902: 'Official Mitre table supplies 40 individual installment rows, excludes annual totals, and keeps Base de Pagamento distinct from ex-date. Literal MTRE3 widget supplies 61 native monthly observations.',
    25534: 'Official OceanPact homepage and quote page acquired. Literal OPCT3 widget supplies 61 native monthly observations. No dividend schedule appears in these acquired pages.',
    18708: 'Official shareholder-notices HTML and current 30/04/2026 AGM PDF acquired. ON/PN values require explicit class evidence; manifest null ISINs are preserved.',
    22608: 'Official Pague Menos schedule has four JSCP rows with separate gross and net unit values and explicit approval/ex/payment dates. Gross amounts preserved independently of net.',
    25070: 'Official Plano&Plano home and quote page acquired; literal PLPL3 widget supplies 61 native monthly observations. Homepage does not publish a native dividend schedule.',
    24821: 'Official Rede DOr Portuguese schedule provides 25 events with gross per-share values rounded as published. No extra precision invented. Literal RDOR3 widget provides 61 native monthly observations.',
    25780: 'Official PetroReconcavo table provides 13 events, including three installments labeled only Dez/26, Dez/27, Dez/28. No day assigned to those installments; Data Com not converted to ex-date.',
    4669: 'FCA-declared Riachuelo origin robots could not be validated. No retry and no GUAR3/RIAA3 predecessor or suffix mapping guessed.',
    24694: 'Official Grupo SBF home and quote page acquired; literal SBFG3 widget provides 61 native monthly observations. Dividend history not present in these acquired pages.',
    24260: 'Official SmartFit schedule provides 11 gross JCP rows. Deliberation month labels remain native cells; no declaration day inferred. Payment dates with two-digit years normalized within the explicit 2020-2026 table schema.',
    10880: 'Official RI explicitly identifies SOND5 as preferred class A and SOND6 as preferred class B. Published 2026 and December 2025 cash-notice links return 302 to index.php; redirect not followed.',
    25208: 'Official Track&Field table explicitly labels gross per PN, giving 22 native events. ON economics not substituted. Literal TFCO4 history request timed out; no retry.',
    25518: 'Official Westwing quote widget supplies 61 native monthly observations. RI notices listing acquired for cash/action review; no cash event or payout type inferred from a third-party yield.',
}


def audit(root, manifest, output):
    root = Path(root)
    manifest_body = Path(manifest).read_bytes()
    data = json.loads(manifest_body)
    rows = [r for r in data['records'] if r['task_group'] == GROUP]
    base = root / 'raw/primary-research/all-pending' / SCOPE
    sources = []
    for path in base.glob('response-*.json'):
        receipt = json.loads(path.read_bytes())
        body = verified_cached_body(root, receipt)
        source = {'receipt': receipt, 'metadata_file': path.relative_to(root).as_posix()}
        if receipt['status'] == 200 and not receipt['error'] and 'robots.txt' not in receipt['url']:
            if body.lstrip().startswith((b'<!', b'<html', b'<HTML', b'<head', b'\xef\xbb\xbf<')):
                source['inspection'] = inspect_html(body, receipt)
                if re.search(r'(?i)pardon our interruption|just a moment|access denied|captcha', source['inspection']['title']):
                    source['content_outcome'] = 'inspected_challenge_stop_no_retry'
                    immutable_json(base / ('content-stop-' + hashlib.sha256(receipt['url'].encode()).hexdigest() + '.json'),
                        {'url': receipt['url'], 'status': receipt['status'], 'error': 'public_content_challenge_stop',
                         'source_file': receipt['file'], 'sha256': receipt['sha256'],
                         'inspected_title': source['inspection']['title'], 'retries': 0})
                else:
                    source['content_outcome'] = 'original_acquired_and_inspected'
            else:
                source['content_outcome'] = 'non_html_original_acquired'
                binary = inspect_binary(body)
                if binary:
                    source['document_inspection'] = binary
        else:
            source['content_outcome'] = 'access_or_policy_outcome_preserved'
        sources.append(source)
    skips = [json.loads(f.read_bytes()) for f in base.glob('skip-*.json')]
    prior = []
    # Existing issuer research is read-only, including sources on stopped hosts.
    for path in (root / 'raw/primary-research').rglob('*receipt.json'):
        if 'test' in str(path).lower() or path.stat().st_size > 200000:
            continue
        try:
            receipt = json.loads(path.read_bytes())
            if not isinstance(receipt, dict) or not receipt.get('file') or not receipt.get('url'):
                continue
            host = urlsplit(receipt['url']).hostname
            matched = [r for r in rows if r['ticker'] == receipt.get('ticker')
                       or (r['cvm'] in SEEDS and host == urlsplit(SEEDS[r['cvm']]).hostname)]
            if not matched:
                continue
            body = verified_cached_body(root, receipt)
            item = {'receipt': receipt, 'metadata_file': path.relative_to(root).as_posix(),
                    'tickers': [r['ticker'] for r in matched], 'original_hash_verified': True,
                    'access': 'existing_local_cache_read_only', 'publication_time_inferred': False}
            if body.lstrip().startswith((b'<!', b'<html', b'<HTML')):
                item['inspection'] = inspect_html(body, receipt)
            prior.append(item)
        except (ValueError, KeyError, OSError):
            continue
    results = []
    for identity in rows:
        cvm, ticker = identity['cvm'], identity['ticker']
        matches = []
        for source in sources:
            receipt = source['receipt']
            host = urlsplit(receipt['url']).hostname
            discovery = receipt.get('discovery') or {}
            if (discovery.get('issuer_cvm') == cvm
                    or (cvm in SEEDS and host == urlsplit(SEEDS[cvm]).hostname)
                    or host == {25984: 'ri.cba.com.br', 25712: 'ri.gpssa.com.br'}.get(cvm)):
                matches.append(source)
        cash, prices, declarations, parsing_errors = [], [], [], []
        for source in matches:
            receipt = source['receipt']
            if receipt['status'] != 200 or receipt['error']:
                continue
            body = verified_cached_body(root, receipt)
            try:
                if 'inspection' in source and source.get('content_outcome') != 'inspected_challenge_stop_no_retry':
                    cash.extend(asdict(row) for row in parse_supported_table(body, receipt, identity))
                    declarations.extend(parse_altona_declarations(body, receipt, identity))
                if '/mzstockinfo/' in receipt['url'] and '/history?' in receipt['url']:
                    if (receipt.get('discovery') or {}).get('ticker') == ticker:
                        prices.append(parse_widget_history(body, receipt, identity))
            except ValueError as error:
                parsing_errors.append({'url': receipt['url'], 'error': str(error)})
        local = [item for item in prior if ticker in item['tickers']]
        outcomes = [s for s in skips if s.get('cvm') == cvm]
        baseline_file = identity.get('source_inventory_record', {}).get('identity_file')
        baseline_cache = {'file': baseline_file, 'access': 'existing_local_cache_read_only',
                          'retrieved_at': None, 'published_at': None,
                          'original_acquisition_receipt_not_reconstructed': True}
        if baseline_file:
            path = (root / baseline_file).resolve()
            if path.is_relative_to(root.resolve()) and path.is_file() and not path.is_symlink():
                original = path.read_bytes()
                cached_identity = json.loads(original)
                baseline_cache.update({'sha256': hashlib.sha256(original).hexdigest(), 'bytes': len(original),
                    'status': 'existing_cached_identity_bytes_inspected',
                    'json_shape': list(cached_identity) if isinstance(cached_identity, dict) else type(cached_identity).__name__,
                    'no_null_ISIN_completed_from_this_cache': True})
            else:
                baseline_cache['status'] = 'manifest_cache_path_not_available'
        else:
            baseline_cache['status'] = 'manifest_does_not_name_identity_cache'
        has_prices = any(p['has_61_consecutive_months'] for p in prices)
        if cvm in (21610, 17329):
            status = 'explicitly_forbidden_origin_not_requested'
        elif cash and has_prices:
            status = 'typed_cash_and_61_month_prices_recovered_for_integration_review'
        elif cash:
            status = 'typed_cash_recovered_independent_price_history_incomplete'
        elif any(s.get('inspection') and s.get('content_outcome') != 'inspected_challenge_stop_no_retry' for s in matches):
            status = 'issuer_original_inspected_requires_specific_document_or_adapter'
        else:
            status = 'specific_access_or_robots_outcome_no_admission'
        payment_dates = sorted(c['payment_date'] for c in cash if c['payment_date'])
        result = {k: identity[k] for k in ('ticker', 'cvm', 'share_class', 'isin', 'name', 'website_declared_in_FCA')}
        result.update({'outcome': status, 'original_pending_reasons': identity['reasons'],
            'specific_inspected_finding': ISSUER_FINDINGS[cvm],
            'source_files': [s['metadata_file'] for s in matches],
            'source_outcomes': [{k: s['receipt'].get(k) for k in ('url', 'status', 'error', 'file', 'sha256', 'retrieved_at', 'published_at', 'location')}
                                | {'content_outcome': s['content_outcome']} for s in matches],
            'skipped_before_resource_request': outcomes,
            'previous_cached_research_files_inspected': [p['metadata_file'] for p in local],
            'baseline_identity_cache_inspected': baseline_cache,
            'cash_events': cash, 'cash_event_count': len(cash),
            'native_declarations_with_unallocated_installments': declarations,
            'first_native_payment_date': payment_dates[0] if payment_dates else None,
            'last_native_payment_date': payment_dates[-1] if payment_dates else None,
            'cash_dates_unspecified': sum(c['payment_date'] is None for c in cash),
            'independent_prices': prices,
            'quote_config_citations': [{'source_file': s['receipt']['file'], 'source_url': s['receipt']['url'],
                'sha256': s['receipt']['sha256'], 'config': s['inspection']['quote_config_citations']}
                for s in matches if s.get('inspection', {}).get('quote_config_citations')],
            'document_citations': [{'source_file': s['receipt']['file'], 'source_url': s['receipt']['url'],
                'sha256': s['receipt']['sha256'], 'inspection': s['document_inspection']}
                for s in matches if s.get('document_inspection')],
            'parsing_errors': parsing_errors, 'model_admission': False,
            'remaining_validation': ['cash_calendar_coverage_and_native_precision_reconciliation',
                'same_CVM_class_ISIN_identity_across_events', 'native_prices_and_share_actions_reconciliation',
                'annual_financial_gate_unchanged', 'publication_date_and_point_in_time_availability_not_inferred']})
        results.append(result)
    candidates = sorted((r for r in results if r['cash_events'] and any(p['has_61_consecutive_months'] for p in r['independent_prices'])),
                        key=lambda r: ({'MTRE3': 0, 'ARML3': 1, 'RDOR3': 2}.get(r['ticker'], 3), -r['cash_event_count']))
    report = {'scope': SCOPE, 'manifest_sha256': hashlib.sha256(manifest_body).hexdigest(),
        'snapshot': data['snapshot'], 'reference': data['reference'], 'assigned_count': len(rows),
        'issuer_count': len({r['cvm'] for r in rows}), 'all_assigned_codes_accounted_for': len(results) == 41,
        'typed_cash_codes': [r['ticker'] for r in results if r['cash_events']],
        'codes_with_independent_61_month_prices': [r['ticker'] for r in results if any(p['has_61_consecutive_months'] for p in r['independent_prices'])],
        'integration_candidates': [{'ticker': r['ticker'], 'cvm': r['cvm'], 'share_class': r['share_class'],
            'isin': r['isin'], 'cash_events': r['cash_event_count'], 'unknown_payment_dates': r['cash_dates_unspecified'],
            'independent_months': 61, 'first_payment': r['first_native_payment_date'],
            'last_payment': r['last_native_payment_date'], 'model_admission': False} for r in candidates[:3]],
        'results': results, 'sources': sources, 'existing_research_inspected': prior,
        'no_B3_network': True, 'model_or_pointer_publication': False, 'gate_relaxation': False,
        'EPS_computed': False, 'prices_imputed': False, 'payment_dates_inferred': False,
        'retrieval_publication_dates_kept_distinct': True,
        'contracts': {'parse_supported_table': '(body: bytes, receipt: dict, identity: manifest_record) -> list[CashEvidence]; exact decimal strings, native cells, row XPath; gross/net unknown retained',
            'parse_altona_declarations': '(body: bytes, receipt: dict, identity: manifest_record) -> list[dict]; native source-class/proposed-value labels, explicit dates, unallocated unequal installments; generic preferred class is not rebound to PNA',
            'widget_request': '(page_body: bytes, page_receipt: dict, script_body: bytes, identity: manifest_record) -> tuple[url, discovery_proof]; exact published ID/ticker/template',
            'parse_widget_history': '(body: bytes, receipt: dict, identity: manifest_record) -> dict; native positive numeric closes, actual monthly dates, 61-month coverage and missing months; never adjusts prices'}}
    immutable_json(Path(output), report)
    digest = hashlib.sha256(Path(output).read_bytes()).hexdigest()
    immutable_json(base / ('recovery-report-' + digest + '.json'), report)
    print(json.dumps({'report_path': str(Path(output).resolve()), 'assigned': len(results),
                      'typed_cash_codes': report['typed_cash_codes'],
                      'prices_61': report['codes_with_independent_61_month_prices'],
                      'integration_candidates': report['integration_candidates']}), flush=True)
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', type=Path, default=Path(__file__).parent / 'data/market')
    parser.add_argument('--manifest', type=Path)
    parser.add_argument('--plan', type=Path)
    parser.add_argument('--widgets', action='store_true')
    parser.add_argument('--audit-output', type=Path)
    parser.add_argument('--cvms', default='')
    parser.add_argument('--follow', action='store_true')
    args = parser.parse_args()
    if args.audit_output:
        audit(args.root, args.manifest, args.audit_output)
    elif args.widgets:
        collect_widgets(args.root, args.manifest)
    elif args.plan:
        collect_plan(args.root, args.plan)
    else:
        collect(args.root, args.manifest,
                {int(c) for c in args.cvms.split(',')} if args.cvms else None, args.follow)


if __name__ == '__main__':
    main()
