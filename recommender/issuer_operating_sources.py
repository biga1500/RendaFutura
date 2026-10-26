"""Bounded issuer-RI research; never fetch exchange widgets or publish model data."""

from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import math
from pathlib import Path
import re
import unicodedata
from urllib.parse import urlsplit
import urllib.error
import urllib.request
import uuid
import zipfile

from lxml import html

from .payment_sources import _bounded_zip, _table_rows


ISSUERS = {
    'B3SA3': {'landing': 'https://ri.b3.com.br/pt-br/proventos/',
              'legal': 'https://ri.b3.com.br/pt-br/termos-e-condicoes/'},
    'AGRO3': {'landing': 'https://ri.brasil-agro.com/informacoes-financeiras/jcp-e-dividendos/',
              'english': 'https://ri.brasil-agro.com/en/financial-information/jcp-and-dividends/',
              'legal': 'https://ri-brasilagro.mz-sites.com/condicoes-de-uso/'},
    'ALUP3': {'landing': 'https://ri.alupar.com.br/pt/dividendos/',
              'workbook': 'https://d169uzu5o4xu1k.cloudfront.net/2c93fb35-d1f1-4125-a660-c9b7378d3259/2026/39f75058-9268-4fcc-92d7-302272f23d8a.xlsx'},
    'TAEE3': {'landing': 'https://ri.taesa.com.br/informacoes-financeiras/dividendos/',
              'workbook': 'https://ri.taesa.com.br/wp-content/uploads/2018/11/Dividendos_2026-site.xlsx'},
}
MAX_BYTES = 12_000_000


def _approved(url):
    parts = urlsplit(url)
    return (parts.scheme == 'https' and not parts.username and not parts.password
            and parts.port in (None, 443)
            and (parts.hostname in {urlsplit(r['landing']).hostname for r in ISSUERS.values()}
                 or url in {r.get('workbook') for r in ISSUERS.values()}))


class _IssuerRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, response, code, message, headers, newurl):
        if not _approved(newurl):
            raise ValueError('unapproved_redirect_no_request_sent')
        return super().redirect_request(request, response, code, message, headers, newurl)


def acquire_source(folder, ticker, kind, *, refresh=False):
    """Download one allowlisted public document into content-addressed raw storage."""
    url = ISSUERS[ticker][kind]
    if not _approved(url):
        raise ValueError('unapproved_issuer_source')
    root = Path(folder) / 'raw/primary-research/operating'
    root.mkdir(parents=True, exist_ok=True)
    if not refresh:
        receipts = [json.loads(p.read_text(encoding='utf-8')) for p in root.glob(f'{ticker}-{kind}-*-source.json')]
        for metadata in sorted(receipts, key=lambda r: r.get('retrieved_at', ''), reverse=True):
            name = metadata.get('cache_name', '')
            if not re.fullmatch(r'[A-Z0-9]+-[a-z]+-[a-f0-9]{64}\.(?:html|xlsx)', name) or metadata.get('url') != url:
                continue
            content = (root / name).read_bytes()
            if hashlib.sha256(content).hexdigest() == metadata.get('sha256'):
                return content, metadata
    request = urllib.request.Request(url, headers={'User-Agent': 'RendaFutura-Academic/1.0',
                                                  'Accept': 'text/html,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
    opener = urllib.request.build_opener(_IssuerRedirect())
    with opener.open(request, timeout=25) as response:
        content = response.read(MAX_BYTES + 1)
        final_url, content_type = response.geturl(), response.headers.get('Content-Type', '')
    if len(content) > MAX_BYTES:
        raise ValueError('bounded_response_size_exceeded')
    if not _approved(final_url):
        raise ValueError('unapproved_final_url')
    digest = hashlib.sha256(content).hexdigest()
    suffix = '.xlsx' if kind == 'workbook' else '.html'
    name = f'{ticker}-{kind}-{digest}{suffix}'
    target = root / name
    if target.exists() and target.read_bytes() != content:
        raise ValueError('immutable_cache_conflict')
    if not target.exists():
        target.write_bytes(content)
    metadata = {'id': f'issuer-ri:{ticker}:{kind}:{digest}', 'ticker': ticker, 'kind': kind,
                'url': url, 'final_url': final_url, 'provider': ticker + ' issuer RI', 'authority': 'primary',
                'content_type': content_type, 'sha256': digest, 'bytes': len(content),
                'retrieved_at': datetime.now(timezone.utc).isoformat(), 'cache_name': name,
                'file': target.relative_to(Path(folder)).as_posix(),
                'access_scope': 'issuer_investor_relations_not_exchange_market_data'}
    receipt = root / f'{ticker}-{kind}-{digest}-source.json'
    if not receipt.exists():
        receipt.write_text(json.dumps(metadata, ensure_ascii=True, indent=2, allow_nan=False), encoding='utf-8')
    return content, metadata


def _norm(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().lower()
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', text).split())


def _date(value, *, english=False, embedded=False):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or '').strip()
    if embedded:
        matches = re.findall(r'\b\d{1,2}/\d{1,2}/\d{4}\b', text)
        if len(matches) != 1:
            return None
        text = matches[0]
    for fmt in ('%Y-%m-%d', '%m/%d/%Y' if english else '%d/%m/%Y'):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _number(value, *, english=False, agro_mixed=False):
    if isinstance(value, bool):
        return None
    try:
        if isinstance(value, (int, float, Decimal)):
            parsed = Decimal(str(value))
        else:
            text = str(value or '').replace('\xa0', '').replace('R$', '').strip()
            if english or agro_mixed and re.fullmatch(r'\d+\.\d{4,}', text):
                if not re.fullmatch(r'(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?', text):
                    return None
                parsed = Decimal(text.replace(',', ''))
            else:
                if not re.fullmatch(r'(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?', text):
                    return None
                parsed = Decimal(text.replace('.', '').replace(',', '.'))
        result = float(parsed)
        return result if parsed.is_finite() and math.isfinite(result) and result >= 0 else None
    except (InvalidOperation, ValueError, OverflowError):
        return None


def _label(value):
    text = _norm(value)
    if text.startswith(('dividendos', 'dividendo', 'dividends')):
        return 'DIVIDENDO'
    if text.startswith(('jcp', 'jscp', 'jcsp', 'juros sobre', 'interest on shareholders')):
        return 'JCP'
    return None


def _columns(values, ticker):
    result = {}
    aliases = {
        'payment_date': {'data pagamento', 'data de pagamento', 'data do pagamento', 'payment date', 'pagamento'},
        'approval_date': {'data da aprovacao', 'data de aprovacao', 'evento data', 'aprovacao', 'announcement date'},
        'ex_date': {'data ex dividendos', 'data ex provento', 'ex dividend'},
        'record_date': {'data record', 'data record1'},
        'label': {'provento', 'tipo', 'type', 'remuneracao'},
        'reported_per_share': {'r acao', 'r share', 'por acao r', 'dps'},
        'reported_per_unit': {'por unit r', 'dpu'},
    }
    for index, value in enumerate(values):
        text = _norm(value)
        for key, names in aliases.items():
            if text in names:
                result[key] = index
        if text.startswith('valor bruto por a'):
            result['reported_per_share'] = index
            result['explicit_gross'] = True
    if 'payment_date' not in result or 'reported_per_share' not in result:
        return None
    if ticker == 'ALUP3' and not any(_norm(v) == 'ato societario' for v in values):
        return None
    return result


def _result(ticker, as_of, source):
    return {'ticker': ticker, 'as_of': str(as_of), 'events': [], 'corporate_actions': [],
            'source_ref': source['id'], 'source_url': source['url'], 'issues': [],
            'financially_confirmed': False, 'cash_model_eligible': False,
            'history_complete': False, 'corporate_action_coverage_complete': False,
            'limitations': ['Issuer history is not independent settlement evidence.',
                            'Approval, ex-date and retrieval are never substituted for publication or record date.',
                            'Observed first/last events do not establish complete cash or corporate-action coverage.',
                            'No share-scale adjustment, formula evaluation, or net-to-gross conversion.']}


def _event(result, raw, source, locator, *, english=False, explicit_gross=False,
           paid_context=None, numeric=None):
    label = _label(raw.get('label'))
    if label is None:
        return
    numeric = numeric or {}
    payment = _date(numeric.get('payment_date', raw.get('payment_date')), english=english)
    approval = _date(numeric.get('approval_date', raw.get('approval_date')), english=english, embedded=True)
    record = _date(numeric.get('record_date', raw.get('record_date')), english=english)
    ex_date = _date(numeric.get('ex_date', raw.get('ex_date')), english=english)
    amount = _number(numeric.get('reported_per_share', raw.get('reported_per_share')),
                     english=english, agro_mixed=result['ticker'] == 'AGRO3')
    issues = []
    if payment is None:
        issues.append('payment_date_not_exact')
    if amount is None:
        issues.append('reported_per_share_unavailable_or_invalid')
    if payment and any(day and payment < day for day in (approval, record)):
        issues.append('payment_before_approval_or_record')
    if record and ex_date and ex_date <= record:
        issues.append('ex_date_not_after_record')
    future = payment is not None and payment > result['as_of']
    status = ('announced' if future else 'invalid' if issues else
              'paid_reported' if paid_context else 'past_scheduled' if payment else 'unresolved')
    event = {'ticker': result['ticker'], 'share_class': 'ON' if result['ticker'] else None, 'label': label,
             'payment_date': payment, 'record_date': record, 'ex_date': ex_date, 'approval_date': approval,
             'reported_per_share': amount, 'gross_per_share': amount if explicit_gross else None,
             'amount_unit': 'BRL/share', 'amount_tax_basis': 'explicit_gross' if explicit_gross else 'unspecified_by_table',
             'reported_per_unit': _number(numeric.get('reported_per_unit', raw.get('reported_per_unit')), english=english),
             'quoted_quantity': 1, 'share_basis': 'issuer_reported_not_share_scale_reconciled',
             'status': status, 'payment_claim_basis': paid_context or 'dated_issuer_table_not_explicit_paid_confirmation',
             'payment_status_observed_at': source['retrieved_at'], 'paid_status_as_of_verified': False,
             'published_at': None, 'availability_date': None, 'available_as_of': None,
             'publication_verified': False, 'financially_confirmed': False, 'cash_model_eligible': False,
             'source_ref': source['id'], 'source_url': source['url'], 'source_row': locator,
             'raw_fields': {k: v.isoformat() if isinstance(v, (date, datetime)) else
                            str(v) if isinstance(v, float) and not math.isfinite(v) else v for k, v in raw.items()},
             'issues': issues}
    result['events'].append(event)


def _finish(result):
    seen, unique = set(), []
    for event in result['events']:
        key = tuple(event[k] for k in ('label', 'approval_date', 'record_date', 'ex_date', 'payment_date', 'reported_per_share'))
        if key in seen:
            result['issues'].append({'code': 'duplicate_source_event', 'source_row': event['source_row']})
            continue
        seen.add(key)
        unique.append(event)
    result['events'] = unique
    observed = [e['payment_date'] for e in unique if e['payment_date'] and e['payment_date'] <= result['as_of']
                and e['reported_per_share'] is not None and not e['issues']]
    result['coverage'] = {'first_observed_payment': min(observed) if observed else None,
                          'last_observed_payment': max(observed) if observed else None,
                          'basis': 'observed_event_span_not_completeness', 'history_complete': False,
                          'event_count': len(unique), 'statuses': dict(Counter(e['status'] for e in unique)),
                          'exact_payment_dates': sum(e['payment_date'] is not None for e in unique),
                          'reported_per_share_values': sum(e['reported_per_share'] is not None for e in unique),
                          'explicit_gross_per_share_values': sum(e['gross_per_share'] is not None for e in unique),
                          'verified_publication_dates': 0}
    result['status'] = ('partial' if result['issues'] or any(e['issues'] for e in unique) else
                        'available' if unique else 'not_found')
    return result


def parse_native_history(content, ticker, as_of, source):
    """Read issuer HTML only. Ex-date is never converted into a record/payment date."""
    result = _result(ticker, as_of, source)
    tree = html.fromstring(content.decode('utf-8'))
    english = '/en/' in source['url']
    excluded = Counter()
    legacy = {}
    for index, table in enumerate(tree.xpath('//table'), 1):
        target = result
        if ticker == 'B3SA3':
            headings = table.xpath('preceding::div[contains(@class,"accordion__item__header")][1]')
            scope = _norm(headings[0].text_content()) if headings else ''
            if scope != 'historico de proventos b3':
                excluded[scope or 'no_explicit_B3_history_heading'] += 1
                if scope not in ('historico de proventos bm fbovespa', 'historico de proventos cetip'):
                    continue
                target = legacy.setdefault(scope, _result(None, as_of, source))
                target['issuer_section'] = headings[0].text_content().strip()
                target['identity_mapping'] = 'not_assigned_to_requested_ticker'
        rows = list(_table_rows(table))
        header = next(((i, _columns(row, ticker)) for i, row in enumerate(rows) if _columns(row, ticker)), None)
        if header is None:
            continue
        start, columns = header
        for offset, values in enumerate(rows[start + 1:], start + 2):
            raw = {key: values[column] if column < len(values) else None
                   for key, column in columns.items() if key != 'explicit_gross'}
            if ticker == 'ALUP3':
                if not raw.get('reported_per_share') or any(_norm(v) in ('total', 'pagamento total de dividendos') for v in values):
                    continue
                raw['label'] = 'Dividendos'
            _event(target, raw, source, f'table:{index}:row:{offset}', english=english,
                   explicit_gross=columns.get('explicit_gross', False))
    if excluded:
        result['excluded_table_scopes'] = dict(excluded)
    if legacy:
        result['other_security_histories'] = [_finish(series) for series in legacy.values()]
    if ticker == 'B3SA3':
        for node in tree.xpath('//p'):
            text = _norm(node.text_content())
            match = re.search(r'desdobramento de (\d+) para (\d+) acoes realizado em (\d{2}) (\d{2}) (\d{4})', text)
            if match:
                old, new, day, month, year = map(int, match.groups())
                if old <= 0 or new <= 0:
                    continue
                result['corporate_actions'].append({'ticker': ticker, 'type': 'split',
                    'effective_date': date(year, month, day).isoformat(), 'ratio': new / old,
                    'published_at': None, 'availability_date': None, 'share_scale_applied': False,
                    'source_ref': source['id'], 'source_url': source['url'], 'basis': 'explicit_issuer_history_note'})
    result['corporate_actions_search'] = ('explicit_issuer_split_note_extracted' if result['corporate_actions'] else
                                          'no_explicit_action_table_in_selected_history_not_proof_of_no_actions')
    return _finish(result)


def parse_workbook(content, ticker, as_of, source):
    """Use native numeric cells or their stored caches, never execute Excel formulas."""
    from openpyxl import load_workbook

    result = _result(ticker, as_of, source)
    with _bounded_zip(content):
        pass
    formulas = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
    try:
        for sheet in formulas:
            if sheet.max_row > 10000 or sheet.max_column > 64:
                raise ValueError('bounded_worksheet_dimensions_exceeded')
            values = [list(row) for row in cached[sheet.title].values]
            raw_values = [list(row) for row in sheet.values]
            # Propagate only cells covered by an explicit workbook merge.
            for merged in sheet.merged_cells.ranges:
                for r in range(merged.min_row - 1, merged.max_row):
                    for c in range(merged.min_col - 1, merged.max_col):
                        values[r][c] = values[merged.min_row - 1][merged.min_col - 1]
                        raw_values[r][c] = raw_values[merged.min_row - 1][merged.min_col - 1]
            header = next(((i, _columns(row, ticker)) for i, row in enumerate(raw_values[:20]) if _columns(row, ticker)), None)
            if header is None:
                continue
            start, columns = header
            paid_context = ('issuer_workbook_Proventos_Pagos_summary' if ticker == 'TAEE3' and
                            any(_norm(cell) == 'proventos pagos' for row in raw_values for cell in row) else None)
            for index in range(start + 1, len(values)):
                raw = {key: raw_values[index][column] for key, column in columns.items() if key != 'explicit_gross'}
                numeric = {key: values[index][column] for key, column in columns.items() if key != 'explicit_gross'}
                if ticker == 'ALUP3':
                    if raw.get('reported_per_share') is None or all(raw.get(k) is None for k in ('payment_date', 'approval_date')):
                        continue
                    if any(_norm(v) in ('total', 'pagamento total de dividendos') for v in raw_values[index]):
                        continue
                    raw['label'] = 'Dividendos'
                before = len(result['events'])
                _event(result, raw, source, f'{sheet.title}:row:{index + 1}', paid_context=paid_context,
                       numeric=numeric, explicit_gross=columns.get('explicit_gross', False))
                if len(result['events']) != before:
                    event = result['events'][-1]
                    cell = sheet.cell(index + 1, columns['reported_per_share'] + 1)
                    event['amount_cell'] = cell.coordinate
                    event['amount_value_basis'] = 'stored_formula_cache' if cell.data_type == 'f' else 'native_numeric_cell'
                    event['date_value_basis'] = {
                        key: 'stored_formula_cache' if sheet.cell(index + 1, columns[key] + 1).data_type == 'f'
                        else 'native_cell_or_explicit_merge'
                        for key in ('payment_date', 'approval_date', 'record_date', 'ex_date') if key in columns}
                    if cell.data_type == 'f' and numeric.get('reported_per_share') is None:
                        event['issues'].append('formula_cache_unavailable')
        result['corporate_actions_search'] = 'no_explicit_action_table_in_selected_workbook_not_proof_of_no_actions'
        return _finish(result)
    finally:
        formulas.close()
        cached.close()


def _crosscheck(primary, alternatives):
    for alternative in alternatives:
        left, right = defaultdict(list), defaultdict(list)
        for event in primary['events']:
            left[event['label'], event['approval_date'], event['reported_per_share']].append(event)
        for event in alternative['events']:
            right[event['label'], event['approval_date'], event['reported_per_share']].append(event)
        for key in left.keys() & right.keys():
            if len(left[key]) != 1 or len(right[key]) != 1:
                continue
            a, b = left[key][0], right[key][0]
            for field in ('payment_date', 'record_date', 'ex_date'):
                if a[field] is None or b[field] is None or a[field] == b[field]:
                    continue
                code = f'conflicting_primary_{field}s'
                primary['issues'].append({'code': code, 'field': field,
                    'source_ref': a['source_ref'], 'alternative_source_ref': b['source_ref'],
                    'approval_date': a['approval_date'], 'value': a[field], 'alternative_value': b[field]})
                a['issues'].append(code)
                if a['status'] != 'announced':
                    a['status'] = 'invalid'
    return _finish(primary)


def collect_operating_sources(folder, as_of, *, tickers=None, refresh=False):
    """Acquire at most seven known documents, retain all dated observations and limitations."""
    folder = Path(folder)
    selected = list(ISSUERS) if tickers is None else list(tickers)
    if len(selected) != len(set(selected)) or any(t not in ISSUERS for t in selected):
        raise ValueError('Only the four explicitly researched issuer identities are supported.')
    result = {'schema_version': 1, 'parser_version': 'issuer_operating_sources_v1',
              'as_of': str(as_of), 'instruments': {}, 'sources': [], 'attempts': []}
    for ticker in selected:
        parsed, landing = {}, None
        for kind in ('landing', 'english', 'workbook'):
            if kind not in ISSUERS[ticker]:
                continue
            try:
                if kind == 'workbook' and (landing is None or ISSUERS[ticker][kind] not in landing.xpath('//a/@href')):
                    raise ValueError('workbook_not_linked_by_downloaded_issuer_landing_no_request_sent')
                content, source = acquire_source(folder, ticker, kind, refresh=refresh)
                if kind == 'workbook':
                    source = {**source, 'discovered_from': ISSUERS[ticker]['landing'],
                              'discovery_source_ref': parsed['landing']['source_ref'], 'download_link_verified': True}
                result['sources'].append(source)
                if kind == 'landing':
                    landing = html.fromstring(content.decode('utf-8'))
                parser = parse_workbook if kind == 'workbook' else parse_native_history
                parsed[kind] = parser(content, ticker, as_of, source)
                result['attempts'].append({'ticker': ticker, 'kind': kind, 'status': 'parsed', 'source_ref': source['id']})
            except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile) as error:
                result['attempts'].append({'ticker': ticker, 'kind': kind, 'status': 'unavailable',
                                           'error_type': type(error).__name__, 'reason': str(error)[:200]})
        primary_kind = 'workbook' if 'workbook' in parsed else 'landing'
        if primary_kind not in parsed:
            result['instruments'][ticker] = {'ticker': ticker, 'status': 'unavailable', 'events': [],
                                             'cash_model_eligible': False, 'history_complete': False}
            continue
        primary = _crosscheck(parsed[primary_kind], [v for k, v in parsed.items() if k != primary_kind])
        primary['selected_source_kind'] = primary_kind
        primary['alternative_observations'] = {k: v for k, v in parsed.items() if k != primary_kind}
        result['instruments'][ticker] = primary
    root = folder / 'raw/primary-research/operating'
    root.mkdir(parents=True, exist_ok=True)
    name = 'operating-sources-' + datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S') + '-' + uuid.uuid4().hex[:8] + '.json'
    result['report_file'] = (root / name).relative_to(folder).as_posix()
    (root / name).write_text(json.dumps(result, ensure_ascii=True, indent=2, allow_nan=False), encoding='utf-8')
    return result
