"""Evidence acquisition and full-window audits for all 39 pending price codes.

This isolated research module never updates inventories, pointers or models.
Public acquisition uses the shared robots-aware collector. Existing evidence
is hash checked and retained, including disagreements and unavailable dates.
"""
import argparse
from collections import Counter
import csv
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import io
import json
import math
from pathlib import Path
import re
import sys
import unicodedata
from urllib.parse import urljoin, urlsplit
import zipfile

from lxml import html

from .bulk_pending_sources import PublicSourceCollector, approved


CACHE = 'raw/primary-research/all-pending/prices'
START, END = date(2021, 9, 30), date(2026, 9, 30)
SNAPSHOT = 'universe-20261025T224724-fda3d4bd'
ACTION_READINGS = (
    {'tickers': ['SLCE3'], 'cvm': 20745,
     'sha256': 'e59acb9da1594861bff8585a9329293673c25cfccc36ccb66184a118c80a41ba',
     'page': 1, 'kind': 'same_species_bonus', 'total_factor_decimal': '1.1',
     'ex_date': '2022-01-03', 'native_ratio': '1 new common share per 10 common shares',
     'semantic_checks': ['1 (uma) nova acao para cada 10 (dez) acoes ordinarias',
                         'a partir do dia 03/01/2022']},
    {'tickers': ['BMEB3', 'BMEB4'], 'cvm': 1325,
     'sha256': 'ba945d4e2c486a708b4ec76a4cb2233acc89d7afe86dddaf7e9c2094b6c0ca9a',
     'page': 1, 'kind': 'simultaneous_reverse_split_and_split',
     'total_factor_decimal': '2', 'ex_date': '2021-12-03',
     'native_ratio': '100 old to 1, then 1 to 200, same species',
     'semantic_checks': ['cada grupo de 100', 'cada acao sera desdobrada em 200',
                         'negociadas a partir do dia 03 de dezembro']},
    {'tickers': ['DEXP3', 'DEXP4'], 'cvm': 16632,
     'sha256': '70935fbb94f00837bd1d52456644dbb7300e19ef78615e81ce333f04241e07fc',
     'page': 2, 'kind': 'same_species_bonus', 'total_factor_decimal': '1.125',
     'ex_date': '2025-12-22', 'native_ratio': '1 new share per 8 existing shares',
     'semantic_checks': ['proporcao de 1 nova acao para cada 8 acoes',
                         'a partir do dia 22 de dezembro de 2025',
                         'novas acoes ordinarias e preferenciais']},
    {'tickers': ['IGTI3'], 'cvm': 8672,
     'sha256': '8f5f80feee1a2d2b0124fbcce3bd683b2f4579b693a5a394737b1b43cfe82189',
     'page': 1, 'kind': 'split', 'total_factor_decimal': '10',
     'ex_date': '2021-12-27', 'native_ratio': '10 shares for each existing same-species share',
     'semantic_checks': ['proporcao de 10 (dez) acoes para cada acao',
                         'desdobramento a partir de 27 de dezembro de 2021']},
    {'tickers': ['COGN3'], 'cvm': 17973,
     'sha256': '4606f00a86e4a533e242a4865b6615b45dfc2b6647802e38f8c14506d513fbb0',
     'page': 6, 'kind': 'same_species_bonus', 'total_factor_decimal': '1.1',
     'ex_date': '2025-12-26', 'native_ratio': '1 bonus share per 10 common shares',
     'semantic_checks': ['proporcao de 1 Acao Bonificada para cada 10 acoes ordinarias',
                         'a partir de 26 de dezembro de 2025']},
    {'tickers': ['KLBN11', 'KLBN3', 'KLBN4'], 'cvm': 12653,
     'sha256': 'ec2f0c8d69753c2b72e49f46a8d93e57aca45c7bd8e4d15ccaf25079c2c06df3',
     'page': 3, 'kind': 'same_species_bonus_with_explicit_unit_allocation',
     'total_factor_decimal': '1.01', 'ex_date': '2025-12-18',
     'native_ratio': '1 new share per 100 same-species shares; 1 ON plus 4 PN per unit',
     'semantic_checks': ['cada especie para cada 100 (cem) acoes da mesma especie',
                         'a partir de 18 de dezembro de 2025',
                         '1 (uma) acao ordinaria e 4 (quatro) acoes preferenciais']},
    {'tickers': ['SLCE3'], 'cvm': 20745,
     'sha256': 'c80b4955e34e544a4010c8b6ddac279c8b8996b1d2caf7ff63e60cd6a43497cf',
     'page': 1, 'kind': 'same_species_bonus', 'total_factor_decimal': '1.125',
     'ex_date': '2026-01-02', 'native_ratio': '1 new common share per 8 common shares',
     'semantic_checks': ['1 (uma) nova acao ordinaria para cada 8 (oito) acoes',
                         'a partir do dia 02/01/2026']},
    {'tickers': ['SBSP3'], 'cvm': 14443,
     'sha256': '1892f7aa498728afd4014b995c207849e282f2a7e4fa61eab6ffe26b1ba79653',
     'page': 1, 'kind': 'split', 'total_factor_decimal': '5',
     'ex_date': '2026-04-29', 'native_ratio': '1 common share becomes 5 common shares',
     'semantic_checks': ['cada 1 (uma) acao ordinaria a representar 5 (cinco) acoes ordinarias',
                         'a partir de 29 de abril de 2026']},
)


def require(value, message):
    if not value:
        raise ValueError(message)


def digest(body):
    return hashlib.sha256(body).hexdigest()


def encoded(value):
    return (json.dumps(value, ensure_ascii=True, indent=2, sort_keys=True,
                       allow_nan=False) + '\n').encode('ascii')


def read_source(root, receipt):
    root = Path(root).resolve()
    relative = receipt['file']
    require(isinstance(relative, str) and not Path(relative).is_absolute()
            and ':' not in relative and '\\' not in relative
            and '..' not in relative.split('/'), 'unsafe_price_evidence_path')
    path = root / relative
    require(path.resolve().is_relative_to(root) and not path.is_symlink(),
            'unsafe_price_evidence_path')
    body = path.read_bytes()
    require(len(body) <= 30_000_000 and digest(body) == receipt['sha256'],
            'price_evidence_hash_mismatch')
    if 'bytes' in receipt:
        require(len(body) == receipt['bytes'], 'price_evidence_size_mismatch')
    return body


def source(root, relative):
    path = Path(root) / relative
    body = path.read_bytes()
    return body, {'file': relative, 'sha256': digest(body), 'bytes': len(body),
                  'access': 'existing_cache_read_only', 'retrieved_at': None,
                  'published_at': None}


def monthly_audit(body, ticker):
    """Keep native close values; an absent month remains an explicit absence."""
    payload = json.loads(body)
    result = payload['chart']['result'][0]
    meta = result['meta']
    require(meta.get('symbol') == ticker + '.SA' and meta.get('currency') == 'BRL'
            and meta.get('dataGranularity') == '1d', 'exact_daily_identity_required')
    stamps = result.get('timestamp', [])
    closes = result['indicators']['quote'][0]['close']
    require(len(stamps) == len(closes), 'daily_length_mismatch')
    seen, daily, invalid = set(), {}, []
    for i, (stamp, close) in enumerate(zip(stamps, closes)):
        day = datetime.fromtimestamp(stamp, timezone.utc).date()
        require(day not in seen, 'duplicate_daily_date')
        seen.add(day)
        if close is None:
            continue
        if isinstance(close, bool) or not isinstance(close, (int, float)) \
                or not math.isfinite(close) or close <= 0:
            invalid.append({'index': i, 'date': day.isoformat(), 'value': close})
            continue
        daily[day] = close
    monthly = {}
    for day, value in sorted(daily.items()):
        if day <= END:
            monthly[day.strftime('%Y-%m')] = {'observed_on': day.isoformat(),
                                             'native_close_BRL': value}
    anchors = []
    for offset in range(61):
        absolute = START.year * 12 + START.month - 1 + offset
        year, month0 = divmod(absolute, 12)
        month = f'{year:04d}-{month0 + 1:02d}'
        anchors.append({'month': month, **monthly.get(month, {
            'observed_on': None, 'native_close_BRL': None})})
    actions = []
    for action in result.get('events', {}).get('splits', {}).values():
        day = datetime.fromtimestamp(action['date'], timezone.utc).date()
        if START <= day <= END:
            actions.append({'date': day.isoformat(), 'original_provider_row': action,
                            'primary_action_verified': False})
    return {'required_months': 61, 'anchors': anchors,
            'observed_required_months': sum(a['observed_on'] is not None for a in anchors),
            'missing_months': [a['month'] for a in anchors if a['observed_on'] is None],
            'first_daily_observation': min(daily).isoformat() if daily else None,
            'last_daily_observation': max(daily).isoformat() if daily else None,
            'invalid_observations': invalid, 'provider_actions_not_primary': actions,
            'adjusted_close_used': False, 'prices_transformed': False}


def inspect_html(body, url, tickers):
    decoded = body.decode('utf-8', errors='replace')
    tree = html.fromstring(decoded)
    content = html.fromstring(decoded)
    for node in content.xpath('//script|//style|//nav|//header|//footer'):
        node.drop_tree()
    text = ' '.join(content.text_content().split())
    links = []
    for node in tree.xpath('//*[@href or @src or @data-url]'):
        for attr in ('href', 'src', 'data-url'):
            if not node.get(attr):
                continue
            target = urljoin(url, node.get(attr))
            label = ' '.join(node.text_content().split())[:180]
            if approved(target) and (re.search(
                    r'cotac|cota%C3|gr.fic|historic|stockinfo|bonific|desdobr|grupam|'
                    r'acionist|relevante|investidor|\.pdf', target + ' ' + label, re.I)
                    or any(t in target + ' ' + label for t in tickers)):
                item = {'url': target, 'label': label, 'attribute': attr}
                if item not in links:
                    links.append(item)
    ids = sorted(set(re.findall(
        r'''(?:stockInfoId|stockinfoId)\s*[:=]\s*["']([0-9a-f-]{36})["']''', decoded)))
    configured = sorted(set(t for value in re.findall(
        r'''tickers\s*:\s*["']([A-Z0-9,]+)["']''', decoded) for t in value.split(',')))
    excerpts = []
    for match in re.finditer(r'bonifica|desdobr|grupamento|negocia|hist.rico|dividendo|'
                             r'juros sobre|a..es preferenciais|AUAU3|BRST3|RENT4|VTRU3|PASS3', text, re.I):
        excerpt = text[max(0, match.start() - 120):match.end() + 230]
        if excerpt not in excerpts:
            excerpts.append(excerpt)
    return {'title': tree.xpath('string(//title)'), 'text_characters': len(text),
            'text_sha256': digest(text.encode()), 'semantic_excerpts': excerpts[:40],
            'published_links': links[:140], 'stock_info_ids': ids,
            'configured_tickers': configured, 'publication_date': None,
            'table_rows': [[' '.join(cell.text_content().split()) for cell in row.xpath('./th|./td')]
                           for row in tree.xpath('//table//tr')],
            'publication_date_policy': 'not_inferred_from_retrieval_or_HTTP_Date'}


def published_history_request(page_body, page_receipt, bundle_body, bundle_receipt, ticker):
    """Instantiate only the issuer's actual configured public widget template."""
    page = inspect_html(page_body, page_receipt['url'], [ticker])
    require(ticker in page['configured_tickers'] and len(page['stock_info_ids']) == 1,
            'exact_ticker_not_in_published_widget_configuration')
    require(bundle_receipt['url'] in [v['url'] for v in page['published_links']],
            'bundle_not_published_by_issuer')
    script = bundle_body.decode('utf-8')
    bases = re.findall(r'API_STOCKINFO_URL:"(https://[^\"]+)"', script)
    formula = 't.API_URL+"/"+t.STOCK_ID+"/history?from="'
    require(len(bases) == 1 and all(v in script for v in (
        formula, 't.ADJUSTED=!1', '"&tickers="+t.TICKER',
        '"&adjusted="+t.ADJUSTED', 'noProvents:"No"')),
        'unverified_public_widget_template')
    url = (f'{bases[0]}/{page["stock_info_ids"][0]}/history?from=2021-09-01'
           f'&to=2026-09-30&tickers={ticker}&adjusted=false')
    require(approved(url), 'unapproved_vendor_template')
    return url, {'kind': 'issuer_published_vendor_template',
                 'source_url': page_receipt['url'], 'page': page_receipt,
                 'bundle': bundle_receipt, 'formula': formula,
                 'ticker': ticker, 'adjusted': False}


def verified_native_history(root, identity, expected_identity, receipt, page, bundle):
    """Reviewable adapter output only; callers do not gain model admission."""
    require(identity == expected_identity and set(identity) == {'ticker', 'cvm', 'share_class', 'isin'}
            and type(identity['cvm']) is int, 'exact_CVM_class_ISIN_identity_required')
    page_body, bundle_body = read_source(root, page), read_source(root, bundle)
    url, proof = published_history_request(page_body, page, bundle_body, bundle, identity['ticker'])
    require(receipt['url'] == url and receipt.get('status', receipt.get('http_status')) == 200
            and not receipt.get('error'), 'native_history_request_proof_mismatch')
    rows, inactive = vendor_rows(read_source(root, receipt), identity['ticker'])
    return {'identity': dict(identity), 'rows': rows, 'inactive_rows': inactive,
            'receipt': receipt, 'discovery_proof': proof, 'model_admission': False,
            'prices_transformed': False, 'corporate_actions_applied': False,
            'ultimate_upstream_independence_verified': False}


def vendor_rows(body, ticker):
    payload = json.loads(body, parse_float=Decimal)
    require(payload.get('success') is True, 'vendor_response_not_successful')
    data = payload['data']
    require(data.get('source') == 'Enfoque' and not data.get('errors'),
            'unverified_vendor_source')
    matches = [r for r in data['tickers'] if r.get('ticker') == ticker]
    require(len(matches) == 1, 'vendor_exact_ticker_required')
    seen, rows, inactive = set(), [], []
    for original in matches[0]['history']:
        day = date.fromisoformat(original['date'])
        require(day not in seen, 'duplicate_vendor_quote_date')
        seen.add(day)
        if not date(2021, 9, 1) <= day <= END:
            raise ValueError('vendor_quote_outside_requested_window')
        close = original['closed']
        require(not isinstance(close, bool) and isinstance(close, (Decimal, int))
                and Decimal(close).is_finite() and close >= 0, 'invalid_vendor_close')
        row = {'date': day.isoformat(), 'native_close_decimal': str(close),
               'native_fields': {k: str(v) if isinstance(v, Decimal) else v
                                 for k, v in original.items()}}
        if close == 0 or original.get('trade') == 0 or original.get('share') == 0:
            inactive.append(row)
        else:
            require(original['low'] <= close <= original['high'], 'vendor_OHLC_inconsistent')
            rows.append(row)
    return sorted(rows, key=lambda r: r['date']), inactive


def compare_vendor(audit, rows):
    bydate = {r['date']: Decimal(r['native_close_decimal']) for r in rows}
    comparisons = []
    for anchor in audit['anchors']:
        day = anchor['observed_on']
        value = bydate.get(day)
        stored = anchor['native_close_BRL']
        comparisons.append({**anchor, 'independent_native_close_decimal':
            str(value) if value is not None else None,
            'same_date_observed': value is not None,
            'native_values_agree_within_existing_band': None if value is None else
            math.isclose(float(value), stored, rel_tol=.006, abs_tol=.015),
            'no_share_adjustment_applied': True})
    return {'required_61_months': comparisons,
            'same_date_monthly_observations': sum(r['same_date_observed'] for r in comparisons),
            'native_month_end_disagreements': [r for r in comparisons
                if r['native_values_agree_within_existing_band'] is False],
            'interpretation': 'raw values only; corporate actions remain separate evidence'}


def full_daily_comparison(body, ticker, rows, audit):
    """Expose every disagreement; replay the original gate as a diagnostic."""
    from .share_basis import factor_after, reconcile_prices
    result = json.loads(body)['chart']['result'][0]
    daily = {datetime.fromtimestamp(stamp, timezone.utc).date(): close
             for stamp, close in zip(result['timestamp'], result['indicators']['quote'][0]['close'])
             if close is not None}
    actions = [{'date': r['date'], 'ratio': float(r['original_provider_row']['numerator']) /
                float(r['original_provider_row']['denominator'])}
               for r in audit['provider_actions_not_primary']]
    monthly_dates = {r['observed_on'] for r in audit['anchors']}
    comparisons, references = [], {}
    for row in rows:
        day = date.fromisoformat(row['date'])
        if not START <= day <= END:
            continue
        raw = float(Decimal(row['native_close_decimal']))
        stored = daily.get(day)
        factor = factor_after(day, actions, END)
        references[day] = raw
        comparisons.append({'date': row['date'], 'native_vendor_close_decimal': row['native_close_decimal'],
            'stored_native_close_BRL': stored, 'reported_share_factor_diagnostic_only': factor,
            'raw_values_compatible': stored is not None and math.isclose(stored, raw,
                rel_tol=.006, abs_tol=.015),
            'provider_action_adjusted_compatible_diagnostic_only': stored is not None and math.isclose(
                stored, raw / factor, rel_tol=.006, abs_tol=.015),
            'is_required_monthly_date': row['date'] in monthly_dates})
    try:
        _, anchors, mode = reconcile_prices(body, ticker, actions, references, END,
            evidence_provider='issuer_published_Enfoque_native_history')
        trial = {'status': 'passed_unchanged_price_gate_diagnostic_only',
                 'independent_dates': len(anchors), 'mode': mode}
    except ValueError as error:
        trial = {'status': 'blocked_by_unchanged_price_gate', 'reason': str(error)}
    return {'independent_exact_dates': len(comparisons), 'all_exact_date_comparisons': comparisons,
            'raw_disagreements': [r for r in comparisons if not r['raw_values_compatible']],
            'provider_action_adjusted_disagreements_diagnostic_only': [r for r in comparisons
                if not r['provider_action_adjusted_compatible_diagnostic_only']],
            'unchanged_price_gate_trial_using_every_native_date': trial,
            'primary_action_chain_still_required': bool(actions),
            'favorable_subset_selected': False, 'model_admission': False}


def cached_receipts(root):
    result = []
    base = Path(root) / 'raw/primary-research'
    for path in base.glob('*/*.json'):
        if path.stat().st_size > 2_000_000:
            continue
        try:
            value = json.loads(path.read_bytes())
        except (OSError, ValueError):
            continue
        if isinstance(value, dict) and value.get('url') and value.get('file') \
                and value.get('sha256'):
            result.append({**value, 'receipt_file': path.relative_to(root).as_posix()})
    return result


def prior_diagnostics(root, ticker):
    paths = (Path(root) / 'raw/primary-research/selected-prices').glob('report-*.json')
    for path in sorted(paths):
        report = json.loads(path.read_bytes())
        matches = [r for r in report.get('records', []) if r.get('ticker') == ticker]
        if matches:
            record = matches[0]
            return {'report': source(root, path.relative_to(root).as_posix())[1],
                    'diagnostic': record['diagnostic'],
                    'next_safe_source_work': record.get('next_safe_source_work'),
                    'focused_closure': next((r for r in report.get(
                        'focused_candidate_evidence_closures', []) if r['ticker'] == ticker), None)}
    for path in sorted((Path(root) / 'raw/primary-research/noncash-price').glob('report-*.json')):
        report = json.loads(path.read_bytes())
        for row in report.get('records', []):
            if row.get('identity', {}).get('ticker') == ticker:
                return {'report': source(root, path.relative_to(root).as_posix())[1],
                        'identity': row['identity'], 'blockers': row['blockers'],
                        'event_boundary_observations': row.get('event_boundary_observations'),
                        'noncash_financial_receipt': row.get('noncash_financial_receipt'),
                        'numerical_backadjustment_hypotheses_not_adopted': True}
    return None


def compact(value):
    return re.sub('[^a-z0-9]', '', unicodedata.normalize('NFKD', value).encode(
        'ascii', errors='ignore').decode().lower())


def reviewed_primary_actions(root, identity, outcomes):
    from pypdf import PdfReader
    actions = []
    for reading in ACTION_READINGS:
        if identity['ticker'] not in reading['tickers']:
            continue
        require(identity['cvm'] == reading['cvm'], 'primary_action_CVM_mismatch')
        matches = [v for v in outcomes if v.get('receipt', {}).get('sha256') == reading['sha256']]
        if not matches:
            continue
        original = matches[0]
        receipt = original['receipt']
        discovery = receipt.get('discovery', original.get('discovery'))
        require(discovery['cvm'] == identity['cvm'] and receipt['url'] == discovery['url'],
                'primary_action_native_index_identity_mismatch')
        index = discovery['index_source']
        with zipfile.ZipFile(io.BytesIO(read_source(root, index))) as archive:
            rows = csv.DictReader(io.StringIO(archive.read(discovery['member']).decode('latin1')),
                                  delimiter=';')
            row = next(r for line, r in enumerate(rows, 2) if line == discovery['csv_line'])
        require(row == discovery['index_row'] and int(row['Codigo_CVM']) == identity['cvm'],
                'primary_action_native_index_row_changed')
        document = PdfReader(io.BytesIO(read_source(root, receipt)))
        text = document.pages[reading['page'] - 1].extract_text()
        require(all(compact(fragment) in compact(text) for fragment in reading['semantic_checks']),
                'primary_action_semantic_reading_changed:' + identity['ticker'])
        actions.append({**reading, 'identity': dict(identity), 'source': receipt,
                        'published_at': row['Data_Entrega'], 'retrieved_at': receipt['retrieved_at'],
                        'native_index_verified': True, 'ratio_and_ex_date_verified': True,
                        'factor_fitted_to_quotes': False, 'applied_to_price': False,
                        'cash_payment_or_bank_settlement_inferred': False,
                        'model_admission': False})
    if identity['ticker'] == 'IRBR3':
        pin = '0bf466235fedea0039db67801a141f6997d3b05521b0d884584f3dca14ca39d6'
        matches = [v for v in outcomes if v.get('receipt', {}).get('sha256') == pin]
        if matches:
            require(identity == {'ticker': 'IRBR3', 'cvm': 24180, 'share_class': 'ON',
                                 'isin': 'BRIRBRACNOR4'}, 'IRBR_native_action_identity_changed')
            receipt = matches[0]['receipt']
            tree = html.fromstring(read_source(root, receipt).decode('utf-8'))
            text = compact(tree.text_content())
            require(all(compact(part) in text for part in (
                '30 (trinta) acoes convertidas em 1 (uma) acao',
                'A partir do dia 25 de janeiro de 2023')), 'IRBR_native_action_text_changed')
            actions.append({'identity': dict(identity), 'source': receipt,
                'kind': 'reverse_split', 'ex_date': '2023-01-25',
                'ratio_numerator': 1, 'ratio_denominator': 30, 'native_ratio': '1/30',
                'published_at': None, 'retrieved_at': receipt['retrieved_at'],
                'ratio_and_ex_date_verified': True, 'factor_fitted_to_quotes': False,
                'applied_to_price': False, 'model_admission': False,
                'locator': 'Grupamento section, literal issuer paragraph'})
    if identity['ticker'] == 'LREN3':
        from .issuer_quote_extension_research import primary_actions
        require(identity == {'ticker': 'LREN3', 'cvm': 8133, 'share_class': 'ON',
                             'isin': 'BRLRENACNOR1'}, 'LREN_cached_action_identity_changed')
        for action in primary_actions(root, 'LREN3'):
            receipt = action['receipt']
            read_source(root, receipt)
            actions.append({**action, 'identity': dict(identity), 'source': receipt,
                'native_ratio': f'{action["new_shares"]} new per {action["per_old_shares"]} old',
                'published_at': receipt['discovery']['row']['Data_Entrega'],
                'retrieved_at': receipt['retrieved_at'], 'ratio_and_ex_date_verified': True,
                'access': 'existing_hash_pinned_semantic_adapter_read_only',
                'factor_fitted_to_quotes': False, 'applied_to_price': False, 'model_admission': False})
    return actions


def native_listing_readings(root, ticker, outcomes):
    pins = {
        'VTRU3': ('fbe2d4f2bdc9ce25bddb1ba34fd2dda1b9372f1e9fe622f73c09d6b31cb477f1',
                  2, '2024-06-10', 'a partir de 10 de junho de 2024', 'new_BRL_listing_after_merger'),
        'PASS3': ('0e40342c17bea8a40ba56e615f9942b3d15b8e0dc7b63943b9a26f3238d8640d',
                  2, '2026-05-11', 'a partir de 11 de maio de 2026', 'secondary_offer_shares_begin_trading'),
        'AUAU3': ('61585e05b7491e963eee9581bca4beb99f9c372f6b60438402f499a286a0279b',
                  None, '2026-01-05', 'a partir de 5 de janeiro de 2026', 'combination_new_ticker'),
    }
    if ticker not in pins:
        return []
    pin, page, day, fragment, kind = pins[ticker]
    matches = [o for o in outcomes if o.get('receipt', {}).get('sha256') == pin]
    if not matches:
        return []
    receipt = matches[0]['receipt']
    raw = read_source(root, receipt)
    if page:
        from pypdf import PdfReader
        text = PdfReader(io.BytesIO(raw)).pages[page - 1].extract_text()
    else:
        text = html.fromstring(raw.decode('utf-8')).text_content()
    require(compact(fragment) in compact(text), 'native_listing_semantics_changed')
    return [{'date': day, 'kind': kind, 'literal_source_fragment': fragment, 'page': page,
             'source': receipt, 'predecessor_history_spliced': False,
             'five_year_history_waived': False}]


def inspect_native_bonus_workbook(body):
    """Read original cells, retaining fractional counts and their source dates."""
    import openpyxl
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    try:
        name = next(n for n in workbook.sheetnames if compact(n) == 'bonificacao')
        sheet = workbook[name]
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                value = cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else cell.value
                cells.append({'cell': cell.coordinate, 'native_value': value,
                              'native_type': cell.data_type})
        return {'sheet': name, 'cells': cells, 'file_modified': False,
                'source_date_not_relabelled_as_ex_date': True,
                'fractional_reported_share_counts_not_rounded': True,
                'formulas_evaluated': False, 'historical_table_not_assumed_complete_to_cutoff': True}
    finally:
        workbook.close()


def primary_index_candidates(root, cvms, sequences=None):
    """CVM native index rows identify issuer, version and actual publication."""
    candidates = []
    for path in sorted((Path(root) / 'raw/primary-research/cvm-documents').glob('index-20*.zip')):
        body, receipt = source(root, path.relative_to(root).as_posix())
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            for member in archive.namelist():
                if not member.endswith('.csv'):
                    continue
                rows = csv.DictReader(io.StringIO(archive.read(member).decode('latin1')), delimiter=';')
                for line, row in enumerate(rows, 2):
                    if int(row['Codigo_CVM']) not in cvms:
                        continue
                    if sequences is not None and not any(
                            f'numSequencia={seq}&' in row['Link_Download'] for seq in sequences):
                        continue
                    if not '2021-09-01' <= row['Data_Entrega'][:10] <= str(END):
                        continue
                    candidates.append({'url': row['Link_Download'], 'cvm': int(row['Codigo_CVM']),
                        'published_at': row['Data_Entrega'], 'index_row': row,
                        'index_source': receipt, 'member': member, 'csv_line': line,
                        'index_subject_is_not_action_ratio_evidence': True})
    return candidates


def classified_outcome(row):
    audit = row['monthly_audit']
    if audit['missing_months']:
        return {'classification': 'incomplete_exact_security_61_month_history',
                'observed': audit['observed_required_months'], 'required': 61,
                'missing_months': audit['missing_months'],
                'predecessor_series_spliced': False}
    histories = [v['vendor_audits'][row['ticker']] for v in row['source_outcomes']
                 if row['ticker'] in v.get('vendor_audits', {})]
    conflicts = [v for v in histories if v['monthly_comparison']['native_month_end_disagreements']]
    if conflicts:
        return {'classification': 'native_vendor_history_recovered_requires_action_or_quote_reconciliation',
                'independent_histories': len(histories),
                'native_disagreement_months': sorted({a['month'] for v in conflicts for a in
                    v['monthly_comparison']['native_month_end_disagreements']}),
                'price_or_action_factor_inferred': False}
    if histories:
        return {'classification': 'native_history_recovered_model_admission_not_granted',
                'independent_histories': len(histories),
                'exact_date_monthly_matches': max(v['monthly_comparison'][
                    'same_date_monthly_observations'] for v in histories)}
    failures = sorted({v.get('inspection_or_acquisition_error') or
                       v.get('receipt', {}).get('error') for v in row['source_outcomes']} - {None})
    return {'classification': 'exact_native_history_not_recovered',
            'specific_acquisition_constraints': failures,
            'original_pending_reasons_retained': row['original_pending_reasons']}


class PriceResearch:
    def __init__(self, root, manifest, *, acquire=True):
        self.root = Path(root).resolve()
        self.manifest_path = Path(manifest).resolve()
        raw = self.manifest_path.read_bytes()
        self.manifest_hash = digest(raw)
        self.manifest = json.loads(raw)
        require(self.manifest['snapshot'] == SNAPSHOT, 'unexpected_pending_baseline')
        self.records = [r for r in self.manifest['records'] if r['task_group'] == 'price_and_actions']
        require(len(self.records) == len({r['ticker'] for r in self.records}) == 39,
                'all_39_price_codes_required')
        self.by_ticker = {r['ticker']: r for r in self.records}
        self.base = self.root / CACHE
        self.base.mkdir(parents=True, exist_ok=True)
        self.receipts = cached_receipts(self.root)
        self.collector = PublicSourceCollector(self.root, 'prices') if acquire else None
        self.extra_stops = set()
        for receipt in self.receipts:
            reason = str(receipt.get('stop_reason', '')) + str(receipt.get('stop', ''))
            if re.search('captcha|challenge|robots.*disallow', reason, re.I):
                self.extra_stops.add(urlsplit(receipt['url']).hostname)
        self.protected = {p: digest((self.root / p).read_bytes())
                          for p in ('current.json', 'inventory/current.json')}

    def save(self, kind, value):
        body = encoded(value)
        relative = f'{CACHE}/{kind}-{digest(body)}.json'
        path = self.root / relative
        if path.exists():
            require(path.read_bytes() == body, 'immutable_price_report_conflict')
        else:
            with path.open('xb') as stream:
                stream.write(body)
        return relative

    def fetch(self, tickers, url, discovery):
        require(tickers and all(t in self.by_ticker for t in tickers), 'outside_assigned_39')
        outcome = {'tickers': tickers, 'url': url, 'discovery': discovery,
                   'model_admission': False}
        candidates = [r for r in self.receipts if r['url'] == url]
        cached = next((r for r in candidates if r.get('status', r.get('http_status')) == 200
                       and not r.get('error') and not r.get('stop')), None)
        try:
            if cached:
                body = read_source(self.root, cached)
                receipt = cached
                outcome['access'] = 'existing_cache_read_only_no_branch_repeated'
            else:
                require(urlsplit(url).hostname not in self.extra_stops,
                        'prior_challenged_origin_no_retry')
                require(self.collector is not None, 'offline_research')
                receipt = self.collector.fetch(url, discovery=discovery)
                body = self.collector.body(receipt)
                outcome['access'] = 'shared_public_collector'
            outcome['receipt'] = receipt
            status = receipt.get('status', receipt.get('http_status'))
            if status == 200 and not receipt.get('error') and not receipt.get('stop'):
                if body.lstrip().startswith((b'<!DOCTYPE', b'<!doctype', b'<html', b'<HTML')):
                    outcome['inspection'] = inspect_html(body, url, tickers)
                elif body.lstrip().startswith(b'{'):
                    payload = json.loads(body)
                    outcome['json_keys'] = list(payload)
                    if '/history?' in url:
                        outcome['vendor_audits'] = {}
                        for ticker in tickers:
                            rows, inactive = vendor_rows(body, ticker)
                            audit = self.audit(ticker)['monthly_audit']
                            outcome['vendor_audits'][ticker] = {
            'rows': len(rows), 'inactive_rows': inactive,
                                'first_date': rows[0]['date'] if rows else None,
                                'last_date': rows[-1]['date'] if rows else None,
                                'monthly_comparison': compare_vendor(audit, rows)}
                elif body.startswith(b'%PDF-'):
                    from pypdf import PdfReader
                    document = PdfReader(io.BytesIO(body))
                    outcome['pdf_pages'] = len(document.pages)
                    outcome['pdf_page_texts'] = [{'page': i + 1, 'text': page.extract_text()}
                                                for i, page in enumerate(document.pages)]
        except (ValueError, KeyError, TypeError) as error:
            outcome['inspection_or_acquisition_error'] = str(error)
        relative = self.save('outcome', outcome)
        return {**outcome, 'outcome_file': relative}

    def audit(self, ticker):
        record = self.by_ticker[ticker]
        relative = record['source_inventory_record']['history']['source_file']
        body, receipt = source(self.root, relative)
        return {'ticker': ticker, 'cvm': record['cvm'], 'share_class': record['share_class'],
                'isin': record['isin'], 'original_pending_reasons': record['reasons'],
                'FCA_declared_website': record.get('website_declared_in_FCA'),
                'FCA_registry_filing': record.get('registry_filing'),
                'source': receipt, 'monthly_audit': monthly_audit(body, ticker),
                'prior_inspected_price_diagnostic': prior_diagnostics(self.root, ticker)}

    def report(self):
        outcomes = [(p.relative_to(self.root).as_posix(), json.loads(p.read_bytes()))
                    for p in self.base.glob('outcome-*.json')]
        represented = {v.get('receipt', {}).get('url') for _, v in outcomes}
        for path in self.base.glob('response-*.json'):
            receipt = json.loads(path.read_bytes())
            cvm = (receipt.get('discovery') or {}).get('cvm')
            if not cvm or receipt['url'] in represented:
                continue
            body = read_source(self.root, receipt)
            outcome = {'tickers': [t for t, r in self.by_ticker.items() if r['cvm'] == cvm],
                       'url': receipt['url'], 'receipt': receipt,
                       'discovery': receipt['discovery'],
                       'access': 'acquired_original_inspected_offline_no_retry',
                       'model_admission': False}
            if body.startswith(b'%PDF-'):
                from pypdf import PdfReader
                document = PdfReader(io.BytesIO(body))
                outcome['pdf_pages'] = len(document.pages)
                outcome['pdf_page_texts'] = [{'page': i + 1, 'text': page.extract_text()}
                                            for i, page in enumerate(document.pages)]
            else:
                outcome['native_PDF_not_returned'] = True
            saved = self.save('outcome', outcome)
            outcomes.append((saved, outcome))
        records = []
        for ticker in sorted(self.by_ticker):
            row = self.audit(ticker)
            row['source_outcomes'] = [{'outcome_file': p, **v} for p, v in outcomes
                                      if ticker in v['tickers']]
            for outcome in row['source_outcomes']:
                discovery = outcome.get('discovery') or {}
                if discovery.get('registry_filing') and discovery['registry_filing'] != row['FCA_registry_filing']:
                    outcome['original_discovery_filing_metadata_mismatch'] = True
                    outcome['authoritative_FCA_filing_from_pinned_manifest'] = row['FCA_registry_filing']
                receipt = outcome.get('receipt')
                if not receipt:
                    continue
                body = read_source(self.root, receipt)
                if 'numSequencia=' in receipt['url'] and not body.startswith(b'%PDF-'):
                    outcome['native_PDF_not_returned'] = True
                if re.search(rb'<html[\s>]', body[:5000], re.I):
                    outcome['inspection'] = inspect_html(body, receipt['url'], [ticker])
                elif '/history?' in receipt['url'] and receipt.get('status', receipt.get('http_status')) == 200:
                    rows, inactive = vendor_rows(body, ticker)
                    price_body = read_source(self.root, row['source'])
                    outcome['full_daily_quote_audit'] = full_daily_comparison(
                        price_body, ticker, rows, row['monthly_audit'])
                elif receipt['sha256'] == 'ada8e87e9809949688138f8a6eea142a84e1493d937432cc090deb0274c349be':
                    outcome['native_bonus_workbook_inspection'] = inspect_native_bonus_workbook(body)
            identity = {k: row[k] for k in ('ticker', 'cvm', 'share_class', 'isin')}
            row['reviewed_primary_actions'] = reviewed_primary_actions(
                self.root, identity, row['source_outcomes'])
            verified_dates = {r['ex_date'] for r in row['reviewed_primary_actions']}
            row['provider_action_dates_without_this_primary_proof'] = [r['date'] for r in
                row['monthly_audit']['provider_actions_not_primary'] if r['date'] not in verified_dates]
            row['native_listing_history_evidence'] = native_listing_readings(
                self.root, ticker, row['source_outcomes'])
            row['research_result'] = classified_outcome(row)
            row['evidence_recovered'] = any(v.get('receipt', {}).get('status',
                v.get('receipt', {}).get('http_status')) == 200
                and not v.get('inspection_or_acquisition_error')
                for v in row['source_outcomes'])
            row['model_admission'] = False
            records.append(row)
        require(all(digest((self.root / p).read_bytes()) == h
                    for p, h in self.protected.items()), 'protected_pointer_changed')
        report = {'scope': 'all39_selected_DY_positive_price_and_actions',
                  'baseline_snapshot': SNAPSHOT, 'manifest_sha256': self.manifest_hash,
                  'window_start': str(START), 'window_end': str(END),
                  'required_monthly_prices': 61, 'records': records,
                  'codes': len(records), 'codes_with_source_outcome': sum(
                      bool(r['source_outcomes']) for r in records),
                  'model_admissions': 0, 'gate_relaxations': 0,
                  'source_or_pointer_publications': 0, 'B3_network_requests': 0,
                  'prices_imputed': 0, 'factor_fitting': False,
                  'protected_pointer_hashes': self.protected,
                  'evidence_recovery_does_not_grant_model_admission': True}
        report['producer_sha256'] = digest(Path(__file__).read_bytes())
        report['tests_sha256'] = digest(Path(__file__).with_name('test_pending_all_prices.py').read_bytes())
        relative = self.save('report', report)
        return {'file': relative, 'codes': len(records),
                'codes_with_source_outcome': report['codes_with_source_outcome'],
                'monthly_counts': {r['ticker']: r['monthly_audit']['observed_required_months']
                                   for r in records}, 'model_admissions': 0}

    def summary(self, report_ref):
        report = json.loads((self.root / report_ref['file']).read_bytes())
        summary = {'detailed_report': str(self.root / report_ref['file']),
                   'codes': 39, 'model_admissions': 0, 'baseline_snapshot': SNAPSHOT,
                   'reusable_adapter': 'recommender.pending_all_prices.verified_native_history',
                   'discovery_adapter': 'recommender.pending_all_prices.published_history_request',
                   'rules': 'exact published issuer configuration, native closed, every date, no imputation',
                   'records': []}
        for row in report['records']:
            summary['records'].append({
                **{k: row[k] for k in ('ticker', 'cvm', 'share_class', 'isin', 'research_result')},
                'monthly_prices': row['monthly_audit']['observed_required_months'],
                'missing_months': row['monthly_audit']['missing_months'],
                'native_primary_actions': [{'date': r['ex_date'], 'ratio': r.get('native_ratio'),
                                           'sha256': r['source']['sha256']}
                                          for r in row['reviewed_primary_actions']],
                'provider_action_dates_without_this_primary_proof': row[
                    'provider_action_dates_without_this_primary_proof'],
                'native_listing_history_evidence': row['native_listing_history_evidence'],
                'outcomes': [{'url': o['url'], 'outcome_file': o['outcome_file'],
                    'status': o.get('receipt', {}).get('status', o.get('receipt', {}).get('http_status')),
                    'error': o.get('inspection_or_acquisition_error') or o.get('receipt', {}).get('error'),
                    'source_file': o.get('receipt', {}).get('file'),
                    'sha256': o.get('receipt', {}).get('sha256'),
                    'native_PDF_not_returned': o.get('native_PDF_not_returned', False)}
                    for o in row['source_outcomes']],
                'full_daily_quote_audits': [{'independent_exact_dates': o['full_daily_quote_audit'][
                    'independent_exact_dates'], 'trial': o['full_daily_quote_audit'][
                    'unchanged_price_gate_trial_using_every_native_date'],
                    'provider_action_adjusted_disagreement_dates_diagnostic_only': [v['date'] for v in
                    o['full_daily_quote_audit']['provider_action_adjusted_disagreements_diagnostic_only']]}
                    for o in row['source_outcomes'] if 'full_daily_quote_audit' in o]})
        body = encoded(summary)
        path = self.manifest_path.parent / ('prices-summary-' + digest(body) + '.json')
        if not path.exists():
            with path.open('xb') as stream:
                stream.write(body)
        return str(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', default='recommender/data/market')
    parser.add_argument('--manifest', default='../work/pending_recheck_20261025/all-157-manifest.json')
    parser.add_argument('--interactive', action='store_true')
    parser.add_argument('--offline', action='store_true')
    parser.add_argument('--summary', action='store_true')
    args = parser.parse_args()
    session = PriceResearch(args.root, args.manifest, acquire=not args.offline)
    if not args.interactive:
        report = session.report()
        if args.summary:
            report['summary'] = session.summary(report)
        print(json.dumps(report))
        return
    print(json.dumps({'ready': True, 'extra_stops': sorted(session.extra_stops),
                      'prior_stops': sorted(session.collector.prior_stops)}), flush=True)
    for line in sys.stdin:
        try:
            command = json.loads(line)
            action = command.get('action', 'fetch')
            if action == 'quit':
                break
            if action == 'report':
                output = session.report()
            elif action == 'widget':
                page = command['page']
                bundle = command['bundle']
                url, proof = published_history_request(read_source(session.root, page), page,
                    read_source(session.root, bundle), bundle, command['ticker'])
                output = session.fetch([command['ticker']], url, proof)
            else:
                output = session.fetch(command['tickers'], command['url'], command.get('discovery'))
            # The immutable outcome retains full inspection, output is bounded.
            if 'inspection' in output:
                output['inspection']['semantic_excerpts'] = output['inspection']['semantic_excerpts'][:12]
            print(json.dumps(output, ensure_ascii=True), flush=True)
        except Exception as error:
            print(json.dumps({'command_error': type(error).__name__, 'detail': str(error)}), flush=True)


if __name__ == '__main__':
    main()
