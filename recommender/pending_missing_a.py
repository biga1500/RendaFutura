"""Independent, bounded primary-source recovery for the missing_sources_a group.

Outputs are research evidence only. No inventory, model, pointer, or admission
state is changed. Literal values, unresolved dates, classes, and duplicates stay
visible for the caller's reconciliation.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import math
from pathlib import Path
import re
import unicodedata
import zipfile
from urllib.parse import urljoin, urlsplit

from lxml import html

from .bulk_pending_sources import PublicSourceCollector, approved


SCOPE = 'missing-a'
GROUP = 'missing_sources_a'
SEEDS = {
    25330: 'https://ri.alliedbrasil.com.br/informacoes-financeiras-e-operacionais/historico-de-proventos/',
    1570: 'https://ri.excelsior.com.br/',
    24627: 'https://ri.blau.com/informacoes-financeiras/dividendos-e-jcp/',
    1309: 'https://ri.bancomercantil.com.br/',
    25500: 'https://ri.bemobi.com.br/nossas-acoes/dividendos/',
    4723: 'https://lindenberg.com.br/ri/cvm/aviso-aos-acionistas/',
    24848: 'https://ri.cea.com.br/',
    16616: 'https://www.naturgy.com.br/corporativo/relacao_com_investidores/',
    4537: 'https://www.grazziotin.com.br/documentos-de-governanca/fatos-relevantes-e-comunicados/',
    2461: 'https://ri.celesc.com.br/servicos-aos-investidores/dividendos/',
    25585: 'https://ri.csnmineracao.com.br/informacoes-financeiras/historico-de-proventos/',
    25526: 'https://www.cruzeirodosuleducacional.edu.br/',
    25100: 'https://ri.cury.net/historico-de-proventos/',
    23795: 'https://www.ri.caixaseguridade.com.br/informacoes-aos-investidores/historico-de-dividendos-2/',
    26026: 'https://desktopsustentabilidade.com.br/compliance',
    25259: 'https://www.enjoei.com.br/',
    5770: 'https://www.eucatex.com.br/ri',
    26050: 'https://ri.unifique.com.br/pt/dividendos-e-jcp/',
    18368: 'https://ri.ctgbr.com.br/',
    25186: 'https://ri.grupomateus.com.br/',
    25402: 'https://ri.hbrrealty.com.br/documentos-cvm/avisos-aos-acionistas/',
    25453: 'https://ri.intelbras.com.br/',
    13285: 'https://josapar.com.br/investidores/aviso-aos-acionistas/',
    22020: 'https://ri.jsl.com.br/governanca-corporativa/dividendos/',
    25062: 'https://ri.lavvi.com.br/central-de-analises/',
    24910: 'https://www.lwsa.tech/',
    21067: 'https://ri.mouradubeux.com.br/servicos-a-investidores/dividendos-2/',
    17450: 'https://ri.rumolog.com/informacoes-financeiras/dividendos/',
    9954: 'https://www.bancoalfa.com.br/',
    20125: 'https://www.saud3.com.br/informacoes-aos-acionistas/remuneracao-aos-acionistas-bradsaude/',
    25704: 'https://ri.boasafrasementes.com.br/',
    25950: 'https://ri.3tentos.com.br/',
    24716: 'https://ri.grupovamos.com.br/governanca-corporativa/dividendos/',
    25763: 'https://ri.vittia.com.br/pt',
    21440: 'https://www.veste.com/',
}

# Each binding below was checked against the rendered original page, including
# the document's distinct notice date. A notice date is not a publication date.
REVIEWED_PDFS = {
    '037189a3cdda0513c5a1a84582754d8dc1664dfb810fa1b91d030869bbc8df6b': {
        'cvm': 4723, 'document_date': '2022-04-19', 'payment_date': '2022-04-26',
        'record_date': '2022-04-04', 'approval_date': '2022-04-04',
        'payment_text': 'cr\u00e9ditos dispon\u00edveis em 26 de abril de 2022',
        'amounts': {'ON': '4,3198767599', 'PN': '4,3198767599'},
        'amount_text': 'R$4,3198767599 por a\u00e7\u00e3o preferencial e R$4,3198767599 por a\u00e7\u00e3o ordin\u00e1ria'},
    '6de079bf9c3bff2ed661230259b0712c74ce4e61626cb079453d0b6a27e70502': {
        'cvm': 13285, 'document_date': '2023-05-04', 'payment_date': '2023-06-15',
        'approval_date': '2023-04-27', 'payment_text': 'em parcela \u00fanica no dia 15/06/2023',
        'amounts': {'ON': '0,82150046685396', 'PN': '0,90365051354'},
        'amount_text': '0,82150046685396 por a\u00e7\u00e3o ordin\u00e1ria e R$ 0,90365051354 por a\u00e7\u00e3o preferencial'},
    '3e22ddc01eb2b13b803884cdcf60e104340c88e559ea737a4b83eac43c545866': {
        'cvm': 21067, 'document_date': '2025-05-21', 'payment_date': '2025-05-30',
        'record_date': '2025-05-21', 'ex_date': '2025-05-22', 'approval_date': '2025-04-24',
        'payment_text': 'pagamento dos dividendos adicionais ser\u00e1 realizado em 30 de maio de 2025',
        'amounts': {'ON': '0,59089087589'},
        'amount_text': 'R$ 0,59089087589 por a\u00e7\u00e3o ordin\u00e1ria'},
    '87a612e8064c242c52f893911e8d2f92f567f25e8d2bb8e93ee10696365deb97': {
        'cvm': 21067, 'document_date': '2026-03-06', 'payment_date': '2026-03-13',
        'record_date': '2025-12-30', 'approval_date': '2025-12-29',
        'payment_text': '13 de mar\u00e7o de 2026', 'aggregate_installments': [1, 2],
        'amounts': {'ON': '1,18367018554'},
        'amount_text': 'R$ 1,18367018554 por a\u00e7\u00e3o ordin\u00e1ria'},
}

QUOTE_PAGES = {
    'BLAU3': 'https://ri.blau.com/servicos-aos-investidores/cotacoes-e-graficos/',
    'BMOB3': 'https://ri.bemobi.com.br/nossas-acoes/cotacoes-e-graficos/',
    'CXSE3': 'https://www.ri.caixaseguridade.com.br/informacoes-aos-investidores/graficos-e-cotacoes/',
}

BEMOBI_MIXED_NOTICES = {
    '6766f21d0b65e5dc6ccbd12f9fe330ac58c1f033c67b8ed55125a97c1e193124': {
        'document_date': '2024-12-11', 'record_date': '2024-12-17',
        'ex_date': '2024-12-18', 'payment_date': '2025-01-07',
        'components': [('DIVIDENDO', '0,19857444', None, 'paragraph (a), page 1'),
                       ('JCP', '0,44527052', '0,37847994', 'paragraph (b), pages 1-2')]},
    '8175a70ec3d71168426c56522cfdeb81a0bec789a4ef085b313b54b20526af6a': {
        'document_date': '2025-12-08', 'record_date': '2025-12-11',
        'ex_date': '2025-12-12', 'payment_date': '2025-12-22',
        'components': [('JCP', '0,500122', None, 'paragraph (a), page 1'),
                       ('DIVIDENDO', '0,381616', None, 'paragraph (b), pages 1-2'),
                       ('DIVIDENDO', '0,708817', None, 'paragraph (c), page 2')]},
}


def norm(value):
    return ' '.join(unicodedata.normalize('NFKD', str(value)).encode(
        'ascii', 'ignore').decode().lower().split())


def exact_date(value):
    """A date must occupy the whole cell: deadlines and ranges stay unresolved."""
    text = str(value).strip()
    if not re.fullmatch(r'\d{1,2}/\d{1,2}/\d{4}', text):
        return None
    try:
        return datetime.strptime(text, '%d/%m/%Y').date().isoformat()
    except ValueError:
        return None


def native_decimal(value):
    """Return an exact decimal string from a Portuguese monetary cell."""
    text = str(value).strip().replace('\xa0', '').replace('R$', '').replace(' ', '')
    if not re.fullmatch(r'(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?', text):
        return None
    try:
        parsed = Decimal(text.replace('.', '').replace(',', '.'))
        return str(parsed) if parsed.is_finite() else None
    except InvalidOperation:
        return None


def _cells(row):
    return [' '.join(cell.text_content().split()) for cell in row.xpath('./th|./td')]


def expanded_rows(table):
    """Expand only explicit HTML spans, retaining original row locators."""
    pending, rows = {}, []
    for row in table.xpath('.//tr'):
        values, inherited = {}, []
        for column, (text, remaining) in list(pending.items()):
            values[column] = text
            inherited.append(column)
            if remaining == 1:
                del pending[column]
            else:
                pending[column] = (text, remaining - 1)
        column = 0
        for cell in row.xpath('./th|./td'):
            while column in values:
                column += 1
            text = ' '.join(cell.text_content().split())
            colspan = min(100, int(cell.get('colspan', '1')))
            rowspan = min(1000, int(cell.get('rowspan', '1')))
            for offset in range(colspan):
                values[column + offset] = text
                if rowspan > 1:
                    pending[column + offset] = (text, rowspan - 1)
            column += colspan
        if values:
            rows.append({'locator': row.getroottree().getpath(row),
                         'cells': [values.get(i, '') for i in range(max(values)+1)],
                         'explicit_rowspan_inherited_columns': inherited})
    return rows


def inspect_html(body, url):
    try:
        decoded = body.decode('utf-8')
        encoding = 'utf-8'
    except UnicodeDecodeError:
        decoded = body.decode('windows-1252', errors='replace')
        encoding = 'windows-1252'
    tree = html.fromstring(decoded, parser=html.HTMLParser(no_network=True))
    published_templates = [{'offset': match.start(), 'text': decoded[max(0, match.start()-160):match.end()+600]}
                           for match in list(re.finditer(
                               r'(?i)fmBase|fmId|stockinfoId|api\.mziq\.com|data-stock|MZFinance|STOCK_ID', decoded))[:30]]
    links = []
    for node in tree.xpath('//a[@href]|//script[@src]|//iframe[@src]'):
        raw = node.get('href') or node.get('src')
        target = urljoin(url, raw)
        if approved(target):
            links.append({'url': target, 'text': ' '.join(node.text_content().split())[:260],
                          'locator': tree.getroottree().getpath(node)})
    tables = []
    for index, table in enumerate(tree.xpath('//table'), 1):
        rows = expanded_rows(table)
        if rows:
            tables.append({'locator': f'(//table)[{index}]', 'rows': rows})
    for node in tree.xpath('//script|//style|//noscript'):
        node.drop_tree()
    text = ' '.join(tree.text_content().split())
    terms = re.compile(r'(?i)dividend\w*|juros sobre (?:o )?capital|bonifica\w*|desdobr\w*|grupamento|pagamento|por a[c\u00e7][a\u00e3]o')
    excerpts = [{'offset': match.start(), 'text': text[max(0, match.start()-140):match.end()+420]}
                for match in list(terms.finditer(text))[:60]]
    return {'kind': 'html', 'text_decoding': encoding, 'title': ' '.join(tree.xpath('//title/text()')),
            'links': links, 'tables': tables, 'semantic_excerpts': excerpts,
            'issuer_published_widget_templates': published_templates,
            'date_literals': sorted(set(re.findall(r'\b\d{1,2}/\d{1,2}/20\d{2}\b', text))),
            'visible_tickers': sorted(set(re.findall(r'\b[A-Z]{4}\d{1,2}\b', text))),
            'publication_date': None, 'publication_basis': 'not_inferred_from_retrieval_or_footer'}


def parse_history(body, receipt, identity):
    """Typed, header-bound observations. This never grants model admission.

    `identity` preserves the exact manifest CVM/class/ISIN; source binding is
    independent and explicitly unresolved when the table does not name it.
    """
    if hashlib.sha256(body).hexdigest() != receipt['sha256']:
        raise ValueError('cash_history_original_hash_mismatch')
    inspected = inspect_html(body, receipt['url'])
    output = []
    for table in inspected['tables']:
        header_index = None
        for index, row in enumerate(table['rows'][:5]):
            headers = [norm(cell) for cell in row['cells']]
            if any('pagamento' in h for h in headers) and any(
                    'acao' in h or '/acao' in h for h in headers):
                header_index = index
                break
        if header_index is None:
            continue
        headers = table['rows'][header_index]['cells']
        normalized = [norm(h) for h in headers]
        amount_columns = [i for i, h in enumerate(normalized)
                          if ('acao' in h and ('valor' in h or 'ordinaria' in h or 'dividendo' in h))]
        for row in table['rows'][header_index+1:]:
            cells = row['cells']
            if len(cells) != len(headers) or not any(exact_date(c) for c in cells):
                continue
            for amount_index in amount_columns:
                decimal = native_decimal(cells[amount_index])
                if decimal is None:
                    continue
                labelled = dict(zip(normalized, cells))
                dates = {'payment_date': None, 'record_date': None, 'ex_date': None,
                         'approval_date': None}
                payment_literal = None
                for header, cell in labelled.items():
                    if 'pagamento' in header:
                        payment_literal = cell
                        dates['payment_date'] = exact_date(cell)
                    elif 'aprovacao' in header or 'deliberacao' in header:
                        dates['approval_date'] = exact_date(cell)
                    elif ('data base' in header or 'composicao aciona' in header
                          or 'posicao aciona' in header):
                        dates['record_date'] = exact_date(cell)
                    elif header.startswith('ex') or 'data ex' in header:
                        dates['ex_date'] = exact_date(cell)
                event_literals = [cells[i] for i, h in enumerate(normalized)
                                  if h in ('provento', 'proventos', 'natureza', 'tipo', 'evento',
                                           'forma de distribuicao', 'beneficio')]
                event_text = norm(' '.join(event_literals))
                has_div = 'dividend' in event_text
                has_jcp = 'jcp' in event_text or 'capital proprio' in event_text
                event = ('DIVIDENDO_E_JCP_NAO_SEGREGADOS' if has_div and has_jcp else
                         'JCP' if has_jcp else 'DIVIDENDO' if has_div else None)
                event_basis = 'explicit_row_event_label'
                if (event is None and identity['ticker'] == 'CXSE3'
                        and receipt['url'] == SEEDS[23795]
                        and any('total de dividendos' in h for h in normalized)):
                    event = 'DIVIDENDO'
                    event_basis = 'explicit_dividend_history_page_and_total_de_dividendos_header'
                source_class = None
                source_label = ' '.join([headers[amount_index], *[
                    cells[i] for i, h in enumerate(normalized) if 'classe' in h]])
                if re.search(r'\bON\b|Ordin[a\u00e1]ria', source_label, re.I):
                    source_class = 'ON'
                elif re.search(r'\bPN\b|Preferencial', source_label, re.I):
                    source_class = 'PN'
                same = [old for old in output if old['literal_cells'] == cells
                        and old['amount_header'] == headers[amount_index]]
                output.append({'identity': {k: identity.get(k) for k in ('ticker', 'cvm', 'share_class', 'isin')},
                               'source_class': source_class,
                               'exact_class_confirmed': source_class == identity.get('share_class'),
                               'event_type': event, 'event_literal': event_literals,
                               'event_type_basis': event_basis,
                               'amount_decimal': decimal, 'amount': float(Decimal(decimal)),
                               'amount_literal': cells[amount_index], 'amount_header': headers[amount_index],
                               'unit': 'BRL/share', 'gross_basis': 'explicit_gross' if 'brut' in normalized[amount_index]
                               else 'explicit_net' if 'liquid' in normalized[amount_index]
                               else 'not_explicit_for_per_share_column',
                               **dates, 'payment_literal': payment_literal,
                               'payment_exact': dates['payment_date'] is not None,
                               'source_duplicate_prior_count': len(same),
                               'literal_headers': headers, 'literal_cells': cells,
                               'explicit_rowspan_inherited_columns': row.get('explicit_rowspan_inherited_columns', []),
                               'source_url': receipt['url'], 'source_file': receipt['file'],
                               'source_sha256': receipt['sha256'], 'locator': row['locator'],
                               'retrieved_at': receipt.get('retrieved_at'), 'published_at': None,
                               'model_admission': False, 'bank_settlement_verified': False})
    return output


def inspect_body(body, receipt):
    if body.lstrip().startswith(b'%PDF'):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(body))
        pages = []
        for index, page in enumerate(reader.pages[:200], 1):
            text = page.extract_text() or ''
            if re.search(r'(?i)dividend|juros sobre|bonifica|desdobr|grupamento|por a[c\u00e7][a\u00e3]o', text):
                pages.append({'page': index, 'text': text})
        return {'kind': 'pdf', 'page_count': len(reader.pages), 'semantic_pages': pages,
                'published_at': None, 'publication_basis': 'not_inferred_from_pdf_metadata'}
    if body.startswith(b'PK\x03\x04'):
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            if (sum(item.file_size for item in archive.infolist()) > 32_000_000
                    or len(archive.infolist()) > 2000):
                raise ValueError('bounded_workbook_archive_limit')
            if 'xl/workbook.xml' not in archive.namelist():
                return {'kind': 'zip', 'members': archive.namelist(), 'numeric_values_interpreted': False}
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
        sheets = []
        for sheet in workbook:
            cells = []
            for row in sheet.iter_rows(max_row=min(sheet.max_row, 2000), max_col=min(sheet.max_column, 50)):
                for cell in row:
                    if cell.value is not None:
                        value = cell.value
                        cells.append({'locator': f"'{sheet.title}'!{cell.coordinate}",
                                      'value': value.isoformat() if isinstance(value, datetime) else value,
                                      'native_type': cell.data_type, 'number_format': cell.number_format,
                                      'formula_not_evaluated': cell.data_type == 'f'})
            sheets.append({'name': sheet.title, 'max_row': sheet.max_row, 'max_column': sheet.max_column,
                           'cells': cells})
        workbook.close()
        return {'kind': 'xlsx', 'sheets': sheets, 'published_at': None,
                'publication_basis': 'HTTP_last_modified_is_not_publication_date',
                'formulas_evaluated': False, 'source_values_changed': False}
    if body.lstrip().startswith((b'{', b'[')):
        return {'kind': 'json', 'value': json.loads(body)}
    return inspect_html(body, receipt['url'])


def parse_reviewed_pdf(body, receipt, identity):
    digest = hashlib.sha256(body).hexdigest()
    if digest != receipt['sha256']:
        raise ValueError('reviewed_pdf_original_hash_mismatch')
    reviewed = REVIEWED_PDFS.get(digest)
    if not reviewed or reviewed['cvm'] != identity['cvm']:
        return []
    share_class = identity['share_class']
    if share_class not in reviewed['amounts']:
        return []
    inspection = inspect_body(body, receipt)
    first = next((page['text'] for page in inspection['semantic_pages'] if page['page'] == 1), '')
    normalized = norm(first)
    # PDF text extraction can split words, so the rendered review is bound to
    # both an immutable digest and the exact native amount, not a loose match.
    literal = reviewed['amounts'][share_class]
    if literal not in first or norm(reviewed['payment_text']) not in normalized:
        raise ValueError('reviewed_pdf_semantic_locator_changed')
    decimal = native_decimal(literal)
    return [{'identity': {key: identity.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')},
             'source_class': share_class, 'exact_class_confirmed': True,
             'event_type': 'DIVIDENDO', 'event_literal': ['dividendos'],
             'amount_decimal': decimal, 'amount': float(Decimal(decimal)), 'amount_literal': literal,
             'unit': 'BRL/share', 'gross_basis': 'dividend_notice_native_amount_no_tax_transformation',
             **{key: reviewed.get(key) for key in ('payment_date', 'record_date', 'ex_date', 'approval_date')},
             'payment_exact': True, 'payment_literal': reviewed['payment_text'],
             'document_date': reviewed['document_date'], 'published_at': None,
             'aggregate_installments': reviewed.get('aggregate_installments'),
             'source_url': receipt['url'], 'source_file': receipt['file'], 'source_sha256': digest,
             'retrieved_at': receipt.get('retrieved_at'), 'locator': 'PDF page 1; cash notice paragraphs',
             'semantic_citations': [reviewed['amount_text'], reviewed['payment_text']],
             'verification': 'rendered_original_page_and_exact_sha256',
             'model_admission': False, 'bank_settlement_verified': False}]


def parse_bemobi_mixed_notice(body, receipt, identity):
    """Return five separately declared native gross components across two PDFs.

    The original mixed table observations remain in research output and must
    not be added again to these overlapping components.
    """
    expected_identity = {'ticker': 'BMOB3', 'cvm': 25500, 'share_class': 'ON', 'isin': 'BRBMOBACNOR7'}
    if any(identity.get(key) != value for key, value in expected_identity.items()):
        return []
    digest = hashlib.sha256(body).hexdigest()
    if digest != receipt['sha256']:
        raise ValueError('mixed_notice_original_hash_mismatch')
    notice = BEMOBI_MIXED_NOTICES.get(digest)
    if notice is None:
        return []
    inspection = inspect_body(body, receipt)
    text = ' '.join(page['text'] for page in inspection['semantic_pages'])
    if not re.search(r'\bticker:\s*bmob3\b', norm(text)):
        raise ValueError('mixed_notice_exact_security_code_missing')
    output = []
    for kind, literal, net, locator in notice['components']:
        if literal not in text or net is not None and net not in text:
            raise ValueError('mixed_notice_native_component_literal_missing')
        decimal = native_decimal(literal)
        output.append({'identity': {key: identity[key] for key in ('ticker', 'cvm', 'share_class', 'isin')},
                       'source_class': 'ON', 'exact_class_confirmed': True,
                       'class_binding_basis': 'exact_BMOB3_in_primary_notice_to_manifest_registered_ON_identity',
                       'event_type': kind, 'amount_literal': literal, 'amount_decimal': decimal,
                       'amount': float(Decimal(decimal)), 'unit': 'BRL/share',
                       'gross_basis': 'explicit_notice_gross_component',
                       'reported_net_amount_decimal': native_decimal(net) if net else None,
                       **{key: notice[key] for key in ('record_date', 'ex_date', 'payment_date', 'document_date')},
                       'approval_date': notice['document_date'], 'payment_exact': True, 'published_at': None,
                       'declared_amount_may_change_for_treasury_revisions': True,
                       'reconciliation_group': 'BMOB3:' + notice['payment_date'],
                       'overlaps_mixed_history_table_do_not_sum_both': True,
                       'source_url': receipt['url'], 'source_file': receipt['file'], 'source_sha256': digest,
                       'retrieved_at': receipt.get('retrieved_at'), 'locator': locator,
                       'verification': 'both_rendered_original_pages_and_exact_sha256',
                       'model_admission': False, 'bank_settlement_verified': False})
    return output


def published_history_request(page_body, page_receipt, bundle_body, bundle_receipt, ticker):
    """Construct the issuer's own public history request from actual JS proof.

    The stock ID and ticker must occur in this issuer's page, and this exact
    bundle URL must be a script published there. No issuer identifier is reused.
    """
    if page_receipt['url'] != QUOTE_PAGES[ticker]:
        raise ValueError('wrong_exact_issuer_quote_page')
    if any(hashlib.sha256(body).hexdigest() != receipt['sha256'] for body, receipt in (
            (page_body, page_receipt), (bundle_body, bundle_receipt))):
        raise ValueError('published_widget_proof_hash_mismatch')
    inspected = inspect_html(page_body, page_receipt['url'])
    if bundle_receipt['url'] not in {link['url'] for link in inspected['links']}:
        raise ValueError('bundle_not_published_by_this_issuer')
    page = page_body.decode('utf-8')
    identifiers = set(re.findall(r'''(?:var|const|let)\s+stockInfoId\s*=\s*['"]([a-f0-9-]{36})['"]''', page))
    identifiers.update(re.findall(r'''stockinfoId:\s*['"]([a-f0-9-]{36})['"]''', page))
    tickers = set(re.findall(r'''tickers:\s*['"]([A-Z0-9,]+)['"]''', page))
    script = bundle_body.decode('utf-8')
    if len(identifiers) != 1 or ticker not in tickers:
        raise ValueError('issuer_stock_id_or_exact_ticker_not_explicit')
    required = ('API_STOCKINFO_URL:"https://api.mziq.com/mzstockinfo"',
                't.API_URL+"/"+t.STOCK_ID+"/history?from="',
                '"&tickers="+t.TICKER', '"&adjusted="+t.ADJUSTED', 't.ADJUSTED=!1',
                'case"btn-period-no-provents":t.ADJUSTED=!1')
    if not all(token in script for token in required):
        raise ValueError('unreviewed_public_history_library_template')
    stock_id = next(iter(identifiers))
    url = (f'https://api.mziq.com/mzstockinfo/{stock_id}/history'
           f'?from=2021-09-01&to=2026-09-30&tickers={ticker}&adjusted=false')
    return url, {'kind': 'issuer_published_public_library_GET_template',
                 'source_url': page_receipt['url'], 'source_file': page_receipt['file'],
                 'source_sha256': page_receipt['sha256'], 'stockinfoId': stock_id, 'ticker': ticker,
                 'bundle_url': bundle_receipt['url'], 'bundle_file': bundle_receipt['file'],
                 'bundle_sha256': bundle_receipt['sha256'], 'adjusted': False,
                 'adjustment_basis': 'actual_widget_no_provents_control_and_request_template',
                 'no_Celesc_or_other_issuer_identifier_reuse': True}


def parse_native_history(body, receipt, identity):
    """Return native daily observations and 61 actual last-observation anchors.

    No adjustment, filling, renaming an adjusted close, or fabricated trading
    date is allowed. A monthly anchor retains the source's actual trade date.
    """
    if hashlib.sha256(body).hexdigest() != receipt['sha256']:
        raise ValueError('native_history_original_hash_mismatch')
    proof = receipt.get('discovery', {})
    if proof.get('ticker') != identity['ticker'] or proof.get('adjusted') is not False:
        raise ValueError('native_history_missing_exact_issuer_and_adjustment_proof')
    payload = json.loads(body)
    if payload.get('success') is not True:
        raise ValueError('native_history_provider_not_success')
    data = payload['data']
    locator_prefix = '$.data'
    if '/graphic?' in receipt['url']:
        data = data['history']
        locator_prefix += '.history'
    matches = [(i, row) for i, row in enumerate(data['tickers']) if row['ticker'] == identity['ticker']]
    if len(matches) != 1 or data.get('errors'):
        raise ValueError('native_history_ambiguous_ticker_or_provider_errors')
    ticker_index, series = matches[0]
    daily, seen, months = [], set(), {}
    for index, row in enumerate(series['history']):
        day, close = row['date'], row['closed']
        if (not re.fullmatch(r'\d{4}-\d{2}-\d{2}', day)
                or datetime.strptime(day, '%Y-%m-%d').date().isoformat() != day
                or day in seen or not isinstance(close, (float, int)) or isinstance(close, bool)
                or not math.isfinite(close) or close <= 0):
            raise ValueError('native_history_invalid_or_duplicate_observation')
        seen.add(day)
        item = {'date': day, 'close': close, 'native_row': row,
                'locator': f'{locator_prefix}.tickers[{ticker_index}].history[{index}].closed'}
        daily.append(item)
        month = day[:7]
        if month not in months or day > months[month]['date']:
            months[month] = item
    expected = [f'{serial // 12:04d}-{serial % 12 + 1:02d}'
                for serial in range(2021 * 12 + 8, 2026 * 12 + 9)]
    anchors = [{'month': month, 'trade_date': months[month]['date'], 'close': months[month]['close'],
                'locator': months[month]['locator']} for month in expected if month in months]
    return {'identity': {key: identity.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')},
            'provider': data.get('source'), 'provider_link': data.get('link'),
            'daily_observations': daily, 'daily_count': len(daily), 'monthly_anchors': anchors,
            'missing_months': [month for month in expected if month not in months],
            '61_consecutive_actual_monthly_anchors': len(anchors) == 61,
            'first_trade_date': min(seen) if seen else None, 'last_trade_date': max(seen) if seen else None,
            'source_url': receipt['url'], 'source_file': receipt['file'], 'source_sha256': receipt['sha256'],
            'retrieved_at': receipt.get('retrieved_at'), 'published_at': None,
            'adjusted_request': False, 'native_values_changed': False, 'imputed_observations': 0,
            'model_admission': False,
            'share_action_reconciliation': 'requires_comparison_with_issuer_actions_and_existing_cache'}


def acquire_priority_quotes(collector, selected, prior):
    root = collector.root
    base = collector.base
    results = []

    def obtain(url, discovery):
        old = next((item for item in prior.get(urlsplit(url).hostname, [])
                    if item['url'] == url and not item.get('error')), None)
        receipt = old or collector.fetch(url, discovery=discovery)
        if receipt.get('status', receipt.get('http_status')) != 200 or receipt.get('error'):
            raise ValueError('public_quote_source_not_success:' + str(receipt.get('error')))
        return receipt, collector.body(receipt)

    for ticker, page_url in QUOTE_PAGES.items():
        identity = next(row for row in selected if row['ticker'] == ticker)
        report_path = base / f'issuer-{identity["cvm"]}-inspection.json'
        issuer = json.loads(report_path.read_bytes())
        result = {'identity': {key: identity.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')},
                  'page_url': page_url, 'model_admission': False, 'outcomes': []}
        try:
            page_receipt = next((doc['receipt'] for doc in issuer['documents'] if doc['receipt']['url'] == page_url), None)
            if page_receipt is None:
                source = next(doc for doc in issuer['documents']
                              if any(link['url'] == page_url for link in doc['inspection'].get('links', [])))
                page_receipt, page_body = obtain(page_url, {'kind': 'issuer_published_link',
                    'source_url': source['receipt']['url'], 'source_file': source['receipt']['file'],
                    'source_sha256': source['receipt']['sha256']})
            else:
                page_body = collector.body(page_receipt)
            result['page_receipt'] = page_receipt
            inspected = inspect_html(page_body, page_url)
            bundle_url = next(link['url'] for link in inspected['links']
                              if re.fullmatch(r'https://cdn\.mziq\.com/stockinfo-viewer/main/latest/js/main\.js(?:\?v=1)?', link['url']))
            bundle_receipt, bundle_body = obtain(bundle_url, {'kind': 'issuer_published_script',
                'source_url': page_url, 'source_file': page_receipt['file'], 'source_sha256': page_receipt['sha256']})
            result['bundle_receipt'] = bundle_receipt
            url, proof = published_history_request(page_body, page_receipt, bundle_body, bundle_receipt, ticker)
            result['request_proof'] = proof
            receipt = collector.fetch(url, discovery=proof)
            result['history_receipt'] = receipt
            if receipt['status'] is None and receipt['error'] == 'TimeoutError':
                # The published chart is a distinct public GET endpoint. Keep
                # the history failure and make at most one chart acquisition.
                template = 'p=t.API_URL+"/"+t.STOCK_ID+"/graphic?from="+r+"&to="+n+"&tickers="+t.TICKER+"&adjusted="+t.ADJUSTED'
                if template in bundle_body.decode('utf-8'):
                    graphic_url = url.replace('/history?', '/graphic?')
                    chart = collector.fetch(graphic_url, discovery={**proof,
                        'kind': 'issuer_published_distinct_graphic_GET_template',
                        'history_failure_retained': True})
                    result['graphic_receipt'] = chart
                    result['outcomes'].append({'url': graphic_url, 'status': chart['status'], 'error': chart['error']})
                    if chart['status'] == 200 and not chart['error']:
                        parsed = parse_native_history(collector.body(chart), chart, identity)
                        result['native_prices'] = {k: value for k, value in parsed.items() if k != 'daily_observations'}
            if receipt['status'] == 200 and not receipt['error']:
                payload = json.loads(collector.body(receipt))
                parsed = parse_native_history(collector.body(receipt), receipt, identity)
                result['native_prices'] = {k: value for k, value in parsed.items() if k != 'daily_observations'}
                result['payload_shape'] = ({'type': 'list', 'length': len(payload), 'first': payload[:2]}
                                           if isinstance(payload, list) else
                                           {'type': type(payload).__name__, 'keys': list(payload)[:20],
                                            'sample': str(payload)[:1600]})
            result['outcomes'].append({'url': url, 'status': receipt['status'], 'error': receipt['error']})
        except (ValueError, OSError, KeyError, StopIteration) as failure:
            result['outcomes'].append({'error': type(failure).__name__ + ': ' + str(failure)})
        (base / f'quote-{ticker}.json').write_text(json.dumps(result, ensure_ascii=True, indent=2), encoding='utf-8')
        results.append(result)
        print(json.dumps({'ticker': ticker, 'outcomes': result['outcomes'],
                          'shape': result.get('payload_shape')}), flush=True)
    return results


def acquire_mixed_bemobi_notices(collector):
    path = collector.base / 'issuer-25500-inspection.json'
    issuer = json.loads(path.read_bytes())
    source = next(doc for doc in issuer['documents'] if doc['receipt']['url'] == SEEDS[25500])
    for row_index in (4, 7):
        link = next(link for link in source['inspection']['links']
                    if link['locator'].endswith(f'/table/tbody/tr[{row_index}]/td[8]/a'))
        try:
            receipt = collector.fetch(link['url'], discovery={'kind': 'issuer_published_exact_mixed_event_notice',
                'source_url': source['receipt']['url'], 'source_file': source['receipt']['file'],
                'source_sha256': source['receipt']['sha256'], 'locator': link['locator']})
            issuer['outcomes'].append({'url': link['url'], 'outcome': 'mixed_event_notice_acquired',
                                      'status': receipt['status'], 'error': receipt['error'], 'receipt': receipt})
            if receipt['status'] == 200 and not receipt['error']:
                inspected = inspect_body(collector.body(receipt), receipt)
                if receipt['url'] not in {doc['receipt']['url'] for doc in issuer['documents']}:
                    issuer['documents'].append({'receipt': receipt, 'inspection': inspected})
        except ValueError as failure:
            issuer['outcomes'].append({'url': link['url'], 'outcome': str(failure)})
    path.write_text(json.dumps(issuer, ensure_ascii=True, indent=2), encoding='utf-8')


def local_metadata(root):
    """Inspect existing public receipts without contacting their origins."""
    sources, stops = defaultdict(list), set()
    for path in (root / 'raw/primary-research').rglob('*.json'):
        if path.stat().st_size > 2_000_000 or SCOPE in path.parts:
            continue
        try:
            value = json.loads(path.read_bytes())
        except (ValueError, OSError):
            continue
        records = [value] if isinstance(value, dict) else value if isinstance(value, list) else []
        if isinstance(value, dict):
            for key in ('responses', 'results', 'sources'):
                if isinstance(value.get(key), list):
                    records.extend(value[key])
        for item in records:
            if not isinstance(item, dict) or not isinstance(item.get('url'), str):
                continue
            host = urlsplit(item['url']).hostname
            reason = ' '.join(str(item.get(key, '')) for key in ('error', 'stop', 'stop_reason', 'reason'))
            status = item.get('status', item.get('http_status', item.get('httpStatus')))
            denied = (status in (401, 403, 429) or item.get('robots_allowed') is False
                    or re.search(r'(?i)captcha|challenge|access.denial|robots.*disallow|HTTP.?40[13]|HTTP.?429', reason))
            if denied:
                stops.add(host)
            if not denied and item.get('file') and item.get('sha256') and item.get('bytes') and status in (200, 'acquired', 'available'):
                sources[host].append({**item, 'prior_receipt_file': path.relative_to(root).as_posix()})
    return sources, stops


def _verified_local_body(root, receipt):
    root = Path(root).resolve()
    path = root / receipt['file']
    if not path.resolve().is_relative_to(root) or path.is_symlink():
        raise ValueError('unsafe_recovery_original_path')
    body = path.read_bytes()
    if len(body) != receipt['bytes'] or hashlib.sha256(body).hexdigest() != receipt['sha256']:
        raise ValueError('recovery_original_integrity_failure')
    return body


def load_cash_rows(market, identity):
    """Offline parent adapter: exact manifest identity -> unreconciled cash rows.

    Mixed HTML totals and PDF components intentionally both remain visible.
    Use `overlaps_mixed_history_table_do_not_sum_both` to avoid double counting.
    A class not named in the source remains unresolved rather than inferred.
    """
    root = Path(market).resolve()
    report = root / 'raw/primary-research/all-pending' / SCOPE / f'issuer-{int(identity["cvm"])}-inspection.json'
    issuer = json.loads(report.read_bytes())
    if identity['ticker'] not in issuer['tickers'] or identity['cvm'] != issuer['cvm']:
        raise ValueError('recovery_issuer_identity_mismatch')
    exact_identity = {key: identity.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')}
    if exact_identity not in issuer.get('identities', []):
        raise ValueError('recovery_manifest_class_or_ISIN_mismatch')
    rows = []
    for document in issuer['documents']:
        receipt = document['receipt']
        body = _verified_local_body(root, receipt)
        if document['inspection']['kind'] == 'html':
            rows.extend(parse_history(body, receipt, identity))
        elif document['inspection']['kind'] == 'pdf':
            rows.extend(parse_reviewed_pdf(body, receipt, identity))
            rows.extend(parse_bemobi_mixed_notice(body, receipt, identity))
    return rows


def load_native_prices(market, identity):
    """Offline parent adapter: exact identity -> daily prices and 61 anchors.

    Returns evidence for comparison, never a synthesized exchange response.
    """
    root = Path(market).resolve()
    report = root / 'raw/primary-research/all-pending' / SCOPE / f'quote-{identity["ticker"]}.json'
    quote = json.loads(report.read_bytes())
    if any(quote['identity'].get(key) != identity.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')):
        raise ValueError('native_price_exact_identity_mismatch')
    expected_url, _ = published_history_request(
        _verified_local_body(root, quote['page_receipt']), quote['page_receipt'],
        _verified_local_body(root, quote['bundle_receipt']), quote['bundle_receipt'], identity['ticker'])
    for key in ('history_receipt', 'graphic_receipt'):
        receipt = quote.get(key, {})
        if receipt.get('status') == 200 and not receipt.get('error'):
            allowed_url = expected_url if key == 'history_receipt' else expected_url.replace('/history?', '/graphic?')
            if receipt['url'] != allowed_url:
                raise ValueError('native_price_response_not_bound_to_actual_issuer_template')
            return parse_native_history(_verified_local_body(root, receipt), receipt, identity)
    raise ValueError('native_price_primary_response_unavailable')


def verify_recovered_evidence(market, selected):
    """Read-only integrity and exact-identity audit for every assigned code."""
    root = Path(market).resolve()
    base = root / 'raw/primary-research/all-pending' / SCOPE
    originals, outcomes, rows = set(), 0, 0
    priorities = []
    for identity in selected:
        issuer = json.loads((base / f'issuer-{identity["cvm"]}-inspection.json').read_bytes())
        assert issuer['cvm'] == identity['cvm'] and identity['ticker'] in issuer['tickers']
        assert issuer['outcomes']
        outcomes += len(issuer['outcomes'])
        for doc in issuer['documents']:
            _verified_local_body(root, doc['receipt'])
            originals.add(doc['receipt']['sha256'])
        cash = load_cash_rows(root, identity)
        assert all(row['model_admission'] is False for row in cash)
        assert all(row['identity'] == {key: identity.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')} for row in cash)
        rows += len(cash)
        if identity['ticker'] in QUOTE_PAGES:
            prices = load_native_prices(root, identity)
            assert prices['61_consecutive_actual_monthly_anchors'] and not prices['missing_months']
            valid = [row for row in cash if row.get('payment_date') and '2021-10-01' <= row['payment_date'] <= '2026-09-30']
            priorities.append({'identity': prices['identity'], 'price_source': prices['provider'],
                               'daily_count': prices['daily_count'], 'monthly_anchor_count': len(prices['monthly_anchors']),
                               'first_trade_date': prices['first_trade_date'], 'last_trade_date': prices['last_trade_date'],
                               'price_source_file': prices['source_file'], 'price_source_sha256': prices['source_sha256'],
                               'native_cash_observations': len(cash), 'exact_payment_rows_in_window': len(valid),
                               'earliest_exact_payment_in_window': min(row['payment_date'] for row in valid),
                               'latest_exact_payment_in_window': max(row['payment_date'] for row in valid),
                               'annual_accounts_available_in_manifest': identity['annual_accounts']['available'],
                               'cash_source_report': f'raw/primary-research/all-pending/{SCOPE}/issuer-{identity["cvm"]}-inspection.json',
                               'specific_remaining_reconciliation': {
                                   'BLAU3': ['per_share_gross_and_historical_class_binding_need_parent_verification',
                                             '2025_complementary_dividend_is_deadline_until_2028_not_exact_payment',
                                             'compare_native_prices_with_existing_actions_and_61_month_cache'],
                                   'BMOB3': ['two_original_mixed_HTML_totals_overlap_five_typed_PDF_components',
                                             '2025_HTML_1.59_is_rounded_relative_to_native_components',
                                             'PDF_values_explicitly_subject_to_treasury_revisions',
                                             'verify_2021_window_start_and_native_price_actions'],
                                   'CXSE3': ['cash_table_per_share_values_have_only_two_decimal_places',
                                             'source_does_not_explicitly_name_class_in_cash_table',
                                             'exact_native_cash_reconciliation_must_not_fill_precision']
                               }[identity['ticker']], 'model_admission': False})
    return {'assigned_codes_checked': len(selected), 'unique_issuer_original_hashes_verified': len(originals),
            'cash_observations_checked': rows, 'every_code_has_specific_outcomes': True,
            'priority_candidates': priorities, 'network_used_for_audit': False, 'model_admissions': 0}


def _score(link):
    text = norm(link['text'] + ' ' + link['url'])
    if re.search(r'/en/|privacy|privacidade|termos|politica|fatores-de-risco|mailing', text):
        return -1
    score = 0
    for pattern, points in ((r'dividend|proventos|remuneracao.*acion', 100),
                            (r'aviso.*acion|comunicados|fatos.relevantes', 65),
                            (r'cotaco|graficos|grafico-de-acoes', 45),
                            (r'central.de.resultados|demonstrac.*financeira|informacoes.financeiras', 30),
                            (r'(^|/)ri(/|$)|relac.*investidor|relacao.*investidor', 25)):
        if re.search(pattern, text):
            score += points
    if '.pdf' in text or 'download.aspx' in text or '/mzfilemanager/' in text:
        score += 20 if score else 0
    if '2021' in text or '2022' in text:
        score += 5
    if 'investidores' in text or re.search(r'https://ri\.[a-z]', text):
        score += 20
    return score


def recover_issuer(collector, identities, prior, page_budget=3, seed_override=None, seed_discovery=None):
    identity = identities[0]
    cvm = identity['cvm']
    seed = seed_override or SEEDS[cvm]
    result = {'cvm': cvm, 'tickers': [r['ticker'] for r in identities],
              'FCA_website': identity.get('website_declared_in_FCA'),
              'seed': seed, 'outcomes': [], 'documents': [], 'rows': [],
              'model_admissions': 0, 'historical_price_months_admitted': 0}
    queued = [(seed, seed_discovery or {'kind': 'FCA_website_or_fresh_official_web_discovery',
                                       'FCA_website': identity.get('website_declared_in_FCA')})]
    seen, page_count = set(), 0
    # An exact prior seed can avoid a fresh request. Other cache pages remain
    # inspectable even where the origin has since stopped allowing acquisition.
    candidates = prior.get(urlsplit(seed).hostname, [])
    cached = {r['url']: r for r in candidates}
    while queued and page_count < page_budget:
        url, discovery = queued.pop(0)
        if url in seen:
            continue
        seen.add(url)
        receipt = cached.get(url)
        try:
            if receipt is None:
                receipt = collector.fetch(url, discovery=discovery)
            body = collector.body(receipt)
        except (ValueError, OSError, KeyError) as failure:
            result['outcomes'].append({'url': url, 'outcome': str(failure), 'discovery': discovery})
            if url == seed:
                for old in candidates[:page_budget]:
                    queued.append((old['url'], {'kind': 'existing_public_cache_read_only'}))
            continue
        page_count += 1
        status = receipt.get('status', receipt.get('http_status'))
        result['outcomes'].append({'url': url, 'outcome': 'cached_inspected' if url in cached else 'fetched',
                                   'status': status, 'error': receipt.get('error'), 'receipt': receipt})
        if status not in (200, 'acquired') or receipt.get('error'):
            continue
        try:
            inspected = inspect_body(body, receipt)
        except Exception as failure:
            result['outcomes'].append({'url': url, 'outcome': 'inspection_failed',
                                       'error': type(failure).__name__ + ': ' + str(failure)[:160]})
            continue
        result['documents'].append({'receipt': receipt, 'inspection': inspected})
        if inspected['kind'] != 'html':
            continue
        for exact_identity in identities:
            result['rows'].extend(parse_history(body, receipt, exact_identity))
        new_links = [link for link in inspected['links'] if _score(link) > 0 and link['url'] not in seen]
        new_links.sort(key=_score, reverse=True)
        for link in new_links[:12]:
            # A vendor file is eligible only through an actual issuer href.
            queued.append((link['url'], {'kind': 'issuer_published_link', 'source_url': url,
                                        'source_file': receipt['file'], 'source_sha256': receipt['sha256'],
                                        'locator': link['locator'], 'link_text': link['text']}))
        queued.sort(key=lambda item: _score({'url': item[0], 'text': item[1].get('link_text', '')}), reverse=True)
    result['pages_inspected_or_attempted'] = page_count
    result['blockers'] = []
    if not result['documents']:
        result['blockers'].append('no_accessible_primary_original_in_bounded_branch')
    if not result['rows']:
        result['blockers'].append('no_header_bound_per_share_cash_rows_in_inspected_pages')
    if any(not row['exact_class_confirmed'] for row in result['rows']):
        result['blockers'].append('cash_table_class_binding_requires_exact_issuer_document')
    result['blockers'].extend(['complete_2021_2026_cash_and_actions_not_certified',
                               '61_month_native_price_independent_reconciliation_not_certified',
                               'source_publication_date_unknown_unless_explicitly_documented'])
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--manifest', type=Path, required=True)
    parser.add_argument('--market', type=Path, required=True)
    parser.add_argument('--report', type=Path, required=True)
    parser.add_argument('--tickers', default='')
    parser.add_argument('--pages', type=int, default=3)
    parser.add_argument('--reinspect', action='store_true')
    parser.add_argument('--supplement-plan', type=Path)
    parser.add_argument('--quotes', action='store_true')
    parser.add_argument('--mixed-notices', action='store_true')
    parser.add_argument('--audit', action='store_true')
    args = parser.parse_args()
    manifest = json.loads(args.manifest.read_bytes())
    selected = [r for r in manifest['records'] if r['task_group'] == GROUP]
    assert len(selected) == 41 and len({r['ticker'] for r in selected}) == 41
    wanted = set(args.tickers.split(',')) - {''}
    grouped = defaultdict(list)
    for row in selected:
        if not wanted or row['ticker'] in wanted:
            grouped[row['cvm']].append(row)
    root = args.market.resolve()
    collector = PublicSourceCollector(root, SCOPE)
    prior, additional_stops = local_metadata(root)
    collector.prior_stops.update(additional_stops)
    base = root / 'raw/primary-research/all-pending' / SCOPE
    results = []
    if args.quotes:
        acquire_priority_quotes(collector, selected, prior)
    if args.mixed_notices:
        acquire_mixed_bemobi_notices(collector)
    plan = json.loads(args.supplement_plan.read_bytes()) if args.supplement_plan else None
    with ThreadPoolExecutor(max_workers=3) as executor:
        tasks = {executor.submit(recover_issuer, collector, rows, prior, args.pages): cvm
                 for cvm, rows in grouped.items()} if not args.reinspect and plan is None and not args.quotes and not args.mixed_notices and not args.audit else {}
        if plan is not None:
            tasks = {executor.submit(recover_issuer, collector, grouped[item['cvm']], prior,
                                     item.get('pages', 1), item['url'], item['discovery']): item['cvm']
                     for item in plan}
        for future in as_completed(tasks):
            result = future.result()
            results.append(result)
            target = base / f'issuer-{result["cvm"]}-inspection.json'
            if plan is not None and target.exists():
                previous = json.loads(target.read_bytes())
                result['previous_seed'] = previous['seed']
                result['outcomes'] = previous['outcomes'] + result['outcomes']
                by_url = {doc['receipt']['url']: doc for doc in previous['documents'] + result['documents']}
                result['documents'] = list(by_url.values())
                result['rows'] = previous['rows'] + result['rows']
            target.write_text(json.dumps(result, ensure_ascii=True, indent=2), encoding='utf-8')
            print(json.dumps({'cvm': result['cvm'], 'tickers': result['tickers'],
                              'documents': len(result['documents']), 'rows': len(result['rows']),
                              'outcomes': [{k: o.get(k) for k in ('url', 'outcome', 'status', 'error')}
                                           for o in result['outcomes']]}), flush=True)
    all_results = [json.loads(p.read_bytes()) for p in sorted(base.glob('issuer-*-inspection.json'))]
    if args.reinspect:
        for result in all_results:
            result['identities'] = [{key: row.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')}
                                    for row in selected if row['cvm'] == result['cvm']]
            for outcome in result['outcomes']:
                if 'robots' in outcome.get('outcome', ''):
                    robots_url = 'https://' + urlsplit(outcome['url']).hostname + '/robots.txt'
                    robots_path = base / ('response-' + hashlib.sha256(robots_url.encode()).hexdigest() + '.json')
                    if robots_path.exists():
                        outcome['robots_receipt'] = json.loads(robots_path.read_bytes())
            result['rows'] = []
            for document in result['documents']:
                body = collector.body(document['receipt'])
                document['inspection'] = inspect_body(body, document['receipt'])
                if document['inspection']['kind'] == 'html':
                    for identity in selected:
                        if identity['cvm'] == result['cvm']:
                            result['rows'].extend(parse_history(body, document['receipt'], identity))
                elif document['inspection']['kind'] == 'pdf':
                    for identity in selected:
                        if identity['cvm'] == result['cvm']:
                            result['rows'].extend(parse_reviewed_pdf(body, document['receipt'], identity))
                            result['rows'].extend(parse_bemobi_mixed_notice(body, document['receipt'], identity))
            result['blockers'] = [b for b in result['blockers'] if b not in (
                'no_header_bound_per_share_cash_rows_in_inspected_pages',
                'cash_table_class_binding_requires_exact_issuer_document')]
            if not result['rows']:
                result['blockers'].append('no_header_bound_per_share_cash_rows_in_inspected_pages')
            if any(not row['exact_class_confirmed'] for row in result['rows']):
                result['blockers'].append('cash_table_class_binding_requires_exact_issuer_document')
            (base / f'issuer-{result["cvm"]}-inspection.json').write_text(
                json.dumps(result, ensure_ascii=True, indent=2), encoding='utf-8')
    coverage = []
    for row in selected:
        match = next((r for r in all_results if r['cvm'] == row['cvm']), None)
        coverage.append({'identity': {key: row.get(key) for key in ('ticker', 'cvm', 'share_class', 'isin')},
                         'inspected': match is not None,
                         'documents': len(match['documents']) if match else 0,
                         'interpretable_rows': sum(r['identity']['ticker'] == row['ticker'] for r in match['rows']) if match else 0,
                         'outcomes': match['outcomes'] if match else [],
                         'blockers': match['blockers'] if match else ['issuer_not_processed'],
                         'model_admission': False})
    report = {'scope': SCOPE, 'snapshot': manifest['snapshot'], 'reference': manifest['reference'],
              'created_at': datetime.now(timezone.utc).isoformat(),
              'manifest_sha256': hashlib.sha256(args.manifest.read_bytes()).hexdigest(),
              'coverage_count': sum(r['inspected'] for r in coverage), 'expected_count': 41,
              'issuers_processed': len(all_results), 'coverage': coverage,
              'totals': {'documents': sum(len(r['documents']) for r in all_results),
                         'interpretable_rows': sum(len(r['rows']) for r in all_results), 'model_admissions': 0},
              'issuer_reports': [p.relative_to(root).as_posix() for p in sorted(base.glob('issuer-*-inspection.json'))],
              'quote_reports': [p.relative_to(root).as_posix() for p in sorted(base.glob('quote-*.json'))],
              'adapter_contracts': {
                  'load_cash_rows(market, identity)': 'offline list of native observations; exact Decimal strings, typed dates, source class separate from manifest class, originals verified; mixed totals overlap PDF components',
                  'load_native_prices(market, identity)': 'offline dict with native daily observations, 61 last-actual-observation monthly anchors, exact source ticker, issuer-published adjusted=false proof and Enfoque origin',
                  'parse_history(body, receipt, identity)': 'header-bound HTML observations; original hash required; no payment inferred from deadline or ex-date',
                  'parse_reviewed_pdf(body, receipt, identity)': 'SHA-bound visually inspected CALI/JOPA/MDNE native cash notices',
                  'parse_bemobi_mixed_notice(body, receipt, identity)': 'SHA-bound primary BMOB3 notices; separately declared gross JCP/dividend components with record/ex/payment dates',
                  'published_history_request(page_body, page_receipt, bundle_body, bundle_receipt, ticker)': 'verifies actual issuer ID/ticker and exact published library before constructing public GET; no identifier borrowed',
                  'verify_recovered_evidence(market, selected)': 'read-only all-41 original-byte/hash and exact-identity audit with three concrete parent integration candidates'
              },
              'guardrails': {'no_B3_network': True, 'no_imputation': True, 'no_gate_relaxation': True,
                            'no_model_publication': True, 'no_old_filing_financial_fallback': True,
                            'original_retrieval_and_unknown_publication_distinct': True}}
    if args.audit:
        report['integrity_audit'] = verify_recovered_evidence(root, selected)
        report['priority_candidates'] = report['integrity_audit']['priority_candidates']
    args.report.write_text(json.dumps(report, ensure_ascii=True, indent=2), encoding='utf-8')
    print(json.dumps({'report': str(args.report), 'coverage': report['coverage_count'], 'totals': report['totals']}))


if __name__ == '__main__':
    main()
