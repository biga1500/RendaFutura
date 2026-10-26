"""Strict, offline import of manually obtained histories; no collector side effects.

Interface: load_import_bundle(folder, cutoff) returns ImportBundle. Dates are ISO
strings, amounts are finite floats, identities/ratios are integers. Root tables
prices/payments/actions/fundamentals contain only historical rows at cutoff;
future_payments contains dated future events and is NEVER paid cash. No prices,
cash, missing accounts, share factors or point-in-time availability are inferred.
Provenance contains the validated source manifest, SHA-256 of every consumed file,
missing/empty tables and per-ticker counts/declared coverage. Coverage is a source
claim, not an independent completeness certification or model-readiness decision.

Input: source.json and optional <table>.csv / <table>.xlsx, or tables.xlsx with
named sheets. CSV headers/order are COLUMNS; delimiter and decimal_format are
explicit in source.json. ISO dates or native midnight Excel dates are accepted.
Numeric strings have no thousands separators/exponents; native Excel numbers
are accepted. Formulas, external workbook links and duplicate rows are rejected.
Adjusted prices/cash must declare their respective *_basis_date equal to cutoff.
Per-ticker coverage maps table names to {start, end, complete}; historical row
dates must lie within that range. Fundamental coverage uses published_at, payment
coverage uses payment_date. Missing coverage never establishes zero payments.

CLI: templates FOLDER [--format csv|xlsx]; validate FOLDER --cutoff YYYY-MM-DD
--workspace ROOT --output RELATIVE_OR_ABSOLUTE_PATH. Output must be within ROOT,
outside input. Frozen <sha256>/normalized.json and provenance.json are immutable
by this API: existing bytes must match. No network, credential or model imports.
"""

import argparse
from collections import Counter
import csv
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
from typing import Any, TypedDict
from urllib.parse import parse_qsl, urlsplit
import zipfile

import openpyxl


COLUMNS = {
    'prices': ('ticker', 'date', 'close'),
    'payments': ('ticker', 'record_date', 'payment_date', 'gross_per_share', 'label'),
    'actions': ('ticker', 'effective_date', 'numerator', 'denominator', 'type'),
    'fundamentals': ('ticker', 'cvm', 'share_class', 'fiscal_end', 'published_at',
                     'share_basis_date', 'eps', 'free_cash_per_share',
                     'liabilities_to_assets', 'sector_code', 'scope', 'source_ref'),
}
MAX_FILE_BYTES = 128 * 1024 * 1024
MAX_ROWS = 2_000_000
TICKER = re.compile(r'[A-Z][A-Z0-9-]{0,29}\Z')
SHARE_CLASSES = {'ON', 'PN', 'PNA', 'PNB', 'PNC', 'PND', 'UNIT'}


class PriceRow(TypedDict):
    ticker: str
    date: str
    close: float


class PaymentRow(TypedDict):
    ticker: str
    record_date: str
    payment_date: str
    gross_per_share: float
    label: str


class ActionRow(TypedDict):
    ticker: str
    effective_date: str
    numerator: int
    denominator: int
    type: str


class FundamentalRow(TypedDict):
    ticker: str
    cvm: int
    share_class: str
    fiscal_end: str
    published_at: str
    share_basis_date: str
    eps: float
    free_cash_per_share: float | None
    liabilities_to_assets: float
    sector_code: int
    scope: str
    source_ref: str


class ImportBundle(TypedDict):
    schema_version: int
    cutoff: str
    prices: list[PriceRow]
    payments: list[PaymentRow]
    future_payments: list[PaymentRow]
    actions: list[ActionRow]
    fundamentals: list[FundamentalRow]
    provenance: dict[str, Any]


class ImportValidationError(ValueError):
    """Messages contain schema locations and codes, never untrusted cell values."""


def _fail(location, code):
    raise ImportValidationError(f'{location}: {code}')


def _day(value, location):
    if isinstance(value, datetime):
        if value.tzinfo is not None or value.time() != datetime.min.time():
            _fail(location, 'expected_date_without_time')
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str) or not re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}', value):
        _fail(location, 'expected_ISO_date')
    try:
        date.fromisoformat(value)
    except ValueError:
        _fail(location, 'invalid_date')
    return value


def _url(value, location):
    if not isinstance(value, str) or len(value) > 2048 or any(c.isspace() or ord(c) < 32 for c in value):
        _fail(location, 'invalid_public_https_url')
    try:
        parsed = urlsplit(value)
        port = parsed.port
        query = parse_qsl(parsed.query, keep_blank_values=True, strict_parsing=True)
    except ValueError:
        _fail(location, 'invalid_public_https_url')
    if (parsed.scheme != 'https' or not parsed.hostname or parsed.username is not None
            or parsed.password is not None or parsed.fragment or port not in (None, 443)
            or '\\' in value):
        _fail(location, 'invalid_public_https_url')
    safe_keys = {'year', 'format', 'download', 'id', 'version', 'document', 'lang'}
    if any(key not in safe_keys or not re.fullmatch(r'[A-Za-z0-9._-]{1,64}', item)
           for key, item in query):
        _fail(location, 'query_not_public_allowlist')
    return value


def _number(value, decimal_format, location, *, minimum=None, positive=False):
    if isinstance(value, bool) or value is None or value == '':
        _fail(location, 'expected_finite_number')
    if isinstance(value, str):
        separator = r'\.' if decimal_format == 'dot' else ','
        if len(value) > 100 or not re.fullmatch(r'[+-]?[0-9]+(?:' + separator + r'[0-9]+)?', value):
            _fail(location, 'invalid_decimal_format')
        value = value.replace(',', '.')
    elif not isinstance(value, (int, float, Decimal)):
        _fail(location, 'expected_finite_number')
    try:
        number = float(Decimal(str(value)))
    except (ValueError, InvalidOperation, OverflowError):
        _fail(location, 'expected_finite_number')
    if not math.isfinite(number) or (positive and number <= 0) or (minimum is not None and number < minimum):
        _fail(location, 'number_out_of_range')
    return number


def _integer(value, location, *, minimum=0, maximum=None):
    if isinstance(value, str):
        if not re.fullmatch(r'[0-9]{1,12}', value):
            _fail(location, 'expected_integer')
        value = int(value)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        _fail(location, 'expected_integer')
    limit = maximum if maximum is not None else 999_999_999_999
    if value < minimum or value > limit or not math.isfinite(value) or value != int(value):
        _fail(location, 'integer_out_of_range')
    return int(value)


def _choice(value, choices, location):
    if not isinstance(value, str) or value not in choices:
        _fail(location, 'unsupported_value')
    return value


def _read_file(root, name):
    path = root / name
    if path.is_symlink() or (hasattr(path, 'is_junction') and path.is_junction()):
        _fail(name, 'linked_input_not_allowed')
    if not path.is_file() or not path.resolve().is_relative_to(root):
        _fail(name, 'invalid_input_path')
    if path.stat().st_size > MAX_FILE_BYTES:
        _fail(name, 'file_too_large')
    try:
        body = path.read_bytes()
    except OSError:
        _fail(name, 'input_unreadable')
    if len(body) > MAX_FILE_BYTES:
        _fail(name, 'file_too_large')
    return body


def _json_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            _fail('source.json', 'duplicate_key')
        result[key] = value
    return result


def _manifest(content, cutoff):
    try:
        source = json.loads(content.decode('utf-8-sig'), object_pairs_hook=_json_object,
                            parse_constant=lambda _: _fail('source.json', 'non_finite_value'))
    except (ValueError, UnicodeError) as error:
        if isinstance(error, ImportValidationError):
            raise
        _fail('source.json', 'invalid_json')
    required = {'schema_version', 'source_name', 'source_url', 'retrieved_at', 'currency',
                'price_basis', 'cash_share_basis', 'history_start', 'history_end',
                'tax_basis', 'decimal_format', 'tickers'}
    optional = {'price_basis_date', 'cash_share_basis_date', 'csv_delimiter'}
    if not isinstance(source, dict) or not required.issubset(source) or set(source) - required - optional:
        _fail('source.json', 'missing_or_unknown_fields')
    if type(source['schema_version']) is not int or source['schema_version'] != 1:
        _fail('source.schema_version', 'unsupported_value')
    name = source['source_name']
    if not isinstance(name, str) or not name.strip() or len(name) > 200 or any(ord(c) < 32 for c in name):
        _fail('source.source_name', 'invalid_name')
    _url(source['source_url'], 'source.source_url')
    retrieved = source['retrieved_at']
    try:
        if not isinstance(retrieved, str) or 'T' not in retrieved:
            raise ValueError
        timestamp = datetime.fromisoformat(retrieved.replace('Z', '+00:00'))
        if timestamp.tzinfo is None or timestamp > datetime.now(timezone.utc):
            raise ValueError
    except ValueError:
        _fail('source.retrieved_at', 'expected_nonfuture_timezone_timestamp')
    source['retrieved_at'] = timestamp.astimezone(timezone.utc).isoformat()
    _choice(source['currency'], {'BRL'}, 'source.currency')
    _choice(source['tax_basis'], {'gross_before_individual_taxes'}, 'source.tax_basis')
    _choice(source['decimal_format'], {'dot', 'comma'}, 'source.decimal_format')
    source['csv_delimiter'] = _choice(source.get('csv_delimiter', ','), {',', ';'}, 'source.csv_delimiter')
    for field in ('history_start', 'history_end'):
        source[field] = _day(source[field], 'source.' + field)
    if source['history_start'] > source['history_end'] or source['history_end'] > cutoff:
        _fail('source.history_end', 'invalid_history_window')
    if source['history_end'] > timestamp.astimezone(timezone.utc).date().isoformat():
        _fail('source.history_end', 'history_after_retrieval')
    for field, choices, adjusted in [('price_basis', {'raw', 'split_adjusted'}, 'split_adjusted'),
                                     ('cash_share_basis', {'record_date', 'adjusted'}, 'adjusted')]:
        _choice(source[field], choices, 'source.' + field)
        basis_date = source.get(field + '_date')
        if source[field] == adjusted:
            if _day(basis_date, 'source.' + field + '_date') != cutoff:
                _fail('source.' + field + '_date', 'adjusted_basis_must_equal_cutoff')
            if cutoff > source['retrieved_at'][:10]:
                _fail('source.' + field + '_date', 'basis_after_retrieval')
        elif basis_date not in (None, ''):
            _fail('source.' + field + '_date', 'unadjusted_basis_has_no_adjustment_date')
    tickers = source['tickers']
    if not isinstance(tickers, dict) or not tickers or len(tickers) > 2000:
        _fail('source.tickers', 'expected_nonempty_ticker_map')
    for ticker, spec in tickers.items():
        if not TICKER.fullmatch(ticker) or not isinstance(spec, dict) or set(spec) - {'coverage', 'cvm', 'share_class'}:
            _fail('source.tickers', 'invalid_identity')
        if ('cvm' in spec) != ('share_class' in spec):
            _fail('source.tickers', 'incomplete_identity')
        if 'cvm' in spec:
            spec['cvm'] = _integer(spec['cvm'], 'source.tickers.cvm', minimum=1)
            _choice(spec['share_class'], SHARE_CLASSES, 'source.tickers.share_class')
        coverage = spec.get('coverage')
        if not isinstance(coverage, dict) or set(coverage) - set(COLUMNS):
            _fail('source.tickers.coverage', 'invalid_table_map')
        for table, interval in coverage.items():
            if not isinstance(interval, dict) or set(interval) != {'start', 'end', 'complete'}:
                _fail('source.tickers.coverage', 'invalid_interval_fields')
            start = _day(interval['start'], 'source.tickers.coverage.start')
            end = _day(interval['end'], 'source.tickers.coverage.end')
            if not source['history_start'] <= start <= end <= source['history_end'] or type(interval['complete']) is not bool:
                _fail('source.tickers.coverage', 'invalid_interval')
    return source


def _csv_rows(body, table, delimiter):
    try:
        reader = csv.reader(io.StringIO(body.decode('utf-8-sig'), newline=''), delimiter=delimiter, strict=True)
        headers = next(reader, None)
        if headers != list(COLUMNS[table]):
            _fail(table, 'headers_must_match_template')
        rows = []
        for index, row in enumerate(reader, 2):
            if not row or all(value == '' for value in row):
                continue
            if len(row) != len(headers):
                _fail(f'{table}.row_{index}', 'column_count_mismatch')
            rows.append((index, dict(zip(headers, row))))
            if len(rows) > MAX_ROWS:
                _fail(table, 'too_many_rows')
        return rows
    except (csv.Error, UnicodeError):
        _fail(table, 'invalid_utf8_csv')


def _xlsx_tables(body, table=None):
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            entries = archive.infolist()
            if sum(row.file_size for row in entries) > MAX_FILE_BYTES * 2 or len(entries) > 1000:
                _fail('xlsx', 'workbook_too_large')
            if any('externallinks/' in row.filename.lower().replace('\\', '/')
                   or 'vbaproject' in row.filename.lower()
                   or '..' in PurePosixPath(row.filename.replace('\\', '/')).parts
                   or PurePosixPath(row.filename.replace('\\', '/')).is_absolute()
                   or re.match(r'^[A-Za-z]:', row.filename) for row in entries):
                _fail('xlsx', 'unsafe_workbook_package')
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
        try:
            if table is not None:
                if len(workbook.worksheets) != 1:
                    _fail(table, 'single_table_workbook_requires_one_sheet')
                sheets = {table: workbook.worksheets[0]}
            else:
                if set(workbook.sheetnames) - set(COLUMNS):
                    _fail('xlsx', 'unexpected_sheet')
                sheets = {name: workbook[name] for name in workbook.sheetnames}
            result = {}
            for name, sheet in sheets.items():
                if sheet.max_column is not None and sheet.max_column > len(COLUMNS[name]):
                    _fail(name, 'unexpected_excel_columns')
                rows = []
                for index, cells in enumerate(sheet.iter_rows(), 1):
                    if any(cell.data_type in ('f', 'e') for cell in cells):
                        _fail(f'{name}.row_{index}', 'formula_or_excel_error')
                    values = [cell.value for cell in cells]
                    while values and values[-1] is None and len(values) > len(COLUMNS[name]):
                        values.pop()
                    if index == 1:
                        if values != list(COLUMNS[name]):
                            _fail(name, 'headers_must_match_template')
                    elif any(value is not None for value in values):
                        if len(values) != len(COLUMNS[name]):
                            _fail(f'{name}.row_{index}', 'column_count_mismatch')
                        rows.append((index, dict(zip(COLUMNS[name], values))))
                    if index > MAX_ROWS + 1:
                        _fail(name, 'too_many_rows')
                if sheet.max_row is None or sheet.max_row < 1:
                    _fail(name, 'headers_must_match_template')
                result[name] = rows
            return result
        finally:
            workbook.close()
    except Exception as error:
        if isinstance(error, ImportValidationError):
            raise
        _fail('xlsx', 'invalid_workbook')


def _normalize(table, row, source, cutoff, location):
    ticker = row['ticker']
    if not isinstance(ticker, str) or ticker not in source['tickers']:
        _fail(location + '.ticker', 'undeclared_ticker_identity')
    result = {'ticker': ticker}
    dates = {'prices': ('date',), 'payments': ('record_date', 'payment_date'),
             'actions': ('effective_date',),
             'fundamentals': ('fiscal_end', 'published_at', 'share_basis_date')}[table]
    for field in dates:
        result[field] = _day(row[field], location + '.' + field)
    decimal = source['decimal_format']
    if table == 'prices':
        result['close'] = _number(row['close'], decimal, location + '.close', positive=True)
        anchor = result['date']
        key = (ticker, anchor)
    elif table == 'payments':
        if result['record_date'] > result['payment_date']:
            _fail(location, 'payment_before_record_date')
        result['gross_per_share'] = _number(row['gross_per_share'], decimal, location + '.gross_per_share', positive=True)
        result['label'] = _choice(row['label'], {'DIVIDENDO', 'JCP', 'CUPOM'}, location + '.label')
        anchor = result['payment_date']
        key = (ticker, result['record_date'], anchor, result['label'])
    elif table == 'actions':
        result['numerator'] = _integer(row['numerator'], location + '.numerator', minimum=1)
        result['denominator'] = _integer(row['denominator'], location + '.denominator', minimum=1)
        result['type'] = _choice(row['type'], {'split', 'reverse_split', 'bonus'}, location + '.type')
        greater = result['numerator'] > result['denominator']
        if result['numerator'] == result['denominator'] or (result['type'] == 'reverse_split') == greater:
            _fail(location, 'action_ratio_disagrees_with_type')
        anchor = result['effective_date']
        key = (ticker, anchor, result['type'])
    else:
        result['cvm'] = _integer(row['cvm'], location + '.cvm', minimum=1)
        result['share_class'] = _choice(row['share_class'], SHARE_CLASSES, location + '.share_class')
        result['sector_code'] = _integer(row['sector_code'], location + '.sector_code', maximum=9)
        result['scope'] = _choice(row['scope'], {'consolidated', 'individual'}, location + '.scope')
        result['source_ref'] = _url(row['source_ref'], location + '.source_ref')
        result['eps'] = _number(row['eps'], decimal, location + '.eps')
        free_cash = row['free_cash_per_share']
        if free_cash in (None, '') and result['sector_code'] == 2:
            result['free_cash_per_share'] = None
        else:
            result['free_cash_per_share'] = _number(free_cash, decimal, location + '.free_cash_per_share')
        result['liabilities_to_assets'] = _number(row['liabilities_to_assets'], decimal, location + '.liabilities_to_assets', minimum=0)
        if (result['fiscal_end'] > result['published_at'] or result['share_basis_date'] > cutoff
                or result['share_basis_date'] > source['retrieved_at'][:10]):
            _fail(location, 'invalid_financial_dates')
        anchor = result['published_at']
        key = (ticker, result['fiscal_end'], anchor)
    future = table == 'payments' and anchor > cutoff
    if not future:
        if anchor > cutoff:
            _fail(location, 'future_training_data')
        if not source['history_start'] <= anchor <= source['history_end']:
            _fail(location, 'outside_declared_history')
        coverage = source['tickers'][ticker]['coverage'].get(table)
        if coverage is None:
            _fail(location, 'missing_ticker_table_coverage')
        if not coverage['start'] <= anchor <= coverage['end']:
            _fail(location, 'outside_ticker_coverage')
        if anchor > source['retrieved_at'][:10]:
            _fail(location, 'observation_after_retrieval')
    return result, key, future


def load_import_bundle(folder: str | Path, cutoff: str | date) -> ImportBundle:
    """Read fixed local filenames only; normalize without modifying input or bases."""
    try:
        root = Path(folder).resolve(strict=True)
    except OSError:
        _fail('bundle', 'input_unreadable')
    if not root.is_dir():
        _fail('bundle', 'expected_directory')
    cutoff = _day(cutoff, 'cutoff')
    if cutoff > datetime.now(timezone.utc).date().isoformat():
        _fail('cutoff', 'future_cutoff')
    manifest_bytes = _read_file(root, 'source.json')
    source = _manifest(manifest_bytes, cutoff)
    files = [{'file': 'source.json', 'bytes': len(manifest_bytes), 'sha256': hashlib.sha256(manifest_bytes).hexdigest()}]
    tables = {}
    for table in COLUMNS:
        for suffix in ('.csv', '.xlsx'):
            name = table + suffix
            if not (root / name).exists() and not (root / name).is_symlink():
                continue
            if table in tables:
                _fail(table, 'ambiguous_table_files')
            body = _read_file(root, name)
            tables[table] = (_csv_rows(body, table, source['csv_delimiter']) if suffix == '.csv'
                             else _xlsx_tables(body, table)[table])
            files.append({'file': name, 'bytes': len(body), 'sha256': hashlib.sha256(body).hexdigest()})
    if (root / 'tables.xlsx').exists() or (root / 'tables.xlsx').is_symlink():
        body = _read_file(root, 'tables.xlsx')
        workbook_tables = _xlsx_tables(body)
        if set(workbook_tables) & set(tables):
            _fail('bundle', 'ambiguous_table_files')
        tables.update(workbook_tables)
        files.append({'file': 'tables.xlsx', 'bytes': len(body), 'sha256': hashlib.sha256(body).hexdigest()})
    result = {'schema_version': 1, 'cutoff': cutoff, **{table: [] for table in COLUMNS}, 'future_payments': []}
    identities = {}
    for table in COLUMNS:
        seen = set()
        for index, row in tables.get(table, []):
            normalized, key, future = _normalize(table, row, source, cutoff, f'{table}.row_{index}')
            if key in seen:
                _fail(f'{table}.row_{index}', 'duplicate_row_identity')
            seen.add(key)
            if table == 'fundamentals':
                ticker = normalized['ticker']
                identity = (normalized['cvm'], normalized['share_class'], normalized['scope'])
                spec = source['tickers'][ticker]
                if (ticker in identities and identities[ticker] != identity) or (
                        'cvm' in spec and (spec['cvm'], spec['share_class']) != identity[:2]):
                    _fail(f'{table}.row_{index}', 'mixed_ticker_identity_or_scope')
                identities[ticker] = identity
            result['future_payments' if future else table].append(normalized)
    sort_fields = {'prices': ('ticker', 'date'),
                   'payments': ('ticker', 'payment_date', 'record_date', 'label'),
                   'future_payments': ('ticker', 'payment_date', 'record_date', 'label'),
                   'actions': ('ticker', 'effective_date', 'type'),
                   'fundamentals': ('ticker', 'published_at', 'fiscal_end')}
    for table, fields in sort_fields.items():
        result[table].sort(key=lambda row: tuple(row[field] for field in fields))
    table_counts = {table: Counter(row['ticker'] for row in result[table]) for table in (*COLUMNS, 'future_payments')}
    per_ticker = {}
    for ticker, spec in sorted(source['tickers'].items()):
        counts = {table: table_counts[table][ticker] for table in (*COLUMNS, 'future_payments')}
        per_ticker[ticker] = {'counts': counts, 'declared_coverage': spec['coverage'],
                              'tables_without_rows': [table for table in COLUMNS if counts[table] == 0],
                              'tables_without_declared_coverage': [table for table in COLUMNS if table not in spec['coverage']]}
        if ticker in identities:
            cvm, share_class, scope = identities[ticker]
            per_ticker[ticker]['identity'] = {'cvm': cvm, 'share_class': share_class, 'scope': scope}
    result['provenance'] = {'source': source, 'input_files': sorted(files, key=lambda row: row['file']),
                            'missing_tables': [table for table in COLUMNS if table not in tables],
                            'empty_tables': [table for table in COLUMNS if table in tables and not tables[table]],
                            'per_ticker': per_ticker, 'completeness_certified': False,
                            'cash_imputed': False, 'share_adjustments_applied': False}
    return result


def _canonical(value):
    return (json.dumps(value, ensure_ascii=True, allow_nan=False, sort_keys=True, separators=(',', ':')) + '\n').encode('utf8')


def _output_directory(output, workspace, input_folder=None):
    root = Path(workspace).resolve(strict=True)
    output = Path(output)
    target = (root / output).resolve()
    if not root.is_dir() or not target.is_relative_to(root):
        _fail('output', 'outside_explicit_workspace')
    if input_folder is not None and target.is_relative_to(Path(input_folder).resolve(strict=True)):
        _fail('output', 'must_be_outside_input_bundle')
    target.mkdir(parents=True, exist_ok=True)
    return target


def _write_new(path, body):
    if path.is_symlink() or (hasattr(path, 'is_junction') and path.is_junction()):
        _fail('output', 'linked_output_not_allowed')
    try:
        with path.open('xb') as stream:
            stream.write(body)
            stream.flush()
            os.fsync(stream.fileno())
        path.chmod(0o444)
    except FileExistsError:
        if not path.is_file() or path.read_bytes() != body:
            _fail('output', 'existing_frozen_content_mismatch')


def freeze_import_bundle(folder, cutoff, output, *, workspace) -> Path:
    """Validate and freeze under explicit workspace/<output>/<content SHA-256>."""
    target = _output_directory(output, workspace, folder)
    bundle = load_import_bundle(folder, cutoff)
    normalized = _canonical(bundle)
    digest = hashlib.sha256(normalized).hexdigest()
    frozen = target / digest
    if frozen.is_symlink() or (hasattr(frozen, 'is_junction') and frozen.is_junction()):
        _fail('output', 'linked_output_not_allowed')
    frozen.mkdir(exist_ok=True)
    _write_new(frozen / 'normalized.json', normalized)
    _write_new(frozen / 'provenance.json', _canonical(bundle['provenance']))
    return frozen


TEMPLATE_README = """Importacao local de historicos RendaFutura

Arquivos vazios: nenhum ativo ou pagamento de exemplo foi inventado.
Preencha source.json com a fonte publica real e os intervalos documentados.
schema_version: 1; currency: BRL; tax_basis: gross_before_individual_taxes.
source_url e source_ref: HTTPS publico, sem usuario, senha ou fragmento.
Queries aceitas apenas: year, format, download, id, version, document, lang;
valores alfanumericos simples. Nao inclua tokens nem chaves de API.
retrieved_at: timestamp ISO com fuso, nunca futuro.
history_start/history_end: datas ISO YYYY-MM-DD, ate o cutoff da validacao.
decimal_format: dot ou comma; sem separadores de milhares ou notacao cientifica.
csv_delimiter: virgula ou ponto e virgula. CSV obrigatoriamente UTF-8.
Datas: ISO YYYY-MM-DD, ou celulas de data nativas do Excel sem horario.
Nao use formulas, links externos, linhas duplicadas ou datas futuras no treino.

price_basis: raw ou split_adjusted. Para split_adjusted, price_basis_date deve
ser exatamente o cutoff. Para raw, deixe price_basis_date vazio.
cash_share_basis: record_date ou adjusted. Para adjusted, cash_share_basis_date
deve ser exatamente o cutoff. Para record_date, deixe essa data vazia.
O importador preserva as bases declaradas; nao converte valores por acao.

tickers: objeto cujas chaves sao os tickers reais presentes nos arquivos.
Cada ticker tem coverage: objeto por tabela, com start, end e complete booleano.
Os intervalos ficam dentro de history_start/history_end. Para fundamentals,
cobertura usa published_at; para payments, payment_date historico.
Pode informar cvm (inteiro positivo) e share_class juntos para conferir identidade.
Classes: ON, PN, PNA, PNB, PNC, PND, UNIT. Nao misture emissor, classe ou scope
para um mesmo ticker. Cobertura declarada nao certifica integralidade.

Tabelas opcionais: prices, payments, actions, fundamentals. Ausencia e tabela
vazia sao reportadas, nunca convertidas em caixa zero ou dados completos.
Use uma tabela por CSV/XLSX, ou tables.xlsx com abas nomeadas como as tabelas.
Nao forneca a mesma tabela em dois arquivos. Preserve os cabecalhos e sua ordem.

payments: record_date e payment_date obrigatorios. Data desconhecida bloqueia
o evento, nunca vira zero. Rotulos: DIVIDENDO, JCP, CUPOM. Principal, amortizacao
e resgate nao sao renda. Pagamentos datados apos cutoff ficam em future_payments,
separados de payments; nao comprovam pagamento e nao entram como caixa recebido.
actions: type split, reverse_split ou bonus. numerator/denominator sao inteiros
positivos e representam quantidade total depois / antes do evento (inclusive bonus).
fundamentals: fiscal_end <= published_at <= cutoff; share_basis_date <= cutoff.
eps e free_cash_per_share podem ser negativos. FCF vazio somente sector_code 2;
nesse caso vira null, nunca zero. sector_code inteiro entre 0 e 9.
scope: consolidated ou individual. liabilities_to_assets >= 0, pode exceder 1.

Interface Python: load_import_bundle(folder, cutoff) retorna dict com schema_version,
cutoff, prices, payments, future_payments, actions, fundamentals e provenance.
Datas normalizadas sao strings ISO; valores monetarios floats finitos; identidades
e fatores sao inteiros. Provenance inclui source, input_files (nome/bytes/sha256),
missing_tables, empty_tables e per_ticker (contagens, identidade e cobertura).
O importador nao certifica 60 meses comuns, adequacao ao modelo ou dados point-in-time.
A integracao deve conferir essas condicoes e reconciliar precos, caixa e fundamentos.

CLI: python -m recommender.historical_import validate PASTA --cutoff YYYY-MM-DD
--workspace RAIZ_DE_SAIDA --output SUBPASTA_DE_SAIDA
Saida: SUBPASTA_DE_SAIDA/SHA256/normalized.json e provenance.json. Conteudo
congelado e verificado por hash; nao sobrescreve versoes diferentes nem altera input.

Tesouro: arquivo OFICIAL de cupons (nao misture arquivos de vencimentos/resgates):
https://www.tesourotransparente.gov.br/ckan/dataset/f30db6e4-6123-416c-b094-be8dfc823601/resource/b04812f4-c453-4298-a12e-0ad190998243/download/cupomjurostesourodireto.csv
Cabecalho original, separado por ponto e virgula:
Tipo Titulo;Vencimento do Titulo;Data Resgate;PU;Quantidade;Valor
Preserve o CSV original. Ele NAO tem o esquema payments acima: a integracao do
Tesouro deve validar tipo+vencimento, Data Resgate e PU, conferir Quantidade*PU
contra Valor e excluir principal. Este importador nao adivinha essa conversao.
"""


def emit_templates(folder, *, format='csv') -> Path:
    """Create a new template directory only; never overwrite user data."""
    _choice(format, {'csv', 'xlsx'}, 'templates.format')
    target = Path(folder)
    if target.exists() or target.is_symlink():
        _fail('templates', 'destination_already_exists')
    target.mkdir(parents=True)
    manifest = {'schema_version': 1, 'source_name': '', 'source_url': '', 'retrieved_at': '',
                'currency': 'BRL', 'price_basis': '', 'price_basis_date': '',
                'cash_share_basis': '', 'cash_share_basis_date': '', 'history_start': '', 'history_end': '',
                'tax_basis': 'gross_before_individual_taxes', 'decimal_format': 'dot',
                'csv_delimiter': ',', 'tickers': {}}
    (target / 'source.json').write_bytes(_canonical(manifest))
    (target / 'README.md').write_text(TEMPLATE_README, encoding='utf8')
    if format == 'csv':
        for table, headers in COLUMNS.items():
            with (target / (table + '.csv')).open('x', encoding='utf8', newline='') as stream:
                csv.writer(stream, lineterminator='\n').writerow(headers)
    else:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for table, headers in COLUMNS.items():
            workbook.create_sheet(table).append(headers)
        workbook.save(target / 'tables.xlsx')
        workbook.close()
    return target.resolve()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    templates = commands.add_parser('templates')
    templates.add_argument('folder', type=Path)
    templates.add_argument('--format', choices=('csv', 'xlsx'), default='csv')
    validate = commands.add_parser('validate')
    validate.add_argument('folder', type=Path)
    validate.add_argument('--cutoff', required=True)
    validate.add_argument('--workspace', type=Path, required=True)
    validate.add_argument('--output', type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        if args.command == 'templates':
            emit_templates(args.folder, format=args.format)
            result = {'status': 'templates_created', 'format': args.format}
        else:
            frozen = freeze_import_bundle(args.folder, args.cutoff, args.output, workspace=args.workspace)
            provenance = json.loads((frozen / 'provenance.json').read_text(encoding='utf8'))
            result = {'status': 'validated', 'bundle_sha256': frozen.name,
                      'files': ['normalized.json', 'provenance.json'],
                      'missing_tables': provenance['missing_tables'], 'empty_tables': provenance['empty_tables'],
                      'row_counts': {table: sum(row['counts'][table] for row in provenance['per_ticker'].values())
                                     for table in (*COLUMNS, 'future_payments')},
                      'completeness_certified': False}
    except ImportValidationError as error:
        print(json.dumps({'status': 'invalid', 'error': str(error)}))
        return 2
    except (OSError, ValueError, TypeError, OverflowError):
        print(json.dumps({'status': 'invalid', 'error': 'input_or_output_unreadable'}))
        return 2
    print(json.dumps(result))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
